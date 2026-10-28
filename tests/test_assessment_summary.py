import io
from datetime import date, datetime, timezone
from decimal import Decimal
from uuid import UUID

import openpyxl
import pytest

from api.services.assessment_summary import (
    EMPTY_PROJECT,
    EMPTY_RANGE,
    MonthError,
    EMPTY_VIEW,
    build_tailored_workbook,
    build_workbook,
    REMEASURE_BANNER,
    parse_month,
    project_total_label,
    range_label,
    summary_filename,
    summary_months,
    summary_payload,
    summary_project_total,
    summary_sites,
)

# ---------------------------------------------------------------------------
# summary_months and build_workbook: rows as list_assessment_flags returns them
# ---------------------------------------------------------------------------

RATES = {"cost_per_sf_4in": Decimal("11.54"), "cost_per_sf_7in": Decimal("12.64")}
MAIN = UUID("20000000-0000-4000-8000-000000000001")
ELM = UUID("20000000-0000-4000-8000-000000000002")
OCT_15 = datetime(2026, 10, 15, 18, 0, tzinfo=timezone.utc)
NOV_2 = datetime(2026, 11, 2, 15, 0, tzinfo=timezone.utc)
PROJECT = {"project_id": "SE384", "project_name": "Sidewalks Queens", "registration_code": "2024123456",
           "borough": "Queens", **RATES}


def flag(property_id: UUID, thickness: str, gross: str, quantity: str, assessable: bool = True,
         constructed: date = None, approved: datetime = OCT_15, address: str = None) -> dict:
    """
    Build a row as list_assessment_flags returns one.
    Takes its property, thickness, gross area, quantity, whether surveyed assessable, the property's Date Constructed, when its IDR was approved and the address.
    Returns the row.
    """
    return {"idr_id": UUID("1d000000-0000-4000-8000-000000000001"), "re_signed_at": approved,
            "property_id": property_id, "property_kind": "sidewalk",
            "address": address or ("138-22 227 STREET" if property_id == MAIN else "350 Main St, Brooklyn"),
            "intersection_street_1": None, "intersection_street_2": None, "quadrant": None, "facing_street": None,
            "date_constructed": constructed, "surveyed_thickness": thickness, "surveyed_is_assessable": assessable,
            "computed_quantity": Decimal(gross), "quantity": Decimal(quantity)}


# SE384 as it stands: 138-22 227 Street (a non-assessable 4", an assessable 4", 7") and 350 Main St (4", 7")
SE384 = [flag(MAIN, "4in", "90.25", "90.25", assessable=False), flag(MAIN, "4in", "160.69", "113.19"),
         flag(MAIN, "7in", "171.81", "94.14"), flag(ELM, "4in", "180", "180"), flag(ELM, "7in", "45", "45")]


def tab(content: bytes, name: str):
    """
    Open a workbook and find a tab.
    Takes the file's bytes and the tab's name.
    Returns the worksheet.
    """
    return openpyxl.load_workbook(io.BytesIO(content))[name]


def values(sheet, row: int) -> list:
    """
    Read one row's thirteen cells.
    Takes the sheet and the row number.
    Returns the values.
    """
    return [sheet.cell(row, column).value for column in range(1, 14)]


class TestGrouping:
    def test_se384_is_one_month_with_its_two_properties(self):
        months = summary_months(SE384, RATES)
        assert list(months) == [date(2026, 10, 1)]
        main, elm = months[date(2026, 10, 1)]
        assert (main.four_assessable, main.four_non_assessable, main.four_total) == (
            Decimal("113.19"), Decimal("137.75"), Decimal("250.94"))
        assert main.four_amount == Decimal("113.19") * Decimal("11.54")
        assert round(main.grand_total, 2) == Decimal("2496.14")
        assert round(elm.grand_total, 2) == Decimal("2646.00") and elm.four_non_assessable == 0
        assert main.remarks == [] and elm.remarks == []

    def test_rows_are_ordered_by_address_any_case(self):
        flags = [flag(ELM, "4in", "1", "1", address="b street"), flag(MAIN, "4in", "1", "1", address="A Street")]
        assert [row.address for row in summary_months(flags, RATES)[date(2026, 10, 1)]] == ["A Street", "b street"]

    def test_two_months_are_two_lists_with_nothing_shared(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "20", "20", approved=NOV_2)]
        months = summary_months(flags, RATES)
        assert list(months) == [date(2026, 10, 1), date(2026, 11, 1)]
        assert months[date(2026, 10, 1)][0].four_assessable == 10
        assert months[date(2026, 11, 1)][0].four_assessable == 20

    def test_the_date_constructed_decides_the_month_over_the_approval(self):
        months = summary_months([flag(MAIN, "4in", "10", "10", constructed=date(2026, 9, 28))], RATES)
        assert list(months) == [date(2026, 9, 1)]
        assert months[date(2026, 9, 1)][0].used_approval_date is False

    def test_without_one_the_approval_day_in_new_york_decides(self):
        # 02:00 UTC on Nov 1 is Oct 31 in New York
        months = summary_months([flag(MAIN, "4in", "10", "10", approved=datetime(2026, 11, 1, 2, 0, tzinfo=timezone.utc))],
                                RATES)
        assert list(months) == [date(2026, 10, 1)]
        assert months[date(2026, 10, 1)][0].used_approval_date is True

    def test_7in_assessed_as_4in_counts_in_the_4in_columns_with_a_remark(self):
        row = summary_months([flag(MAIN, "7in_as_4in", "50", "50")], RATES)[date(2026, 10, 1)][0]
        assert (row.four_assessable, row.four_total, row.seven_total) == (Decimal(50), Decimal(50), 0)
        assert row.four_amount == Decimal(50) * Decimal("11.54")
        assert row.remarks == ['50.00 SF built at 7" billed at 4"']

    def test_a_non_assessable_area_is_in_non_assessable_and_not_in_the_amount(self):
        row = summary_months([flag(MAIN, "4in", "25", "20")], RATES)[date(2026, 10, 1)][0]
        assert (row.four_assessable, row.four_non_assessable) == (Decimal(20), Decimal(5))
        assert row.four_amount == Decimal(20) * Decimal("11.54")

    def test_a_missing_rate_leaves_its_amount_blank_with_a_remark(self):
        row = summary_months(SE384[:3], {**RATES, "cost_per_sf_4in": None})[date(2026, 10, 1)][0]
        assert row.four_amount is None and row.seven_amount is not None
        assert row.grand_total == row.seven_amount
        assert row.remarks == ['4" rate not set']

    def test_the_range_keeps_its_months_only(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "20", "20", approved=NOV_2)]
        assert list(summary_months(flags, RATES, first=date(2026, 11, 1))) == [date(2026, 11, 1)]
        assert list(summary_months(flags, RATES, last=date(2026, 10, 1))) == [date(2026, 10, 1)]
        assert summary_months(flags, RATES, last=date(2026, 9, 1)) == {}


class TestSites:
    def test_a_site_approved_in_two_months_is_one_row_summed(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "20", "20", approved=NOV_2), flag(ELM, "7in", "5", "5")]
        rows = summary_sites(flags, RATES)
        assert [row.address for row in rows] == ["138-22 227 STREET", "350 Main St, Brooklyn"]
        assert rows[0].four_assessable == Decimal("30") and rows[0].four_amount == Decimal("30") * RATES["cost_per_sf_4in"]
        assert rows[1].seven_total == Decimal("5")

    def test_the_range_is_of_days_and_sums_only_the_flags_inside_it(self):
        flags = [flag(MAIN, "4in", "10", "10", constructed=date(2026, 10, 3)),
                 flag(MAIN, "4in", "20", "20", approved=NOV_2)]
        assert summary_sites(flags, RATES, first=date(2026, 10, 3), last=date(2026, 10, 3))[0].four_total == Decimal("10")
        assert summary_sites(flags, RATES, first=date(2026, 10, 4))[0].four_total == Decimal("20")
        assert summary_sites(flags, RATES, last=date(2026, 10, 2)) == []

    def test_it_keeps_only_the_sites_asked_for(self):
        assert [row.property_id for row in summary_sites(SE384, RATES, property_ids={ELM})] == [ELM]

    def test_a_row_carries_what_the_page_searches_by(self):
        flags = [{**flag(MAIN, "4in", "10", "10"), "block": "123", "lot": "45", "violation_no": "V-9",
                  "community_board": "QN12"}]
        row = summary_sites(flags, RATES)[0]
        assert (row.kind, row.block, row.lot, row.violation_no, row.community_board) == (
            "sidewalk", "123", "45", "V-9", "QN12")

    def test_the_payload_rounds_to_the_cent_and_totals_the_rows(self):
        payload = summary_payload(summary_sites(SE384, RATES), RATES, date(2026, 10, 1), None)
        assert payload["range"] == {"from": date(2026, 10, 1), "to": None}
        assert payload["cost_per_sf_4in"] == 11.54 and len(payload["sites"]) == 2
        main = payload["sites"][0]
        assert main["display_address"] == "138-22 227 STREET" and main["four_assessable"] == 113.19
        assert main["four_amount"] == 1306.21 and main["grand_total"] == 2496.14
        assert payload["totals"]["four_assessable"] == 293.19 and payload["totals"]["grand_total"] == 5142.14

    def test_without_a_rate_its_amounts_are_none_in_the_payload(self):
        rates = {**RATES, "cost_per_sf_7in": None}
        payload = summary_payload(summary_sites(SE384, rates), rates, None, None)
        assert payload["sites"][0]["seven_amount"] is None and payload["totals"]["seven_amount"] is None
        assert payload["cost_per_sf_7in"] is None


class TestTailoredWorkbook:
    @pytest.mark.parametrize("first, last, name", [
        (date(2026, 10, 1), date(2026, 10, 31), "2026-10-01 to 2026-10-31"), (date(2026, 10, 1), None, "From 2026-10-01"),
        (None, date(2026, 10, 31), "Through 2026-10-31"), (None, None, "All dates")])
    def test_one_tab_named_for_the_range(self, first, last, name):
        assert range_label(first, last) == name
        content = build_tailored_workbook(PROJECT, summary_sites(SE384, RATES), first, last)
        book = openpyxl.load_workbook(io.BytesIO(content))
        assert book.sheetnames == [name] and book[name]["A2"].value == name

    def test_it_holds_the_sites_and_their_totals(self):
        sheet = tab(build_tailored_workbook(PROJECT, summary_sites(SE384, RATES), None, None), "All dates")
        assert values(sheet, 3)[:2] == ["Property Address", "CB"]
        assert sheet["A4"].value == "138-22 227 STREET" and sheet["A5"].value == "350 Main St, Brooklyn"
        assert sheet["A6"].value == "Totals" and round(Decimal(str(sheet["K6"].value)), 2) == Decimal("5142.14")

    def test_no_sites_says_so(self):
        sheet = tab(build_tailored_workbook(PROJECT, [], None, date(2026, 9, 30)), "Through 2026-09-30")
        assert sheet["A4"].value == EMPTY_VIEW

    def test_its_file_is_named_tailored(self):
        assert summary_filename("SE384", datetime(2026, 10, 16, 15, 0, tzinfo=timezone.utc), tailored=True) == \
            "SE384_assessment_summary_tailored_2026-10-16.xlsx"


class TestParseMonth:
    def test_a_month_is_its_first_day(self):
        assert parse_month("2026-10", "from") == date(2026, 10, 1)
        assert parse_month(None, "from") is None and parse_month("  ", "to") is None

    @pytest.mark.parametrize("value", ["10-2026", "2026-13", "2026-1", "2026/10", "Oct 2026"])
    def test_anything_else_is_refused(self, value):
        with pytest.raises(MonthError, match="from must be YYYY-MM"):
            parse_month(value, "from")


class TestWorkbook:
    def test_se384_has_one_tab_with_the_header_rows_and_totals(self):
        content = build_workbook(PROJECT, summary_months(SE384, RATES), filtered=False, total_rows=[])
        book = openpyxl.load_workbook(io.BytesIO(content))
        assert book.sheetnames == ["Oct 2026", "Project Total"]
        sheet = book["Oct 2026"]
        assert sheet["A1"].value == "Sidewalks Queens · Project ID: SE384 · Reg: 2024123456 · Borough: Queens"
        assert (sheet["A2"].value, sheet["C2"].value, sheet["G2"].value) == (
            "October 2026", '4" rate: $11.54', '7" rate: $12.64')
        assert values(sheet, 3)[:3] == ["Property Address", "CB", '4" Assessable (SF)']
        assert values(sheet, 3)[11:13] == ["Outside PL (SF)", "Remarks"]
        assert sheet.freeze_panes == "A4"
        main = values(sheet, 4)
        assert main[0] == "138-22 227 STREET" and main[1] is None  # no CB set
        assert [round(Decimal(str(v)), 2) for v in main[2:11]] == [
            Decimal("113.19"), Decimal("137.75"), Decimal("250.94"), Decimal("1306.21"), Decimal("94.14"),
            Decimal("77.67"), Decimal("171.81"), Decimal("1189.93"), Decimal("2496.14")]
        elm = values(sheet, 5)
        assert elm[0] == "350 Main St, Brooklyn" and elm[3] == "—"  # no non-assessable square feet
        totals = values(sheet, 6)
        assert totals[0] == "Totals" and totals[1] is None
        assert sheet["A6"].font.bold and sheet["A6"].fill.fgColor.rgb.endswith("D9D9D9")
        assert round(Decimal(str(totals[2])), 2) == Decimal("293.19")
        assert round(Decimal(str(totals[3])), 2) == Decimal("137.75")
        assert round(Decimal(str(totals[6])), 2) == Decimal("139.14")
        assert round(Decimal(str(totals[10])), 2) == Decimal("5142.14")
        assert sheet["C4"].number_format == "#,##0.00" and sheet["F4"].number_format == "$#,##0.00"
        assert sheet["J4"].number_format == "$#,##0.00" and sheet["K4"].number_format == "$#,##0.00"
        assert sheet["A8"].value == "2 rows used IDR approval date (no page-level Date Constructed)."
        assert len(content) < 30_000

    def test_cells_keep_full_precision(self):
        content = build_workbook(PROJECT, summary_months([flag(MAIN, "4in", "1.005", "1.005")], RATES), False, [])
        assert tab(content, "Oct 2026")["C4"].value == pytest.approx(1.005)

    def test_each_row_carries_its_propertys_community_board(self):
        flags = [{**flag(MAIN, "4in", "10", "10"), "community_board": "QN13"}]
        sheet = tab(build_workbook(PROJECT, summary_months(flags, RATES), False, []), "Oct 2026")
        assert sheet["B3"].value == "CB" and sheet["B4"].value == "QN13"

    def test_with_dates_typed_there_is_no_footer_note(self):
        flags = [flag(MAIN, "4in", "10", "10", constructed=date(2026, 10, 3))]
        sheet = tab(build_workbook(PROJECT, summary_months(flags, RATES), False, []), "Oct 2026")
        assert sheet["A7"].value is None

    def test_a_clip_at_the_property_line_has_its_column_and_a_note(self):
        flags = [{**flag(MAIN, "4in", "100", "100"), "clipped_sf": Decimal("25.50")}, flag(MAIN, "4in", "10", "10"),
                 {**flag(MAIN, "4in", "40", "40", assessable=False), "clipped_sf": Decimal("5")}]
        sheet = tab(build_workbook(PROJECT, summary_months(flags, RATES), False, []), "Oct 2026")
        row = values(sheet, 4)
        assert row[2] == pytest.approx(84.5) and row[3] == pytest.approx(65.5) and row[11] == pytest.approx(25.5)
        assert sheet["L4"].number_format == "#,##0.00" and sheet["L5"].value == pytest.approx(25.5)
        assert sheet["A8"].value.startswith("Outside PL (SF): the part of the flags outside the sidewalk strip")

    def test_without_a_clip_the_column_reads_a_dash_and_there_is_no_note(self):
        flags = [flag(MAIN, "4in", "10", "10", constructed=date(2026, 10, 3))]
        sheet = tab(build_workbook(PROJECT, summary_months(flags, RATES), False, []), "Oct 2026")
        assert sheet["L4"].value == "—" and sheet["A7"].value is None

    def test_the_page_carries_the_clip_too(self):
        flags = [{**flag(MAIN, "4in", "100", "100"), "clipped_sf": Decimal("25.50")}]
        payload = summary_payload(summary_sites(flags, RATES), RATES, None, None)
        assert payload["sites"][0]["outside_pl"] == 25.5 and payload["totals"]["outside_pl"] == 25.5
        assert payload["sites"][0]["four_assessable"] == 74.5

    def test_a_missing_rate_reads_as_a_dash_in_the_header(self):
        project = {**PROJECT, "cost_per_sf_7in": None}
        sheet = tab(build_workbook(project, summary_months(SE384, {**RATES, "cost_per_sf_7in": None}), False, []), "Oct 2026")
        assert sheet["G2"].value == '7" rate: —'
        assert sheet["J4"].value is None and sheet["M4"].value == '7" rate not set'

    def test_two_months_are_two_tabs(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "20", "20", approved=NOV_2)]
        book = openpyxl.load_workbook(io.BytesIO(build_workbook(PROJECT, summary_months(flags, RATES), False, [])))
        assert book.sheetnames == ["Oct 2026", "Nov 2026", "Project Total"]
        assert book["Nov 2026"]["A2"].value == "November 2026"

    @pytest.mark.parametrize("filtered, message", [(False, EMPTY_PROJECT), (True, EMPTY_RANGE)])
    def test_nothing_is_one_no_activity_tab(self, filtered, message):
        book = openpyxl.load_workbook(io.BytesIO(build_workbook(PROJECT, {}, filtered, [])))
        assert book.sheetnames == ["No activity", "Project Total"]
        assert book["No activity"]["A4"].value == message and book["No activity"]["A4"].font.italic
        assert book["No activity"]["A1"].value.startswith("Sidewalks Queens")


def test_the_file_is_named_for_the_project_and_the_day_in_new_york():
    # 03:00 UTC on Oct 17 is still Oct 16 in New York
    assert summary_filename("SE384", datetime(2026, 10, 17, 3, 0, tzinfo=timezone.utc)) == \
        "SE384_assessment_summary_2026-10-16.xlsx"


class TestRemeasureAndProjectTotal:
    """Slice Y3b-2: the warning for measurements waiting to be re-measured, and the Project Total tab."""

    def pending(self, flag_row: dict) -> dict:
        """A flag whose measurement waits to be re-measured."""
        return {**flag_row, "needs_remeasure": True}

    def test_a_row_is_pending_when_any_of_its_flags_measurements_is(self):
        rows = summary_sites([self.pending(flag(MAIN, "4in", "25", "25")), flag(MAIN, "4in", "10", "10"),
                              flag(ELM, "4in", "5", "5")], RATES)
        assert [(row.address, row.remeasure_pending) for row in rows] == [
            ("138-22 227 STREET", True), ("350 Main St, Brooklyn", False)]
        payload = summary_payload(rows, RATES, None, None)
        assert [site["remeasure_pending"] for site in payload["sites"]] == [True, False]
        assert payload["totals"]["remeasure_pending"] is True
        assert summary_payload(summary_sites(SE384, RATES), RATES, None, None)["totals"]["remeasure_pending"] is False

    def test_a_tab_with_a_pending_row_gets_the_banner_and_the_column(self):
        flags = [self.pending(flag(MAIN, "4in", "25", "25")), flag(ELM, "4in", "5", "5")]
        sheet = tab(build_workbook(PROJECT, summary_months(flags, RATES), False, summary_sites(flags, RATES)), "Oct 2026")
        assert sheet["A3"].value == REMEASURE_BANNER and sheet["A3"].font.bold and "A3:N3" in sheet.merged_cells
        assert sheet.cell(4, 14).value == "Remeasure Pending" and sheet.freeze_panes == "A5"
        assert [sheet.cell(row, 14).value for row in (5, 6, 7)] == ["⚠", None, "⚠"]  # the totals row: any
        assert sheet.cell(7, 1).value == "Totals"

    def test_a_tab_without_one_is_as_before(self):
        sheet = tab(build_workbook(PROJECT, summary_months(SE384, RATES), False, summary_sites(SE384, RATES)), "Oct 2026")
        assert sheet["A3"].value == "Property Address" and sheet.freeze_panes == "A4" and sheet.max_column == 13

    def test_the_project_total_is_every_month_summed_per_property_last(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "5", "5", approved=NOV_2), flag(ELM, "7in", "4", "4")]
        content = build_workbook(PROJECT, summary_months(flags, RATES), False, summary_project_total(flags, RATES))
        book = openpyxl.load_workbook(io.BytesIO(content))
        assert book.sheetnames == ["Oct 2026", "Nov 2026", "Project Total"]
        sheet = book["Project Total"]
        assert sheet["A2"].value == "Project Total"
        assert values(sheet, 4)[:6] == ["138-22 227 STREET", None, 15, "—", 15, pytest.approx(173.10)]
        assert values(sheet, 5)[0] == "350 Main St, Brooklyn" and values(sheet, 6)[0] == "Totals"
        assert values(sheet, 6)[10] == pytest.approx(223.66)  # 15 x 11.54 + 4 x 12.64

    def test_the_project_total_sums_the_same_months_as_the_tabs(self):
        flags = [flag(MAIN, "4in", "10", "10"), flag(MAIN, "4in", "5", "5", approved=NOV_2)]
        [row] = summary_project_total(flags, RATES, last=date(2026, 10, 1))
        assert row.four_total == Decimal("10")
        [row] = summary_project_total(flags, RATES, first=date(2026, 11, 1))
        assert row.four_total == Decimal("5")
        assert project_total_label(date(2026, 10, 1), date(2026, 12, 1)) == "Project Total, October 2026 to December 2026"
        assert project_total_label(date(2026, 10, 1), None) == "Project Total from October 2026"
        assert project_total_label(None, date(2026, 12, 1)) == "Project Total through December 2026"

    @pytest.mark.parametrize("filtered, message", [(False, EMPTY_PROJECT), (True, EMPTY_RANGE)])
    def test_nothing_still_has_a_project_total_saying_why(self, filtered, message):
        book = openpyxl.load_workbook(io.BytesIO(build_workbook(PROJECT, {}, filtered, [])))
        assert book.sheetnames == ["No activity", "Project Total"] and book["Project Total"]["A4"].value == message

    def test_the_tailored_download_warns_too(self):
        rows = summary_sites([self.pending(flag(MAIN, "4in", "25", "25"))], RATES)
        sheet = openpyxl.load_workbook(io.BytesIO(build_tailored_workbook(PROJECT, rows, None, None))).active
        assert sheet["A3"].value == REMEASURE_BANNER and sheet.cell(5, 14).value == "⚠"
