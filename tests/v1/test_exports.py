import io
import re
import textwrap
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, time, timezone
from decimal import Decimal
from pathlib import Path
from typing import Optional
from unittest.mock import patch
from uuid import UUID

import openpyxl
import pytest

from api.queries.projects import get_project_contractor_name
from api.queries.report_attachments import list_uploaded_attachments_for_reports
from api.services import export, export_ac, export_attachments, export_conc_mix, export_swcb
from api.services.export import generate_idr_export
from api.services.export_common import (
    fill_lines, fit_pay_description, paragraphs, pay_item_rows, pay_item_slices, truncate_to_lines,
)
from api.services.export_common import DRAFT_MARKER, REPORT_CONT_TEXT, typed_value
from api.services.export_general import GEN_FRONT_PAY_ITEMS
from api.services.xlsx_template import WorkbookTemplate

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "report_forms.xlsx"
SOURCE_TEMPLATE = ROOT / "templates" / "report_forms_source.xltx"

# Drawing part -> cells under a checkbox rectangle the cleanup makes transparent (Conc Fr, Conc Bk)
CHECKBOX_DRAWINGS = {
    "xl/drawings/drawing10.xml": ["E29", "K29", "S29", "Z29"],
    "xl/drawings/drawing11.xml": ["Z37", "C52"],
}

# AC Fr's drawing, which loses the five collapsed rectangles in its top row, and AC Bk's, whose "Attached Pages" box
# is opened like Conc Fr / Conc Bk's
AC_FR_DRAWING = "xl/drawings/drawing8.xml"
AC_FR_DELETED_SHAPES = ["Rectangle 1", "Rectangle 2", "Rectangle 3", "Rectangle 4", "Rectangle 5"]
AC_BK_DRAWING = "xl/drawings/drawing9.xml"
AC_BK_ATTACHED_PAGES_BOX = "C48"
CLEANED_DRAWINGS = {*CHECKBOX_DRAWINGS, AC_FR_DRAWING, AC_BK_DRAWING}

# Conc Mix drawing -> cells under its checkboxes, which are four-line groups with no fill (the cleanup leaves them)
CONC_MIX_DRAWING = "xl/drawings/drawing5.xml"
CONC_MIX_CHECKBOXES = ["F22", "N22", "X22", "AF22", "O25", "T25"]  # Curb, Sidewalk, Conc Base, Structural, Ready Mix, Other


def checkbox_fill_and_outline(drawing_xml: str, cell: str) -> tuple[str, str]:
    """
    Read a checkbox rectangle's own fill and its outline, from the shape anchored on a cell.
    Takes the drawing XML and the cell.
    Returns (the shape properties before its outline, the outline onwards, without its opening "<a:ln").
    """
    column, row = cell_position(cell)
    shape = re.search(rf"<xdr:twoCellAnchor\b[^>]*><xdr:from><xdr:col>{column}</xdr:col><xdr:colOff>\d+"
                      rf"</xdr:colOff><xdr:row>{row}</xdr:row>.*?</xdr:twoCellAnchor>", drawing_xml, re.DOTALL).group(0)
    fill, outline = re.search(r"<xdr:spPr\b[^>]*>(.*?)</xdr:spPr>", shape, re.DOTALL).group(1).split("<a:ln", 1)
    return fill, outline


def cell_position(coordinate: str) -> tuple[int, int]:
    """
    Convert a cell reference to the zero-based column and row a drawing anchor uses.
    Takes a reference like 'E29'.
    Returns (4, 28).
    """
    letters = re.match(r"[A-Z]+", coordinate).group(0)
    column = 0
    for letter in letters:
        column = column * 26 + ord(letter) - 64
    return column - 1, int(coordinate[len(letters):]) - 1

# The sample project the DDC template shipped with; none of it may survive the cleanup
SAMPLE_VALUES = ["SER200220", "20151410922", "Jewett", "LaPeruta", "Jay Patel", "Staten Island", "STATEN ISLAND"]

IDR_ID = UUID("8b2f887b-1c2d-4e3f-9a0b-1c2d3e4f5a6b")
REPORTER = UUID("7f3c2a9e-1b4d-4c8a-9e2f-3a5b6c7d8e90")

SUBMITTED_IDR = {
    "idr_id": IDR_ID,
    "project_id": "HWS0023",
    "reporter_uuid": REPORTER,
    "report_date": date(2026, 9, 30),  # a Wednesday
    "work_start_time": time(7, 0),
    "work_end_time": time(15, 30),
    "inspector_start_time": time(6, 45),
    "inspector_end_time": None,
    "temp_low": Decimal("45.0"),
    "temp_high": Decimal("62.5"),
    "weather_am": "Cloudy",
    "weather_pm": "Rain",
    "total_pages": 3,
    "has_dismissed_auto_general": False,
    "status": "submitted",
    "submitted_at": datetime(2026, 9, 30, 17, 0),
    "created_at": datetime(2026, 9, 30, 7, 0),
    "updated_at": datetime(2026, 9, 30, 17, 0),
}

DRAFT_IDR = {**SUBMITTED_IDR, "status": "draft", "submitted_at": None, "total_pages": None}

PROJECT = {
    "project_id": "HWS0023",
    "project_name": "S/W Queens 2025",
    "project_description": "Installation of Curb, Sidewalk & Ped-Ramp <Queens>",
    "registration_code": "2024123457",
    "borough": "Queens",
    "status": "Construction; Active",
}

USER = {"user_id": REPORTER, "email": "KhanG@magnoleng.pc", "first_name": "Genghis", "last_name": "Khan"}

GENERAL = {
    "report_type": "GEN",
    "page_number": 1,
    "report_data": {"description": "Poured curb along Main St.\n\nInspected forms before the pour."},
}


@contextmanager
def patched_export(idr=SUBMITTED_IDR, project=PROJECT, contractor="Benny Bowers Contracting Co.", user=USER,
                   general=GENERAL, main_reports=None, reports=None, attachments=None, files=None, signature=None):
    """
    Patch the queries generate_idr_export reads, and Storage, so it runs without a database or a bucket.
    Takes the IDR, project, contractor name, user, General report (or None), non-General main reports, all the
    IDR's reports (where the SWCB report is found), their uploaded attachment rows, {storage path: file bytes}
    (a path that's missing, or maps to an exception, fails its download) and the signature file's bytes (None, or
    an exception, fails its download).
    Yields a dict of the mocks.
    """
    def download(path):
        result = (files or {}).get(path, FileNotFoundError(path))
        if isinstance(result, Exception):
            raise result
        return result

    def download_signature(path, bucket):
        if signature is None or isinstance(signature, Exception):
            raise signature or FileNotFoundError(path)
        return signature

    with (
        patch.object(export, "get_idr_by_id", return_value=idr) as gi,
        patch.object(export, "get_project_by_id", return_value=project) as gp,
        patch.object(export, "get_project_contractor_name", return_value=contractor) as gc,
        patch.object(export, "get_user_by_id", return_value=user) as gu,
        patch.object(export, "get_general_report", return_value=general) as gg,
        patch.object(export, "list_non_general_main_reports", return_value=main_reports or []) as lm,
        patch.object(export, "list_reports_for_idr", return_value=reports or []) as lr,
        patch.object(export, "list_uploaded_attachments_for_reports", return_value=attachments or []) as la,
        patch.object(export_attachments, "download_file", side_effect=download) as df,
        patch.object(export, "download_file", side_effect=download_signature) as ds,
    ):
        yield {"idr": gi, "project": gp, "contractor": gc, "user": gu, "general": gg, "main": lm, "reports": lr,
               "attachments": la, "download": df, "signature": ds}


def exported_workbook(**overrides) -> openpyxl.Workbook:
    """
    Run generate_idr_export on stubbed data and open the result read-only, for checking cell values.
    Takes patched_export's keyword overrides.
    Returns the workbook (read-only mode parses only the sheets a test touches; a full load takes seconds).
    """
    with patched_export(**overrides):
        result = generate_idr_export(IDR_ID)
    return openpyxl.load_workbook(io.BytesIO(result.content), read_only=True)


@pytest.fixture(scope="module")
def full_export() -> tuple[bytes, openpyxl.Workbook]:
    """
    The default export, fully loaded once for the checks read-only mode can't make (styles, print setup, visibility).
    Takes nothing.
    Returns (the export's bytes, the loaded workbook).
    """
    with patched_export():
        content = generate_idr_export(IDR_ID).content
    return content, openpyxl.load_workbook(io.BytesIO(content))


def drawing_parts(package_bytes: bytes) -> list[str]:
    """
    List an .xlsx package's drawing parts.
    Takes the package's bytes.
    Returns the drawing part names, sorted.
    """
    return sorted(n for n in zipfile.ZipFile(io.BytesIO(package_bytes)).namelist() if n.startswith("xl/drawings/drawing"))


# ---------------------------------------------------------------------------
# The cleaned template (templates/report_forms.xlsx)
# ---------------------------------------------------------------------------

class TestTemplate:
    def test_contract_info_values_are_blank(self):
        sheet = openpyxl.load_workbook(TEMPLATE, read_only=True)["Contract Info"]
        assert [sheet[f"C{row}"].value for row in range(2, 8)] == [None] * 6
        assert sheet["B2"].value == "Contract No.:"  # the labels stay

    def test_no_sample_project_data_left_anywhere_in_the_package(self):
        package = zipfile.ZipFile(TEMPLATE)
        xml = "".join(package.read(n).decode("utf-8", "ignore") for n in package.namelist() if n.endswith(".xml"))
        assert [value for value in SAMPLE_VALUES if value in xml] == []

    def test_header_formulas_still_read_contract_info(self):
        sheet = openpyxl.load_workbook(TEMPLATE, read_only=True)["Gen Fr"]
        assert sheet["G8"].value == "='Contract Info'!C2"
        assert sheet["F12"].value == "='Contract Info'!C5"

    def test_cleanup_kept_every_logo_and_shape(self):
        assert drawing_parts(TEMPLATE.read_bytes()) == drawing_parts(SOURCE_TEMPLATE.read_bytes())
        source, cleaned = zipfile.ZipFile(SOURCE_TEMPLATE), zipfile.ZipFile(TEMPLATE)
        for name in source.namelist():
            if name.startswith("xl/media/") or (name.startswith("xl/drawings/") and name not in CLEANED_DRAWINGS):
                assert cleaned.read(name) == source.read(name), name
        # The cleaned drawings only lose their checkboxes' fill, and AC Fr its collapsed rectangles: every other
        # shape is still there
        names = lambda package, part: re.findall(r'<xdr:cNvPr id="\d+" name="([^"]+)"', package.read(part).decode())
        for part in CLEANED_DRAWINGS:
            kept = [n for n in names(source, part) if part != AC_FR_DRAWING or n not in AC_FR_DELETED_SHAPES]
            assert names(cleaned, part) == kept, part

    def test_conc_checkboxes_are_transparent_with_their_outlines(self):
        cleaned = zipfile.ZipFile(TEMPLATE)
        for drawing, cells in CHECKBOX_DRAWINGS.items():
            xml = cleaned.read(drawing).decode()
            for cell in cells:
                fill, outline = checkbox_fill_and_outline(xml, cell)
                assert "<a:noFill/>" in fill and "<a:solidFill>" not in fill, cell
                assert '<a:srgbClr val="000000"/>' in outline, cell

    def test_ac_bk_attached_pages_box_is_transparent_with_its_outline(self):
        xml = zipfile.ZipFile(TEMPLATE).read(AC_BK_DRAWING).decode()
        fill, outline = checkbox_fill_and_outline(xml, AC_BK_ATTACHED_PAGES_BOX)
        assert "<a:noFill/>" in fill and "<a:solidFill>" not in fill
        assert '<a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/>' in "<a:ln" + outline  # thin black, as before
        source = zipfile.ZipFile(SOURCE_TEMPLATE).read(AC_BK_DRAWING).decode()
        assert '<a:srgbClr val="FFFFFF"/>' in checkbox_fill_and_outline(source, AC_BK_ATTACHED_PAGES_BOX)[0]

    def test_ac_fr_has_no_shapes_in_its_top_row(self):
        # Row 1 is the strip "DRAFT - Not for Submission" goes across; the five rectangles there were collapsed
        anchored_rows = lambda xml: [int(r) for r in re.findall(r"<xdr:from><xdr:col>\d+</xdr:col><xdr:colOff>\d+"
                                                                r"</xdr:colOff><xdr:row>(\d+)</xdr:row>", xml)]
        source = zipfile.ZipFile(SOURCE_TEMPLATE).read(AC_FR_DRAWING).decode()
        cleaned = zipfile.ZipFile(TEMPLATE).read(AC_FR_DRAWING).decode()
        assert anchored_rows(source).count(0) == 5
        assert 0 not in anchored_rows(cleaned)
        assert cleaned.count("<xdr:pic>") == source.count("<xdr:pic>") == 2  # the banner and the logo stay

    def test_conc_mix_checkboxes_have_no_fill_to_hide_a_stamped_x(self):
        xml = zipfile.ZipFile(TEMPLATE).read(CONC_MIX_DRAWING).decode()
        for cell in CONC_MIX_CHECKBOXES:
            column, row = cell_position(cell)
            anchor = re.search(rf"<xdr:twoCellAnchor\b[^>]*><xdr:from><xdr:col>{column}</xdr:col><xdr:colOff>\d+"
                               rf"</xdr:colOff><xdr:row>{row}</xdr:row>.*?</xdr:twoCellAnchor>", xml, re.DOTALL).group(0)
            properties = re.findall(r"<xdr:(?:grpSpPr|spPr)\b[^>]*>(.*?)</xdr:(?:grpSpPr|spPr)>", anchor, re.DOTALL)
            assert properties, cell
            # Each shape's own fill comes before its outline; the group and its lines must carry none
            assert all("<a:solidFill>" not in p.split("<a:ln", 1)[0] for p in properties), cell


# ---------------------------------------------------------------------------
# generate_idr_export: Gen Fr header
# ---------------------------------------------------------------------------

def general_with(**report_data) -> dict:
    """
    Build the General report row the export reads, with the given report_data.
    Takes report_data fields as keyword arguments.
    Returns the row (page 1).
    """
    return {"report_type": "GEN", "page_number": 1, "report_data": report_data}


def visible_sheets(content: bytes) -> list[str]:
    """
    List an export's visible sheets in tab order, read straight from workbook.xml (fast, no full load).
    Takes the .xlsx bytes.
    Returns the sheet names without state="hidden".
    """
    workbook = zipfile.ZipFile(io.BytesIO(content)).read("xl/workbook.xml").decode()
    return [re.search(r'name="([^"]+)"', tag).group(1) for tag in re.findall(r"<sheet [^>]*/>", workbook)
            if 'state="hidden"' not in tag]


def export_bytes(**overrides) -> bytes:
    """
    Run generate_idr_export on stubbed data.
    Takes patched_export's keyword overrides.
    Returns the .xlsx bytes.
    """
    with patched_export(**overrides):
        return generate_idr_export(IDR_ID).content


# Long enough to fill Gen Fr, Gen Bk's comment lines and Report Cont, and still be cut
LONG_DESCRIPTION = " ".join(f"word{n}" for n in range(400))
CASCADE = general_with(description=LONG_DESCRIPTION, comments="Visitor from DEP at 10am.")


@pytest.fixture(scope="module")
def full_cascade() -> tuple[bytes, openpyxl.Workbook]:
    """
    An export whose description runs onto Gen Bk and Report Cont, fully loaded once for style and order checks.
    Takes nothing.
    Returns (the export's bytes, the loaded workbook).
    """
    content = export_bytes(general=CASCADE)
    return content, openpyxl.load_workbook(io.BytesIO(content))


class TestGeneralFrontHeader:
    def test_returns_a_valid_workbook_named_for_the_idr(self):
        with patched_export():
            result = generate_idr_export(IDR_ID)
        assert result.filename == f"IDR_{IDR_ID}_2026-09-30.xlsx"
        assert openpyxl.load_workbook(io.BytesIO(result.content), read_only=True).sheetnames[0] == "Contract Info"

    def test_stamps_contract_info(self):
        sheet = exported_workbook()["Contract Info"]
        assert [sheet[f"C{row}"].value for row in range(2, 8)] == [
            "HWS0023", "2024123457", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens",
            "Benny Bowers Contracting Co.", None,  # no Resident Engineer in the data model yet
        ]

    def test_writes_project_values_over_the_contract_info_formulas(self):
        sheet = exported_workbook()["Gen Fr"]
        # A formula would read back as "='Contract Info'!C2"; the export leaves plain values for previewers
        assert [sheet[c].value for c in ("G8", "P8", "I10", "F12", "F14")] == [
            "HWS0023", "2024123457", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens",
            "Benny Bowers Contracting Co.",
        ]

    def test_stamps_the_header_fields(self):
        sheet = exported_workbook()["Gen Fr"]
        assert sheet["AI4"].value == datetime(2026, 9, 30)
        assert sheet["AH8"].value == 1 and sheet["AM8"].value == 3
        assert sheet["AG10"].value == "( Start 07:00 End 15:30 )"
        assert sheet["AG12"].value == "( Start 06:45 End ________ )"  # one of two times known keeps the other blank
        assert sheet["AD13"].value == "Low  45" and sheet["AK13"].value == "High  62.5"
        assert sheet["AD17"].value == "Cloudy" and sheet["AK17"].value == "Rain"
        assert sheet["H17"].value == "Genghis Khan"
        assert sheet["AH6"].value is None  # no I.R. No. in the data model yet

    def test_empty_fields_clear_the_template_placeholders(self):
        idr = {**SUBMITTED_IDR, "work_start_time": None, "work_end_time": None, "inspector_start_time": None,
               "temp_low": None, "temp_high": None, "weather_am": None, "weather_pm": "  "}
        sheet = exported_workbook(idr=idr, contractor=None, user=None)["Gen Fr"]
        assert sheet["AG10"].value is None and sheet["AG12"].value is None
        assert sheet["AD17"].value is None and sheet["AK17"].value is None
        assert sheet["F14"].value is None and sheet["H17"].value is None
        # Low / High are the boxes' labels, not placeholders, so they stay
        assert sheet["AD13"].value == "Low" and sheet["AK13"].value == "High"

    def test_highlights_only_the_day_of_week(self, full_export):
        sheet = full_export[1]["Gen Fr"]
        highlighted = [c for c in ("AI5", "AJ5", "AK5", "AL5", "AM5", "AN5", "AO5") if sheet[c].fill.fill_type == "solid"]
        assert highlighted == ["AL5"]
        assert sheet["AL5"].value == "W" and sheet["AL5"].border.left.style == "medium"

    def test_sets_letter_portrait_one_page_print_setup(self, full_export):
        sheet = full_export[1]["Gen Fr"]
        assert sheet.page_setup.paperSize == 1  # US Letter
        assert sheet.page_setup.orientation == "portrait"
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
        assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)

    def test_composes_a_general_when_the_idr_has_none(self):
        children = [
            {"report_type": "SWCB", "report_data": {"description": "Formed sidewalk."}},
            {"report_type": "CONC_MIX", "report_data": {"description": "Addendum-type, left out."}},
        ]
        sheet = exported_workbook(general=None, main_reports=children)["Gen Fr"]
        assert sheet["B22"].value == "Sidewalk, Curb, Concrete Base: Formed sidewalk."
        assert sheet["B23"].value.startswith("See the individual reports")
        assert sheet["AH8"].value is None and sheet["AM8"].value is None  # not one of the IDR's numbered pages

    def test_unknown_idr_raises_not_found(self):
        with patched_export(idr=None), pytest.raises(export.IdrNotFoundError):
            generate_idr_export(IDR_ID)

    def test_draft_idr_exports(self):
        with patched_export(idr={**DRAFT_IDR}):
            result = generate_idr_export(IDR_ID)
        assert openpyxl.load_workbook(io.BytesIO(result.content), read_only=True)["Gen Fr"]["H17"].value == "Genghis Khan"


# ---------------------------------------------------------------------------
# Pay items (Gen Fr)
# ---------------------------------------------------------------------------

def pay_item(n: int, **overrides) -> dict:
    """
    Build a saved pay item.
    Takes a number to make it distinct and field overrides.
    Returns the pay item dict.
    """
    return {"itemNo": f"4.{n:02d} AAS", "budgetCode": "12345", "payQuantity": "312.50", "unit": "S.F.",
            "description": f"Sidewalk {n}", **overrides}


class TestPayItems:
    def test_stamps_each_column_with_the_unit_in_pay_quantity_and_chk_blank(self):
        general = general_with(payItems=[pay_item(1, quantityChk="RM")])
        sheet = exported_workbook(general=general)["Gen Fr"]
        assert [sheet[f"{c}39"].value for c in "BGNSX"] == ["4.01 AAS", "12345", "312.50 S.F.", None, "Sidewalk 1"]
        assert sheet["B40"].value is None

    def test_pay_quantity_without_a_unit_is_just_the_number(self):
        general = general_with(payItems=[pay_item(1, unit=""), pay_item(2, payQuantity="", unit="S.F.")])
        sheet = exported_workbook(general=general)["Gen Fr"]
        assert sheet["N39"].value == "312.50"
        assert sheet["N40"].value == "S.F."

    def test_more_items_than_rows_continue_on_the_next_page(self):
        general = general_with(payItems=[pay_item(n) for n in range(14)])
        book = exported_workbook(general=general)
        sheet = book["Gen Fr"]
        assert sheet["B49"].value == "4.10 AAS"  # 11 items fit, then the note on the 12th row
        assert [sheet[f"{c}50"].value for c in "BGNS"] == [None] * 4
        assert sheet["X50"].value == "Pay items continued on next page"
        assert [book["Gen Fr 2"][f"B{row}"].value for row in (39, 40, 41, 42)] == [
            "4.11 AAS", "4.12 AAS", "4.13 AAS", None]

    def test_non_list_pay_items_are_none(self):
        assert pay_item_rows(None, 12) == [] and pay_item_rows({"a": 1}, 12) == [] and pay_item_rows(["x"], 12) == []

    def test_a_long_description_wraps_on_a_taller_row(self):
        long = 'Thermoplastic Reflectorized Pavement Markings (4" Wide)'  # 55 characters: two lines
        content = export_bytes(general=general_with(payItems=[pay_item(1, description=long)]))
        cell = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]["X39"]
        assert cell.value == long
        assert cell.alignment.wrap_text is True
        assert cell.alignment.horizontal == "center"  # the template's centring is kept
        assert gen_front_row_height(content, 39) == 2 * 12.75

    def test_a_short_description_wraps_but_keeps_the_row_height(self):
        content = export_bytes(general=general_with(payItems=[pay_item(1, description="Plastic Barrels")]))
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]["X39"].alignment.wrap_text is True
        assert gen_front_row_height(content, 39) == 18.0

    def test_descriptions_shrink_to_8pt_before_being_cut(self):
        level_1 = "Plastic Barrels"
        level_2 = 'Thermoplastic Reflectorized Pavement Markings (4" Wide)'
        level_3 = "Reinforced Concrete Pavement (Full Width Pavement) with dowels and tie bars at all joints"
        level_4 = ('Thermoplastic Reflectorized Pavement Markings (4" Wide) including surface preparation, primer, '
                   "layout and removal of existing markings where shown")
        items = [pay_item(n, description=d) for n, d in enumerate((level_1, level_2, level_3, level_4))]
        content = export_bytes(general=general_with(payItems=items))
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]
        cells = [sheet[f"X{row}"] for row in (39, 40, 41, 42)]
        assert [c.font.sz for c in cells] == [10, 10, 8, 8]
        assert [gen_front_row_height(content, row) for row in (39, 40, 41, 42)] == [18.0, 25.5, 22.5, 22.5]
        assert [c.value for c in cells[:3]] == [level_1, level_2, level_3]  # not cut
        assert cells[3].value.endswith("preparation...") and len(textwrap.wrap(cells[3].value, 48)) == 2
        assert all(c.alignment.wrap_text for c in cells)

    def test_fit_pay_description_levels(self):
        assert fit_pay_description("Plastic Barrels", GEN_FRONT_PAY_ITEMS) == ("Plastic Barrels", 1, 10)
        assert fit_pay_description("word " * 15, GEN_FRONT_PAY_ITEMS)[1:] == (2, 10)
        assert fit_pay_description("word " * 18, GEN_FRONT_PAY_ITEMS)[1:] == (2, 8)
        text, lines, size = fit_pay_description("word " * 40, GEN_FRONT_PAY_ITEMS)
        assert (lines, size) == (2, 8) and text.endswith("word...")

    def test_truncation_ends_on_a_word_and_handles_one_long_word(self):
        cut = truncate_to_lines("alpha beta gamma delta epsilon zeta", width=12, max_lines=2, suffix="...")
        assert cut == "alpha beta gamma..."
        long_word = truncate_to_lines("X" * 130, width=48, max_lines=2, suffix="...")
        assert long_word == "X" * 93 + "..." and len(textwrap.wrap(long_word, 48)) == 2

    def test_empty_pay_item_rows_keep_the_template_style(self):
        content = export_bytes(general=general_with(payItems=[pay_item(1)]))
        template_style = WorkbookTemplate(TEMPLATE).cell_style("Gen Fr", "X40")
        assert gen_front_cell_style(content, "X40") == template_style
        assert gen_front_row_height(content, 40) == 18.0


def gen_front_xml(content: bytes) -> str:
    """
    Read Gen Fr's worksheet XML from an export.
    Takes the .xlsx bytes.
    Returns the XML (Gen Fr is sheet4.xml in the template package).
    """
    return zipfile.ZipFile(io.BytesIO(content)).read("xl/worksheets/sheet4.xml").decode()


def gen_front_row_height(content: bytes, row: int) -> float:
    """
    Read one Gen Fr row's height from an export (read-only openpyxl doesn't expose row heights).
    Takes the .xlsx bytes and the row number.
    Returns the height in points.
    """
    return float(re.search(rf'<row r="{row}"[^>]*?\sht="([\d.]+)"', gen_front_xml(content)).group(1))


def gen_front_cell_style(content: bytes, coordinate: str) -> str:
    """
    Read one Gen Fr cell's style index from an export.
    Takes the .xlsx bytes and the cell reference.
    Returns the style index.
    """
    return re.search(rf'<c r="{coordinate}"[^>]*?\ss="(\d+)"', gen_front_xml(content)).group(1)


# ---------------------------------------------------------------------------
# Gen Bk: work force, equipment, safety, comments
# ---------------------------------------------------------------------------

class TestGenBack:
    def test_gen_bk_is_shown_even_with_nothing_on_it(self):
        # Gen Bk carries the certification and signature lines, so every General prints it
        content = export_bytes(general=general_with(description="Poured curb."))
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]

    def test_gen_bk_is_shown_with_workforce_only(self):
        content = export_bytes(general=general_with(description="Poured curb.", workforce={"laborers": "4"}))
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Bk"]
        assert sheet["G46"].value == 4
        assert [sheet[f"B{row}"].value for row in range(3, 7)] == [None] * 4  # no text cascaded onto it

    def test_workforce_counts_including_legacy_keys(self):
        general = general_with(workforce={"superintendent": "1", "foreman": "2", "laborers": "6", "flaggers": "x2"})
        content = export_bytes(general=general)
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Bk"]
        assert [sheet[f"G{row}"].value for row in range(43, 48)] == [1, 2, None, 6, "x2"]
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]

    def test_added_trades_use_printed_rows_then_blank_rows_then_a_count(self):
        trades = [{"label": "Masons", "count": "3"}, {"label": "Chauffeurs", "count": "1"},
                  {"label": "Surveyors", "count": "2"}, {"label": "Welders", "count": "1"}]
        sheet = exported_workbook(general=general_with(additionalWorkforce=trades))["Gen Bk"]
        assert sheet["G50"].value == 3  # Masons' own row
        assert (sheet["B52"].value, sheet["G52"].value) == ("Chauffeurs", 1)
        assert (sheet["B53"].value, sheet["G53"].value) == ("+2 more (see ICID)", None)

    def test_equipment_model_and_number(self):
        general = general_with(equipment={"backhoe": {"model": "CAT 420", "number": "1"}, "excavator": {"model": "", "number": ""}})
        sheet = exported_workbook(general=general)["Gen Bk"]
        assert (sheet["N44"].value, sheet["V44"].value) == ("CAT 420", 1)
        assert (sheet["N49"].value, sheet["V49"].value) == (None, None)

    def test_added_equipment_shares_printed_rows_and_uses_the_blank_row(self):
        extras = [{"label": "Crane", "model": "Grove", "number": "1"},
                  {"label": "Roller – Static", "model": "BW120", "number": "1"},
                  {"label": "Roller – Dynamic", "model": "CB24", "number": "2"},
                  {"label": "Sweepers", "model": "Elgin", "number": "1"}]
        sheet = exported_workbook(general=general_with(additionalEquipment=extras))["Gen Bk"]
        assert (sheet["N45"].value, sheet["V45"].value) == ("Grove", 1)
        # Both rollers on the Roller row, in its two Model / No. pairs, each keeping its variant name
        assert (sheet["N47"].value, sheet["V47"].value) == ("Roller – Static BW120", 1)
        assert (sheet["Y47"].value, sheet["AG47"].value) == ("Roller – Dynamic CB24", 2)
        assert (sheet["I53"].value, sheet["N53"].value, sheet["V53"].value) == ("Sweepers", "Elgin", 1)

    def test_safety_checklist_marks_y_or_n_and_writes_remarks(self):
        general = general_with(
            safetyChecks={"plasticBarrels": "Y", "fencing": "N", "plates": "NA", "arrowBoard": True, "timberCurbs": None},
            safetyRemarks={"fencing": "Gap at gate", "plates": "None on site", "generalSafety": "Good"},
        )
        sheet = exported_workbook(general=general)["Gen Bk"]
        row = lambda r: (sheet[f"N{r}"].value, sheet[f"P{r}"].value, sheet[f"R{r}"].value)
        assert row(30) == ("X", None, None)                      # plasticBarrels: Y
        assert row(36) == (None, "X", "Gap at gate")             # fencing: N
        assert row(37) == (None, None, "N/A — None on site")      # plates: N/A has no column on the form
        assert row(38) == ("X", None, None)                      # arrowBoard: older reports' True
        assert row(32) == (None, None, None)                     # timberCurbs: unanswered
        assert row(34) == (None, None, "Good")                   # a remark without an answer

    def test_comments_go_on_the_back_without_continuing(self):
        content = export_bytes(general=general_with(description="Short.", comments="Visitor from DEP at 10am."))
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert workbook["Gen Bk"]["B3"].value == "Visitor from DEP at 10am."
        assert workbook["Gen Bk"]["B4"].value is None
        assert workbook["Gen Bk"]["Z27"].value is None
        assert workbook["Gen Fr"]["AC36"].value is None  # the description fit on the front
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]


# ---------------------------------------------------------------------------
# Description of Work overflow: Gen Fr → Gen Bk → Report Cont
# ---------------------------------------------------------------------------

def words(count: int, start: int = 0) -> str:
    """
    Make filler text of distinct words.
    Takes how many words and the first word's number.
    Returns the text.
    """
    return " ".join(f"word{n}" for n in range(start, start + count))


class TestDescriptionCascade:
    def test_level_1_fits_on_the_front(self, full_export):
        content, workbook = full_export
        sheet = workbook["Gen Fr"]
        assert sheet["B22"].value == "Poured curb along Main St."
        assert sheet["AC36"].value is None
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]  # Gen Bk always prints; Report Cont doesn't
        # The template centres the first line; every description line is left-aligned
        assert [sheet[f"B{row}"].alignment.horizontal for row in (22, 23, 34)] == ["left"] * 3

    def test_level_2_continues_on_the_back_and_ticks_reverse_page_used(self):
        content = export_bytes(general=general_with(description=words(110), comments="Visitor."))
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert workbook["Gen Fr"]["AC36"].value == "X"
        back = [workbook["Gen Bk"][f"B{row}"].value for row in range(3, 7)]
        assert back[0] == "Description of work (continued):"
        assert back[-2:] == ["Comments:", "Visitor."]
        assert workbook["Gen Bk"]["Z27"].value is None
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk"]

    def test_level_3_continues_on_report_cont_and_ticks_continued(self, full_cascade):
        content, workbook = full_cascade
        assert workbook["Gen Fr"]["AC36"].value == "X"
        assert workbook["Gen Bk"]["Z27"].value == "X"
        assert workbook["Report Cont"]["B21"].value.startswith("word")
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Report Cont"]

    def test_text_past_report_cont_is_cut_with_a_note(self, full_cascade):
        sheet = full_cascade[1]["Report Cont"]
        assert sheet["B45"].value.endswith("… (continued in ICID)")
        assert len(sheet["B45"].value) <= REPORT_CONT_TEXT.line_chars

    def test_report_cont_header_is_stamped_and_placeholders_cleared(self, full_cascade):
        sheet = full_cascade[1]["Report Cont"]
        assert sheet["I10"].value == "9/30/26"
        assert sheet["L11"].fill.fill_type == "solid"  # Wednesday
        assert sheet["U10"].value is None and sheet["AA10"].value == "Sheet No.:"
        assert [sheet[c].value for c in ("G14", "P14", "F17", "H19")] == ["HWS0023", "2024123457", "Queens", "Genghis Khan"]

    def test_continuation_lines_are_left_aligned(self, full_cascade):
        workbook = full_cascade[1]
        assert workbook["Gen Bk"]["B3"].alignment.horizontal == "left"  # the template centres these
        assert workbook["Report Cont"]["B21"].alignment.horizontal == "left"

    def test_report_cont_prints_after_the_general_pages(self, full_cascade):
        workbook = full_cascade[1]
        names = workbook.sheetnames
        assert names.index("Gen Fr") < names.index("Gen Bk") < names.index("Report Cont")
        assert workbook.active.title == "Gen Fr"
        # The one sheet-scoped name (AC Fr's print area) still points at AC Fr after the move
        xml = zipfile.ZipFile(io.BytesIO(full_cascade[0])).read("xl/workbook.xml").decode()
        assert re.search(r'localSheetId="(\d+)"', xml).group(1) == str(names.index("AC Fr"))

    def test_control_characters_do_not_break_the_file(self):
        general = general_with(description="Bad \x01char & <tag>")
        assert exported_workbook(general=general)["Gen Fr"]["B22"].value == "Bad char & <tag>"


class TestTextLayout:
    def test_fill_lines_wraps_and_leaves_the_rest_queued(self):
        queue = [words(30), "Second paragraph."]
        lines = fill_lines(queue, capacity=2, width=40)
        assert len(lines) == 2 and all(len(line) <= 40 for line in lines)
        assert queue[0].startswith("word") and queue[1] == "Second paragraph."

    def test_each_paragraph_starts_a_new_line(self):
        assert fill_lines(["One.", "Two."], capacity=5, width=40) == ["One.", "Two."]

    def test_paragraphs_ignores_blank_lines_and_non_text(self):
        assert paragraphs("A\n\n  B  \n") == ["A", "B"]
        assert paragraphs(None) == [] and paragraphs({"a": 1}) == []


# ---------------------------------------------------------------------------
# The XML writer
# ---------------------------------------------------------------------------

def written(workbook: WorkbookTemplate) -> openpyxl.Workbook:
    """
    Serialize a WorkbookTemplate and open it read-only.
    Takes the workbook.
    Returns the loaded workbook.
    """
    return openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)


class TestWorkbookTemplate:
    def test_a_value_replaces_a_formula(self):
        workbook = WorkbookTemplate(TEMPLATE)
        assert workbook.cell_style("Gen Fr", "G8") is not None
        workbook.set_cell("Gen Fr", "G8", "HWS0023")
        sheet = written(workbook)["Gen Fr"]
        assert sheet["G8"].value == "HWS0023"
        assert sheet["P8"].value == "='Contract Info'!C3"  # untouched neighbours keep their formulas

    def test_left_alignment_keeps_the_other_alignment_settings(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.align_left("Gen Fr", "AD17")  # a centred, wrapped cell
        cell = written(workbook)["Gen Fr"]["AD17"]
        assert cell.alignment.horizontal == "left"
        assert cell.alignment.wrap_text is True

    def test_derived_styles_are_shared(self):
        workbook = WorkbookTemplate(TEMPLATE)
        style = workbook.cell_style("Gen Fr", "B23")
        assert workbook.left_aligned_style(style) == workbook.left_aligned_style(style)

    def test_wrap_text_style_on_the_default_style(self):
        workbook = WorkbookTemplate(TEMPLATE)
        style = workbook.wrap_text_style(None)
        styles = zipfile.ZipFile(io.BytesIO(workbook.to_bytes())).read("xl/styles.xml").decode()
        cell_xfs = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.DOTALL)
        assert re.search(r'count="(\d+)"', cell_xfs.group(0)).group(1) == str(int(style) + 1)
        new_xf = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", cell_xfs.group(1), re.DOTALL)[int(style)]
        assert 'applyAlignment="1"' in new_xf and '<alignment wrapText="1"/>' in new_xf

    def test_wrap_style_is_shared_and_keeps_left_alignment(self):
        workbook = WorkbookTemplate(TEMPLATE)
        left = workbook.left_aligned_style(workbook.cell_style("Gen Fr", "B23"))
        assert workbook.wrap_text_style(left) == workbook.wrap_text_style(left)
        workbook.set_style("Gen Fr", "B23", workbook.wrap_text_style(left))
        cell = written(workbook)["Gen Fr"]["B23"]
        assert (cell.alignment.horizontal, cell.alignment.wrap_text) == ("left", True)

    def test_font_size_style_points_at_a_new_8pt_font(self):
        workbook = WorkbookTemplate(TEMPLATE)
        style = workbook.font_size_style(None, 8)
        assert workbook.font_size_style(None, 8) == style  # shared, not appended twice
        styles = zipfile.ZipFile(io.BytesIO(workbook.to_bytes())).read("xl/styles.xml").decode()
        fonts_block = re.search(r"<fonts\b[^>]*>(.*?)</fonts>", styles, re.DOTALL)
        fonts = re.findall(r"<font\b[^>]*?(?:/>|>.*?</font>)", fonts_block.group(1), re.DOTALL)
        assert re.search(r'count="(\d+)"', fonts_block.group(0)).group(1) == str(len(fonts))
        xf = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.DOTALL).group(1), re.DOTALL)[int(style)]
        new_font = fonts[int(re.search(r'fontId="(\d+)"', xf).group(1))]
        assert '<sz val="8"/>' in new_font and '<name val="Arial"/>' in new_font
        assert 'applyFont="1"' in xf
        assert '<sz val="10"/>' in fonts[0]  # the shared default font is left alone

    def test_set_font_size_keeps_the_cell_wrapping(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.wrap_cell("Gen Fr", "X39")
        workbook.set_font_size("Gen Fr", "X39", 8)
        cell = written(workbook)["Gen Fr"]["X39"]
        assert (cell.font.sz, cell.alignment.wrap_text, cell.alignment.horizontal) == (8, True, "center")

    def test_set_row_height_fixes_the_height(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.set_row_height("Gen Fr", 39, 25.5)
        assert gen_front_row_height(workbook.to_bytes(), 39) == 25.5

    def test_the_calculation_chain_is_dropped(self):
        package = zipfile.ZipFile(io.BytesIO(WorkbookTemplate(TEMPLATE).to_bytes()))
        assert "xl/calcChain.xml" not in package.namelist()
        assert "calcChain" not in package.read("xl/_rels/workbook.xml.rels").decode()
        assert "calcChain" not in package.read("[Content_Types].xml").decode()


# ---------------------------------------------------------------------------
# WorkbookTemplate.clone_sheet
# ---------------------------------------------------------------------------

def cloned(source: str = "Conc Mix", new_name: str = "Conc Mix 2") -> WorkbookTemplate:
    """
    Load the template and clone one of its sheets.
    Takes the source sheet name and the clone's name.
    Returns the workbook holding the clone.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    workbook.clone_sheet(source, new_name)
    return workbook


def clone_parts(content: bytes, name: str) -> tuple[str, str]:
    """
    Find a sheet's worksheet part and the drawing part it shows, through the package's relationships.
    Takes the .xlsx bytes and the sheet name.
    Returns (the worksheet part name, the drawing part name).
    """
    package = zipfile.ZipFile(io.BytesIO(content))
    rel_id = re.search(rf'<sheet name="{name}"[^>]*r:id="(rId\d+)"', package.read("xl/workbook.xml").decode()).group(1)
    rels = package.read("xl/_rels/workbook.xml.rels").decode()
    part = "xl/" + re.search(rf'Id="{rel_id}"[^>]*Target="([^"]+)"', rels).group(1)
    sheet_rels = package.read(part.replace("worksheets/", "worksheets/_rels/") + ".rels").decode()
    return part, "xl/drawings/" + re.search(r'Target="\.\./drawings/([^"]+)"', sheet_rels).group(1)


class TestCloneSheet:
    def test_the_clone_is_added_as_the_last_tab(self):
        names = written(cloned()).sheetnames
        assert names[-1] == "Conc Mix 2" and names.index("Conc Mix") == 7 and len(names) == 39

    def test_the_clone_starts_byte_identical_to_its_source(self):
        content = cloned().to_bytes()
        package = zipfile.ZipFile(io.BytesIO(content))
        (source, _), (clone, _) = clone_parts(content, "Conc Mix"), clone_parts(content, "Conc Mix 2")
        assert (source, clone) == ("xl/worksheets/sheet8.xml", "xl/worksheets/sheet39.xml")
        assert package.read(clone) == package.read(source)

    def test_the_clone_has_its_own_drawing_with_the_same_shapes(self):
        content = cloned().to_bytes()
        package = zipfile.ZipFile(io.BytesIO(content))
        (_, source), (_, clone) = clone_parts(content, "Conc Mix"), clone_parts(content, "Conc Mix 2")
        assert (source, clone) == ("xl/drawings/drawing5.xml", "xl/drawings/drawing23.xml")
        assert package.read(clone) == package.read(source)
        # Its images are the source's (shared media), through its own relationships part
        rels = "xl/drawings/_rels/drawing{}.xml.rels"
        assert package.read(rels.format(23)) == package.read(rels.format(5))
        # and its own printer settings
        sheet_rels = package.read("xl/worksheets/_rels/sheet39.xml.rels").decode()
        assert "../printerSettings/printerSettings37.bin" in sheet_rels

    def test_cells_stamped_on_one_sheet_stay_off_the_other(self):
        workbook = cloned()
        workbook.set_cell("Conc Mix", "B28", "T-1")
        workbook.set_cell("Conc Mix 2", "B29", "T-12")
        book = written(workbook)
        assert (book["Conc Mix"]["B28"].value, book["Conc Mix"]["B29"].value) == ("T-1", None)
        assert (book["Conc Mix 2"]["B28"].value, book["Conc Mix 2"]["B29"].value) == (None, "T-12")

    def test_a_box_ticked_on_the_clone_leaves_the_sources_box_alone(self):
        workbook = cloned()
        workbook.set_cell("Conc Mix 2", "F22", "X")  # Curb's box, ticked the way export_conc_mix ticks it
        workbook.center_across("Conc Mix 2", ["F22", "G22"])
        content = workbook.to_bytes()
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Mix 2"]["F22"].value, book["Conc Mix"]["F22"].value) == ("X", None)
        assert book["Conc Mix"]["F22"].alignment.horizontal is None
        package, template = zipfile.ZipFile(io.BytesIO(content)), zipfile.ZipFile(TEMPLATE)
        assert package.read("xl/worksheets/sheet8.xml") == template.read("xl/worksheets/sheet8.xml")

    def test_the_clone_is_registered_in_the_content_types(self):
        types = zipfile.ZipFile(io.BytesIO(cloned().to_bytes())).read("[Content_Types].xml").decode()
        worksheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
        drawing = "application/vnd.openxmlformats-officedocument.drawing+xml"
        assert f'<Override PartName="/xl/worksheets/sheet39.xml" ContentType="{worksheet}"/>' in types
        assert f'<Override PartName="/xl/drawings/drawing23.xml" ContentType="{drawing}"/>' in types

    def test_the_clone_gets_a_new_sheet_id_and_relationship(self):
        package = zipfile.ZipFile(io.BytesIO(cloned().to_bytes()))
        workbook = package.read("xl/workbook.xml").decode()
        assert '<sheet name="Conc Mix 2" sheetId="57" r:id="rId43"/>' in workbook  # template tops out at 56 / rId42
        rels = package.read("xl/_rels/workbook.xml.rels").decode()
        assert 'Id="rId43" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" ' \
               'Target="worksheets/sheet39.xml"' in rels

    def test_the_sources_print_area_is_copied_onto_the_clone(self):
        book = openpyxl.load_workbook(io.BytesIO(cloned("AC Fr", "AC Fr 2").to_bytes()))
        assert book["AC Fr"].print_area == "'AC Fr'!$A$1:$AQ$63"
        assert book["AC Fr 2"].print_area == "'AC Fr 2'!$A$1:$AQ$63"
        # A sheet with no scoped names (Conc Mix) clones without any
        xml = zipfile.ZipFile(io.BytesIO(cloned().to_bytes())).read("xl/workbook.xml").decode()
        assert xml.count("<definedName ") == 1

    def test_fit_to_letter_page_sets_up_the_clone_alone(self):
        workbook = cloned()
        workbook.fit_to_letter_page("Conc Mix 2")
        book = openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))
        assert (book["Conc Mix 2"].page_setup.fitToWidth, book["Conc Mix 2"].page_setup.fitToHeight) == (1, 1)
        assert book["Conc Mix"].sheet_properties.pageSetUpPr is None or \
            not book["Conc Mix"].sheet_properties.pageSetUpPr.fitToPage

    def test_show_only_shows_the_source_and_its_clone(self):
        workbook = cloned()
        workbook.show_only(["Conc Mix", "Conc Mix 2"])
        assert visible_sheets(workbook.to_bytes()) == ["Conc Mix", "Conc Mix 2"]
        assert written(workbook).active.title == "Conc Mix"

    def test_move_sheet_puts_the_clone_after_its_source(self):
        workbook = cloned("AC Fr", "AC Fr 2")
        workbook.move_sheet("AC Fr 2", after="AC Fr")
        content = workbook.to_bytes()
        names = openpyxl.load_workbook(io.BytesIO(content), read_only=True).sheetnames
        assert names.index("AC Fr 2") == names.index("AC Fr") + 1 and len(names) == 39
        book = openpyxl.load_workbook(io.BytesIO(content))  # each print area still on its own sheet
        assert (book["AC Fr"].print_area, book["AC Fr 2"].print_area) == (
            "'AC Fr'!$A$1:$AQ$63", "'AC Fr 2'!$A$1:$AQ$63")

    def test_cloning_needs_an_existing_source_and_a_free_name(self):
        workbook = WorkbookTemplate(TEMPLATE)
        with pytest.raises(ValueError, match="no sheet named"):
            workbook.clone_sheet("Conc Mixx", "Conc Mix 2")
        with pytest.raises(ValueError, match="already exists"):
            workbook.clone_sheet("Conc Mix", "Conc Fr")
        assert workbook.to_bytes() == WorkbookTemplate(TEMPLATE).to_bytes()  # a refused clone changes nothing

    def test_cloning_leaves_every_other_part_as_it_was(self):
        template = zipfile.ZipFile(TEMPLATE)
        package = zipfile.ZipFile(io.BytesIO(cloned().to_bytes()))
        changed = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/calcChain.xml"}
        assert [n for n in template.namelist() if n not in changed and package.read(n) != template.read(n)] == []
        added = sorted(set(package.namelist()) - set(template.namelist()))
        assert added == ["xl/drawings/_rels/drawing23.xml.rels", "xl/drawings/drawing23.xml",
                         "xl/printerSettings/printerSettings37.bin", "xl/worksheets/_rels/sheet39.xml.rels",
                         "xl/worksheets/sheet39.xml"]


# ---------------------------------------------------------------------------
# export_swcb.render (not wired into the export yet): the Conc Fr header
# ---------------------------------------------------------------------------

def swcb_render_bytes(idr: dict = SUBMITTED_IDR) -> bytes:
    """
    Run export_swcb.render on a fresh template and serialize the result.
    Takes the IDR row (the submitted one by default).
    Returns the .xlsx bytes.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_swcb.render(workbook, idr, PROJECT, "Benny Bowers Contracting Co.", inspector="Genghis Khan",
                               page_number=2)
    workbook.show_only(pages)  # as the dispatcher does with the pages render returns
    return workbook.to_bytes()


@pytest.fixture(scope="module")
def full_swcb() -> tuple[bytes, openpyxl.Workbook]:
    """
    The rendered SWCB workbook, fully loaded once for print-setup checks.
    Takes nothing.
    Returns (the bytes, the loaded workbook).
    """
    content = swcb_render_bytes()
    return content, openpyxl.load_workbook(io.BytesIO(content))


class TestSwcbHeader:
    def test_header_fields_land_on_the_conc_fr_cells(self):
        sheet = openpyxl.load_workbook(io.BytesIO(swcb_render_bytes()), read_only=True)["Conc Fr"]
        assert [sheet[c].value for c in ("G8", "P8", "I10", "F12", "F14", "H17")] == [
            "HWS0023", "2024123457", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens",
            "Benny Bowers Contracting Co.", "Genghis Khan",
        ]
        assert sheet["AI4"].value == "9/30/26"  # a General-formatted cell here, so the date goes in as text
        assert sheet["AL5"].fill.fill_type == "solid"  # Wednesday
        assert sheet["AH6"].value is None  # no I.R. No. in the data model yet
        assert (sheet["AH8"].value, sheet["AM8"].value) == (2, 3)
        assert sheet["AG10"].value == "( Start 07:00 End 15:30 )"
        assert sheet["AG12"].value == "( Start 06:45 End ________ )"
        assert (sheet["AD13"].value, sheet["AK13"].value) == ("Low  45", "High  62.5")

    def test_weather_shares_its_box_with_the_am_pm_label(self):
        sheet = openpyxl.load_workbook(io.BytesIO(swcb_render_bytes()), read_only=True)["Conc Fr"]
        assert (sheet["AD15"].value, sheet["AK15"].value) == ("AM\nCloudy", "PM\nRain")
        assert sheet["AD15"].alignment.wrap_text is True
        # Missing weather leaves just the box's label
        idr = {**SUBMITTED_IDR, "weather_am": None}
        assert openpyxl.load_workbook(io.BytesIO(swcb_render_bytes(idr)), read_only=True)["Conc Fr"]["AD15"].value == "AM"

    def test_no_contract_info_formulas_left_on_conc_fr(self):
        sheet = openpyxl.load_workbook(io.BytesIO(swcb_render_bytes()), read_only=True)["Conc Fr"]
        formulas = [c.coordinate for row in sheet.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        assert formulas == []

    def test_only_conc_fr_and_conc_bk_are_visible(self):
        content = swcb_render_bytes()
        assert visible_sheets(content) == ["Conc Fr", "Conc Bk"]
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True).active.title == "Conc Fr"

    def test_both_pages_print_on_one_letter_page(self, full_swcb):
        for name in ("Conc Fr", "Conc Bk"):
            sheet = full_swcb[1][name]
            assert sheet.page_setup.paperSize == 1 and sheet.page_setup.orientation == "portrait", name
            assert sheet.sheet_properties.pageSetUpPr.fitToPage is True, name
            assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1), name

    def test_the_calculation_chain_is_dropped(self, full_swcb):
        package = zipfile.ZipFile(io.BytesIO(full_swcb[0]))
        assert "xl/calcChain.xml" not in package.namelist()
        assert "calcChain" not in package.read("xl/_rels/workbook.xml.rels").decode()


# ---------------------------------------------------------------------------
# export_swcb.render: the Conc Fr body
# ---------------------------------------------------------------------------

def swcb_body(**report_data) -> tuple[bytes, openpyxl.Workbook]:
    """
    Render an SWCB report with the given report_data onto a fresh template.
    Takes report_data fields as keyword arguments.
    Returns (the .xlsx bytes, the workbook opened read-only).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    export_swcb.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data=report_data)
    content = workbook.to_bytes()
    return content, openpyxl.load_workbook(io.BytesIO(content), read_only=True)


def conc_front_row_height(content: bytes, row: int) -> float:
    """
    Read one Conc Fr row's height from an export (Conc Fr is sheet19.xml in the template package).
    Takes the .xlsx bytes and the row number.
    Returns the height in points.
    """
    xml = zipfile.ZipFile(io.BytesIO(content)).read("xl/worksheets/sheet19.xml").decode()
    return float(re.search(rf'<row r="{row}"[^>]*?\sht="([\d.]+)"', xml).group(1))


def matrix(**lines) -> dict:
    """
    Build an Inspection Matrix with the given lines (others absent).
    Takes item key -> {base, sidewalk, curb}.
    Returns the inspectionMatrix dict.
    """
    return lines


OPERATION_CELLS = ("S29", "K29", "E29", "Z29")  # base, sidewalk, curb, structural


class TestSwcbFront:
    def test_description_fills_five_left_aligned_lines_then_continues_on_conc_bk(self):
        _, workbook = swcb_body(description=" ".join(f"word{n}" for n in range(120)))
        sheet = workbook["Conc Fr"]
        lines = [sheet[f"B{row}"].value for row in range(23, 28)]
        assert all(lines) and lines[0].startswith("word0 ")
        assert "continued in ICID" not in lines[-1]
        assert sheet["B22"].value is None and sheet["B28"].value is None
        assert sheet["B23"].alignment.horizontal == "left"
        assert workbook["Conc Bk"]["C20"].value == "Description of work (continued):"

    def test_base_is_ticked_from_a_base_answer(self):
        sheet = swcb_body(inspectionMatrix=matrix(subgradeCompacted={"base": "N", "sidewalk": None, "curb": None}))[1]["Conc Fr"]
        assert [sheet[c].value for c in OPERATION_CELLS] == ["X", None, None, None]

    def test_sidewalk_is_ticked_from_write_in_text(self):
        lines = matrix(otherCuringMethods={"base": "", "sidewalk": "Wet burlap", "curb": None})
        sheet = swcb_body(inspectionMatrix=lines)[1]["Conc Fr"]
        assert [sheet[c].value for c in OPERATION_CELLS] == [None, "X", None, None]

    def test_curb_is_ticked_only_from_columns_that_take_curb(self):
        # Curb answers on items without a Curb box (rows 42 and 43) don't count; one on Rebar does
        lines = matrix(sidewalkFoundationPlaced={"base": None, "sidewalk": None, "curb": "Y"},
                       roadwayStoneBasePlaced={"base": None, "sidewalk": None, "curb": "Y"})
        assert [swcb_body(inspectionMatrix=lines)[1]["Conc Fr"][c].value for c in OPERATION_CELLS] == [None] * 4
        lines["rebarInstalled"] = {"base": None, "sidewalk": None, "curb": "NA"}
        assert [swcb_body(inspectionMatrix=lines)[1]["Conc Fr"][c].value for c in OPERATION_CELLS] == [None, None, "X", None]

    def test_structural_ticks_its_box_with_a_small_centred_x(self):
        sheet = swcb_body(structural=True)[1]["Conc Fr"]
        assert [sheet[c].value for c in OPERATION_CELLS] == [None, None, None, "X"]
        assert sheet["Z29"].font.sz == 6
        assert (sheet["Z29"].alignment.horizontal, sheet["Z29"].alignment.vertical) == ("center", "center")
        assert swcb_body(structural="yes")[1]["Conc Fr"]["Z29"].value is None  # only a real true ticks it

    def test_subcontractor_goes_on_the_operation_line_and_shrinks_to_fit(self):
        sheet = swcb_body(subcontractor="  Acme Concrete Corp.  ")[1]["Conc Fr"]
        assert sheet["AJ29"].value == "Acme Concrete Corp."
        assert sheet["AJ29"].alignment.shrink_to_fit is True

    def test_detailed_activity_rows(self):
        activity = {"excavation": {"fromStation": "9+50", "toStation": "10+00", "remarks": "Trench"},
                    "formPrep": {"fromStation": "", "toStation": "", "remarks": "Forms set"},
                    "pour": {"fromStation": "10+00", "toStation": "10+40", "remarks": "Curb pour"}}
        sheet = swcb_body(activity=activity)[1]["Conc Fr"]
        assert [[sheet[f"{c}{row}"].value for c in ("L", "R", "X")] for row in (33, 34, 35)] == [
            ["9+50", "10+00", "Trench"], [None, None, "Forms set"], ["10+00", "10+40", "Curb pour"],
        ]
        assert [sheet[f"{c}36"].value for c in ("L", "R", "X")] == [None, None, None]  # the spare row stays blank

    def test_matrix_answers_go_in_the_chosen_box(self):
        lines = matrix(subgradeCompacted={"base": "Y", "sidewalk": "N", "curb": "NA"})
        sheet = swcb_body(inspectionMatrix=lines)[1]["Conc Fr"]
        boxes = ("X", "Z", "AB", "AD", "AF", "AH", "AJ", "AL", "AN")
        assert [sheet[f"{c}40"].value for c in boxes] == ["X", None, None, None, "X", None, None, None, "X"]

    def test_sidewalk_foundation_row_stamps_base_and_sidewalk(self):
        # The template's row 42 has Base and Sidewalk boxes; Curb is merged away, so a Curb answer is dropped
        lines = matrix(sidewalkFoundationPlaced={"base": "Y", "sidewalk": "N", "curb": "NA"})
        sheet = swcb_body(inspectionMatrix=lines)[1]["Conc Fr"]
        assert [sheet[f"{c}42"].value for c in ("X", "Z", "AB", "AD", "AF", "AH", "AJ")] == ["X", None, None, None, "X", None, None]
        assert [sheet[c].value for c in OPERATION_CELLS] == ["X", "X", None, None]  # Base and Sidewalk ticked, not Curb

    def test_roadway_stone_base_row_stamps_only_base(self):
        lines = matrix(roadwayStoneBasePlaced={"base": "NA", "sidewalk": "Y", "curb": "Y"})
        sheet = swcb_body(inspectionMatrix=lines)[1]["Conc Fr"]
        assert [sheet[f"{c}43"].value for c in ("X", "Z", "AB", "AD", "AJ")] == [None, None, "X", None, None]

    def test_other_curing_methods_text_goes_in_each_column_box(self):
        lines = matrix(otherCuringMethods={"base": "Wet burlap", "sidewalk": "", "curb": "Plastic sheet"})
        sheet = swcb_body(inspectionMatrix=lines)[1]["Conc Fr"]
        assert [sheet[c].value for c in ("X45", "AD45", "AJ45")] == ["Wet burlap", None, "Plastic sheet"]
        assert sheet["X45"].alignment.shrink_to_fit is True

    def test_pay_items_use_conc_fr_columns_and_its_wider_description_budget(self):
        on_one_line_here = 'Corner Steel Faced Concrete Curb (18" Deep)'  # 43 chars: two lines on Gen Fr's 40, one here
        long = ('Thermoplastic Reflectorized Pavement Markings (4" Wide) including surface preparation, primer, '
                "layout and removal of existing markings")
        items = [pay_item(1, description=on_one_line_here, quantityChk="RM"), pay_item(2, description=long)]
        content, workbook = swcb_body(payItems=items)
        sheet = workbook["Conc Fr"]
        assert [sheet[f"{c}49"].value for c in "BFKPU"] == ["4.01 AAS", "12345", "312.50 S.F.", None, on_one_line_here]
        assert (sheet["U49"].font.sz, conc_front_row_height(content, 49)) == (10, 15.0)  # one line: template's 15 pt
        assert (sheet["U50"].font.sz, conc_front_row_height(content, 50)) == (8, 22.5)  # two lines at 8 pt
        assert sheet["B51"].value is None


# ---------------------------------------------------------------------------
# export_swcb.render: the Conc Bk body
# ---------------------------------------------------------------------------

class TestSwcbBack:
    def test_standard_workforce_rows(self):
        workforce = {"superintendent": "1", "foreman": "2", "operators": "", "laborers": "6", "flaggers": "x2"}
        sheet = swcb_body(workforce=workforce)[1]["Conc Bk"]
        assert [sheet[f"G{row}"].value for row in range(5, 10)] == [1, 2, None, 6, "x2"]  # "foreman": older key
        assert [sheet[f"G{row}"].value for row in (10, 11, 12)] == [None, None, None]  # Teamsters, Surveyors, Masons

    def test_added_trades_use_printed_rows_then_blank_rows(self):
        trades = [{"label": "Masons", "count": "3"}, {"label": "Carpenters", "count": "2"},
                  {"label": "Surveyors", "count": "1"}, {"label": "Chauffeurs", "count": "4"}]
        sheet = swcb_body(additionalWorkforce=trades)[1]["Conc Bk"]
        assert (sheet["G11"].value, sheet["G12"].value) == (1, 3)  # Surveyors and Masons are printed here
        assert [(sheet[f"B{row}"].value, sheet[f"G{row}"].value) for row in (13, 14, 15)] == [
            ("Carpenters", 2), ("Chauffeurs", 4), (None, None),
        ]

    def test_too_many_added_trades_end_with_a_count(self):
        trades = [{"label": f"Trade {n}", "count": "1"} for n in range(7)]
        sheet = swcb_body(additionalWorkforce=trades)[1]["Conc Bk"]
        assert [sheet[f"B{row}"].value for row in range(13, 17)] == ["Trade 0", "Trade 1", "Trade 2", "Trade 3"]
        assert (sheet["B17"].value, sheet["G17"].value) == ("+3 more (see ICID)", None)

    def test_standard_equipment_in_the_left_pair(self):
        equipment = {"backhoe": {"model": "CAT 420", "number": "1"}, "truckDump": {"model": "Mack", "number": "2"},
                     "compressor": {"model": "", "number": ""}}
        sheet = swcb_body(equipment=equipment)[1]["Conc Bk"]
        assert (sheet["N6"].value, sheet["V6"].value) == ("CAT 420", 1)
        assert (sheet["N12"].value, sheet["V12"].value) == ("Mack", 2)  # Truck (Dump) is row 12 here
        assert (sheet["N15"].value, sheet["V15"].value) == (None, None)
        assert (sheet["Y6"].value, sheet["AG6"].value) == (None, None)

    def test_added_equipment_on_its_printed_row_second_unit_in_the_right_pair(self):
        extras = [{"label": "Crane", "model": "Grove", "number": "1"}, {"label": "Crane", "model": "Link-Belt", "number": "1"},
                  {"label": "Paving Machine", "model": "Blaw-Knox", "number": "1"},
                  {"label": "Roller – Dynamic", "model": "CB24", "number": "2"}]
        sheet = swcb_body(additionalEquipment=extras)[1]["Conc Bk"]
        assert (sheet["N7"].value, sheet["V7"].value, sheet["Y7"].value, sheet["AG7"].value) == ("Grove", 1, "Link-Belt", 1)
        assert sheet["N8"].value == "Blaw-Knox"
        assert (sheet["N14"].value, sheet["V14"].value) == ("CB24", 2)  # its own row: no variant name added
        assert sheet["I17"].value == " "  # the blank row is untouched

    def test_excavator_and_other_unprinted_equipment_go_on_the_blank_row(self):
        content, workbook = swcb_body(equipment={"excavator": {"model": "PC200", "number": "1"}},
                                      additionalEquipment=[{"label": "Pavement Cutter", "model": "Husqvarna", "number": "1"}])
        sheet = workbook["Conc Bk"]
        # One blank row: the first entry without a printed row gives way to a count of everything that didn't fit
        assert (sheet["I17"].value, sheet["N17"].value, sheet["V17"].value) == ("+2 more (see ICID)", None, None)
        alone = swcb_body(equipment={"excavator": {"model": "PC200", "number": "1"}})[1]["Conc Bk"]
        assert (alone["I17"].value, alone["N17"].value, alone["V17"].value) == ("Excavator", "PC200", 1)

    def test_safety_marks_y_or_n(self):
        sheet = swcb_body(safetyChecks={"plasticBarrels": "Y", "fencing": "N", "arrowBoard": True})[1]["Conc Bk"]
        assert (sheet["N40"].value, sheet["P40"].value) == ("X", None)
        assert (sheet["N46"].value, sheet["P46"].value) == (None, "X")
        assert (sheet["N48"].value, sheet["P48"].value) == ("X", None)  # older reports' True
        assert (sheet["N41"].value, sheet["P41"].value) == (None, None)  # unanswered

    def test_safety_na_leaves_both_boxes_empty(self):
        sheet = swcb_body(safetyChecks={"plates": "NA"}, safetyRemarks={"plates": "None on site"})[1]["Conc Bk"]
        assert (sheet["N47"].value, sheet["P47"].value) == (None, None)
        assert sheet["R47"].value == "N/A — None on site"  # as on Gen Bk: the N/A goes at the start of the remarks

    def test_safety_remarks(self):
        sheet = swcb_body(safetyRemarks={"generalSafety": "Good", "siteCleaned": "  Swept  "})[1]["Conc Bk"]
        assert (sheet["R44"].value, sheet["R49"].value) == ("Good", "Swept")
        assert sheet["R40"].value is None

    def test_comments_go_on_the_remarks_lines_when_the_description_fits(self):
        sheet = swcb_body(description="Poured curb.", comments="Visitor from DEP.")[1]["Conc Bk"]
        assert sheet["C19"].value == "Remarks:"
        assert sheet["C20"].value == "Visitor from DEP."  # no "Comments:" label when only comments continue
        assert sheet["C20"].alignment.horizontal == "left"
        assert [sheet[f"C{row}"].value for row in range(21, 35)] == [None] * 14

    def test_signatures_stay_blank(self):
        sheet = swcb_body(description="Poured curb.")[1]["Conc Bk"]
        assert [sheet[c].value for c in ("C59", "S59", "AE59")] == [None, None, None]
        assert sheet["C60"].value == "Inspector's Signature"


# ---------------------------------------------------------------------------
# export_swcb.render: the description / comments cascade (Conc Fr -> Conc Bk Remarks -> Report Cont)
# ---------------------------------------------------------------------------

def swcb_pages_and_book(report_cont_available: bool = True, **report_data) -> tuple[list[str], openpyxl.Workbook]:
    """
    Render an SWCB report and return the pages render reported along with the result opened read-only.
    Takes whether Report Cont is free, and report_data fields as keyword arguments.
    Returns (render's pages, the workbook).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_swcb.render(workbook, SUBMITTED_IDR, PROJECT, None, inspector="Genghis Khan",
                               report_data=report_data, report_cont_available=report_cont_available)
    return pages, openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)


class TestSwcbCascade:
    def test_level_1_short_description_stays_on_conc_fr(self):
        pages, workbook = swcb_pages_and_book(description="Poured curb along Main St.")
        assert workbook["Conc Fr"]["B23"].value == "Poured curb along Main St."
        assert [workbook["Conc Bk"][f"C{row}"].value for row in range(20, 35)] == [None] * 15
        assert workbook["Conc Bk"]["C52"].value is None
        assert pages == ["Conc Fr", "Conc Bk"]

    def test_level_2_spills_onto_the_conc_bk_remarks(self):
        pages, workbook = swcb_pages_and_book(description=words(80))
        assert all(workbook["Conc Fr"][f"B{row}"].value for row in range(23, 28))
        back = workbook["Conc Bk"]
        assert back["C20"].value == "Description of work (continued):"
        remarks = [back[f"C{row}"].value for row in range(21, 35) if back[f"C{row}"].value]
        assert remarks and remarks[-1].endswith("word79")  # the rest of the description, ending with its last word
        assert back["C52"].value is None  # nothing continues past the Remarks
        assert pages == ["Conc Fr", "Conc Bk"]

    def test_level_3_continues_on_report_cont_and_ticks_attached_pages(self):
        pages, workbook = swcb_pages_and_book(description=words(600))
        back = workbook["Conc Bk"]
        assert all(back[f"C{row}"].value for row in range(20, 35))
        assert back["C52"].value == "X"
        assert back["C52"].font.sz == 6
        assert (back["C52"].alignment.horizontal, back["C52"].alignment.vertical) == ("center", "center")
        cont = workbook["Report Cont"]
        assert cont["B21"].value.startswith("word")
        assert (cont["I10"].value, cont["H19"].value) == ("9/30/26", "Genghis Khan")
        assert cont["B45"].value.endswith("… (continued in ICID)")  # 600 words outrun Report Cont too
        assert pages == ["Conc Fr", "Conc Bk", "Report Cont"]

    def test_comments_follow_the_description_onto_report_cont(self):
        # Enough description to fill the Remarks, so the comments land on Report Cont
        pages, workbook = swcb_pages_and_book(description=words(220), comments="Visitor from DEP at 10am.")
        cont = [workbook["Report Cont"][f"B{row}"].value for row in range(21, 46)]
        assert "Comments:" in cont
        assert cont[cont.index("Comments:") + 1] == "Visitor from DEP at 10am."
        assert pages[-1] == "Report Cont"

    def test_without_report_cont_the_remarks_are_cut_and_nothing_is_ticked(self):
        pages, workbook = swcb_pages_and_book(report_cont_available=False, description=words(600))
        back = workbook["Conc Bk"]
        assert back["C34"].value.endswith("… (continued in ICID)")
        assert len(back["C34"].value) <= 75
        assert back["C52"].value is None
        assert pages == ["Conc Fr", "Conc Bk"]


# ---------------------------------------------------------------------------
# generate_idr_export with an SWCB report (the dispatcher)
# ---------------------------------------------------------------------------

def swcb_report(page_number: int = 2, **report_data) -> dict:
    """
    Build the SWCB report row list_reports_for_idr returns.
    Takes its page number and report_data fields as keyword arguments.
    Returns the row.
    """
    return {"report_id": UUID("1b2c3d4e-5f60-4718-8293-a4b5c6d7e8f9"), "report_type": "SWCB", "is_addendum": False,
            "page_number": page_number, "report_data": report_data}


class TestSwcbExport:
    def test_an_swcb_report_is_exported_on_conc_fr_and_conc_bk(self):
        content = export_bytes(reports=[swcb_report(description="Formed and poured curb.", structural=True)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Fr"]
        assert sheet["B23"].value == "Formed and poured curb."
        assert sheet["Z29"].value == "X"

    def test_the_header_gets_the_inspector_and_the_reports_own_page_number(self):
        content = export_bytes(reports=[swcb_report(page_number=2, description="x")])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Fr"]
        assert sheet["H17"].value == "Genghis Khan"
        assert (sheet["AH8"].value, sheet["AM8"].value) == (2, 3)

    def test_conc_bk_is_shown_for_an_swcb_report_with_no_back_page_content(self):
        content = export_bytes(reports=[swcb_report(description="Poured curb.")])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Bk"]
        assert [sheet[f"C{row}"].value for row in range(20, 35)] == [None] * 15
        assert export_swcb.render(WorkbookTemplate(TEMPLATE), SUBMITTED_IDR, PROJECT, None, report_data={}) == [
            "Conc Fr", "Conc Bk",
        ]

    def test_without_an_swcb_report_conc_fr_stays_hidden(self):
        addendum_only = {**swcb_report(), "is_addendum": True}
        assert visible_sheets(export_bytes(reports=[addendum_only])) == ["Gen Fr", "Gen Bk"]

    def test_report_cont_prints_after_conc_bk_when_the_swcb_uses_it(self):
        content = export_bytes(reports=[swcb_report(description=words(600))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Report Cont"]
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        names = workbook.sheetnames
        assert names.index("Conc Bk") + 1 == names.index("Report Cont")
        assert workbook.active.title == "Gen Fr"
        # The one sheet-scoped name, AC Fr's print area, moved with AC Fr when Report Cont moved past it
        xml = zipfile.ZipFile(io.BytesIO(content)).read("xl/workbook.xml").decode()
        assert re.search(r'localSheetId="(\d+)"', xml).group(1) == str(names.index("AC Fr"))
        assert re.search(r"<definedName [^>]*>([^<]*)</definedName>", xml).group(1) == "'AC Fr'!$A$1:$AQ$63"

    def test_the_general_keeps_report_cont_and_the_swcb_text_is_cut(self):
        content = export_bytes(general=general_with(description=words(600)),
                               reports=[swcb_report(description=words(600))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Report Cont", "Conc Fr", "Conc Bk"]
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert workbook.sheetnames.index("Gen Bk") + 1 == workbook.sheetnames.index("Report Cont")
        back = workbook["Conc Bk"]
        assert back["C34"].value.endswith("… (continued in ICID)")
        assert back["C52"].value is None
        assert workbook["Report Cont"]["B21"].value.startswith("word")  # the General's continuation


# ---------------------------------------------------------------------------
# export_conc_mix.render (not wired into the export yet): the Conc Mix header
# ---------------------------------------------------------------------------

def conc_mix_render(idr: dict = SUBMITTED_IDR, page_number: Optional[int] = 3) -> tuple[list[str], bytes]:
    """
    Run export_conc_mix.render on a fresh template and serialize the result.
    Takes the IDR row and the report's page number (None for an unnumbered draft page).
    Returns (the pages render returned, the .xlsx bytes).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_conc_mix.render(workbook, idr, PROJECT, "Benny Bowers Contracting Co.", inspector="Genghis Khan",
                                   page_number=page_number)
    return pages, workbook.to_bytes()


class TestConcMixHeader:
    def test_header_fields_land_on_the_conc_mix_cells(self):
        sheet = openpyxl.load_workbook(io.BytesIO(conc_mix_render()[1]), read_only=True)["Conc Mix"]
        assert [sheet[c].value for c in ("G8", "P8", "I10", "F12", "F14", "H17")] == [
            "HWS0023", "2024123457", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens",
            "Benny Bowers Contracting Co.", "Genghis Khan",
        ]
        assert sheet["AD8"].value == "9/30/26"  # a General-formatted cell, so the date goes in as text
        assert (sheet["AD10"].value, sheet["AJ10"].value) == (3, 3)
        # Without a page number (a draft) PAGE / OF stay blank
        draft = openpyxl.load_workbook(io.BytesIO(conc_mix_render(DRAFT_IDR, None)[1]), read_only=True)["Conc Mix"]
        assert (draft["AD10"].value, draft["AJ10"].value) == (None, None)

    def test_attachment_to_ir_no_stays_blank(self):
        sheet = openpyxl.load_workbook(io.BytesIO(conc_mix_render()[1]), read_only=True)["Conc Mix"]
        assert sheet["AJ17"].value is None
        assert sheet["Y17"].value == "ATTACHMENT TO I.R. NO.:"  # the label stays

    def test_no_contract_info_formulas_left_on_conc_mix(self):
        # Read the XML: Material Usage's "=" labels (V41:V44, AK41:AK44) are text openpyxl can't tell from formulas
        part = f"xl/worksheets/{SHEET_PARTS['Conc Mix']}"
        assert zipfile.ZipFile(TEMPLATE).read(part).decode().count("<f>") == 5  # G8, P8, I10, F12, F14
        assert "<f>" not in zipfile.ZipFile(io.BytesIO(conc_mix_render()[1])).read(part).decode()

    def test_render_returns_the_conc_mix_page_set_to_one_letter_page(self):
        pages, content = conc_mix_render()
        assert pages == ["Conc Mix"]
        sheet = openpyxl.load_workbook(io.BytesIO(content))["Conc Mix"]
        assert sheet.page_setup.paperSize == 1 and sheet.page_setup.orientation == "portrait"
        assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)

    def test_render_leaves_every_other_sheet_as_the_template_has_it(self):
        rendered = zipfile.ZipFile(io.BytesIO(conc_mix_render()[1]))
        template = zipfile.ZipFile(TEMPLATE)
        others = [n for n in template.namelist()
                  if n.startswith("xl/worksheets/sheet") and n != f"xl/worksheets/{SHEET_PARTS['Conc Mix']}"]
        assert len(others) == 37
        assert [n for n in others if rendered.read(n) != template.read(n)] == []


# ---------------------------------------------------------------------------
# export_conc_mix.render: the Conc Mix body
# ---------------------------------------------------------------------------

def conc_mix_body(**report_data):
    """
    Render a CONC_MIX report with the given report_data and open its Conc Mix sheet.
    Takes report_data fields as keyword arguments.
    Returns the read-only Conc Mix worksheet.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    export_conc_mix.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data=report_data)
    return openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)["Conc Mix"]


def truck(number: int, **fields) -> dict:
    """
    Build one Trucks-table row as the frontend saves it, every field filled.
    Takes the truck's number (used in its values) and any fields to override.
    Returns the row.
    """
    return {"truckOrTicketNo": f"T-{number}", "inspectionSticker": "Y", "loadSizeCy": "10", "endBatch": f"B-{number}",
            "mixingRevs": "70", "startDischTime": "07:30", "endDischTime": "07:50", "slump": "4.5",
            "airContent": "6", "concTemp": "68", "cylinderNumbers": f"C{number}A-C{number}D", **fields}


REMARKS_LINES = [("H", 49)] + [("C", row) for row in range(50, 55)]


def remarks_lines(sheet) -> list:
    """
    Read Conc Mix's six Remarks lines.
    Takes the worksheet.
    Returns the six values (H49, then C50 … C54).
    """
    return [sheet[f"{column}{row}"].value for column, row in REMARKS_LINES]


class TestConcMixBody:
    def test_each_location_of_use_ticks_its_box(self):
        boxes = {"curb": ("F22", "G22"), "sidewalk": ("N22", "O22"), "concreteBase": ("X22", "Y22"),
                 "structural": ("AF22", "AG22")}
        for key, (cell, partner) in boxes.items():
            sheet = conc_mix_body(locationOfUse={**{k: False for k in boxes}, key: True})
            assert [c for c, _ in boxes.values() if sheet[c].value == "X"] == [cell], key
            assert sheet[cell].font.sz == 7, key
            # Centred across the box's two cells (the box straddles the line between them), not merged
            assert [sheet[c].alignment.horizontal for c in (cell, partner)] == ["centerContinuous"] * 2, key
            assert sheet[cell].alignment.vertical == "center", key

    def test_ready_mix_ticks_only_its_box(self):
        sheet = conc_mix_body(mixerType={"type": "readyMix", "otherLabel": "ignored"})
        assert (sheet["O25"].value, sheet["T25"].value, sheet["X25"].value) == ("X", None, None)
        assert sheet["P25"].alignment.horizontal == "centerContinuous"

    def test_other_mixer_ticks_its_box_and_writes_the_type(self):
        sheet = conc_mix_body(mixerType={"type": "other", "otherLabel": " Site batch plant "})
        assert (sheet["O25"].value, sheet["T25"].value, sheet["X25"].value) == (None, "X", "Site batch plant")
        assert sheet["T25"].alignment.horizontal == "centerContinuous"  # Other's box sits inside T25 alone

    def test_trucks_fill_their_rows(self):
        sheet = conc_mix_body(trucks=[truck(1), truck(2), truck(3)])
        columns = ("B", "K", "N", "Q", "T", "W", "Z", "AC", "AH", "AK")
        for number, row in ((1, 28), (2, 29), (3, 30)):
            assert [sheet[f"{c}{row}"].value for c in columns] == [
                f"T-{number}", "10", f"B-{number}", "70", "07:30", "07:50", "4.5", "6", "68", f"C{number}A-C{number}D",
            ]
        assert [sheet[f"{c}31"].value for c in columns] == [None] * 10
        assert (sheet["G31"].value, sheet["I31"].value) == ("Y", "N")  # an empty row keeps its letters

    def test_inspection_sticker_replaces_the_answers_letter(self):
        stickers = ["Y", "N", "NA", None]
        sheet = conc_mix_body(trucks=[truck(i, inspectionSticker=s) for i, s in enumerate(stickers)])
        assert [(sheet[f"G{row}"].value, sheet[f"I{row}"].value) for row in range(28, 32)] == [
            ("X", "N"), ("Y", "X"), ("Y", "N"), ("Y", "N"),
        ]

    def test_class_of_concrete_replaces_the_blank(self):
        assert conc_mix_body(concreteSpecs={"classOfConcrete": "40"})["B41"].value == "Class of Concrete: 40"
        assert conc_mix_body()["B41"].value == "Class of Concrete: ________________"

    def test_slump_and_air_ranges(self):
        sheet = conc_mix_body(concreteSpecs={"slumpMin": "3", "slumpMax": "5", "airMin": "5.5", "airMax": "7.5"})
        assert [sheet[c].value for c in ("H43", "L43", "H44", "L44")] == ["3", "5", "5.5", "7.5"]
        assert (sheet["H45"].value, sheet["L45"].value) == (" ", " ")  # the unlabelled spare row is left alone

    def test_material_usage(self):
        sheet = conc_mix_body(materialUsage={
            "batchReportNo": "BR-77", "noOfTickets": "4", "firstTicketNo": "1001", "lastTicketNo": "1004",
            "quantityDispatched": "40", "quantityReceived": "40", "quantityUsed": "38.5", "quantityWasted": "1.5",
        })
        assert [sheet[f"W{row}"].value for row in range(41, 45)] == ["BR-77", "4", "1001", "1004"]
        assert [sheet[f"AL{row}"].value for row in range(41, 45)] == ["40", "40", "38.5", "1.5"]

    def test_remarks_paragraphs_fill_the_lines_left_aligned(self):
        sheet = conc_mix_body(remarks="Truck 2 held 10 minutes.\n\nCylinders cast on site.\nForms checked.")
        assert remarks_lines(sheet) == [
            "Truck 2 held 10 minutes.", "Cylinders cast on site.", "Forms checked.", None, None, None,
        ]
        assert [sheet[f"{c}{r}"].alignment.horizontal for c, r in REMARKS_LINES[:3]] == ["left"] * 3

    def test_long_remarks_are_cut_on_the_sixth_line(self):
        lines = remarks_lines(conc_mix_body(remarks=words(200)))
        assert all(lines)
        assert len(lines[0]) <= 86 and all(len(line) <= 99 for line in lines[1:])
        assert lines[5].endswith("… (continued in ICID)")



# ---------------------------------------------------------------------------
# export_conc_mix.render: more than 11 trucks continue on clones of Conc Mix
# ---------------------------------------------------------------------------

def conc_mix_sheets(page_number: Optional[int] = None, **report_data) -> tuple[list[str], openpyxl.Workbook]:
    """
    Render a CONC_MIX report with the given report_data and open the result.
    Takes the report's page number and report_data fields as keyword arguments.
    Returns (the sheets render used, the read-only workbook).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_conc_mix.render(workbook, SUBMITTED_IDR, PROJECT, None, page_number=page_number,
                                   report_data=report_data)
    return pages, openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)


def tickets(sheet) -> list:
    """
    Read the Truck or Ticket No column of a Conc Mix sheet's Trucks table.
    Takes the worksheet.
    Returns the eleven values, None where a row is empty.
    """
    return [sheet[f"B{row}"].value for row in range(28, 39)]


class TestConcMixClones:
    def test_eleven_trucks_or_none_need_just_the_one_sheet(self):
        for count in (11, 0):
            pages, book = conc_mix_sheets(trucks=[truck(i) for i in range(1, count + 1)], remarks="Done.")
            assert pages == ["Conc Mix"] and "Conc Mix 2" not in book.sheetnames, count
            assert remarks_lines(book["Conc Mix"]) == ["Done.", None, None, None, None, None], count
        assert tickets(conc_mix_sheets()[1]["Conc Mix"]) == [None] * 11

    def test_twelve_trucks_continue_on_a_second_sheet(self):
        pages, book = conc_mix_sheets(trucks=[truck(i) for i in range(1, 13)], remarks="Pour went well.")
        assert pages == ["Conc Mix", "Conc Mix 2"]
        assert tickets(book["Conc Mix"]) == [f"T-{i}" for i in range(1, 12)]
        assert tickets(book["Conc Mix 2"]) == ["T-12"] + [None] * 10
        assert remarks_lines(book["Conc Mix"]) == ["Pour went well.", "Continued on next page", None, None, None, None]
        # The second sheet is the last: it only says where it continues from
        assert remarks_lines(book["Conc Mix 2"]) == ["(Continued from previous page)", None, None, None, None, None]

    def test_twenty_two_trucks_fill_two_sheets(self):
        pages, book = conc_mix_sheets(trucks=[truck(i) for i in range(1, 23)])
        assert pages == ["Conc Mix", "Conc Mix 2"]
        assert tickets(book["Conc Mix 2"]) == [f"T-{i}" for i in range(12, 23)]
        assert remarks_lines(book["Conc Mix 2"])[:2] == ["(Continued from previous page)", None]

    def test_twenty_three_trucks_take_three_sheets(self):
        pages, book = conc_mix_sheets(trucks=[truck(i) for i in range(1, 24)])
        assert pages == ["Conc Mix", "Conc Mix 2", "Conc Mix 3"]
        assert tickets(book["Conc Mix 3"]) == ["T-23"] + [None] * 10
        # No remarks: the first sheet's only line is the pointer onward
        assert remarks_lines(book["Conc Mix"])[:2] == ["Continued on next page", None]
        assert remarks_lines(book["Conc Mix 2"])[:3] == [
            "(Continued from previous page)", "Continued on next page", None,
        ]
        assert remarks_lines(book["Conc Mix 3"])[:2] == ["(Continued from previous page)", None]

    def test_every_sheet_carries_the_whole_form(self):
        _, book = conc_mix_sheets(
            trucks=[truck(i, inspectionSticker="N") for i in range(1, 13)],
            locationOfUse={"curb": True, "structural": True}, mixerType={"type": "other", "otherLabel": "Mobile"},
            concreteSpecs={"classOfConcrete": "40", "slumpMin": "3", "airMax": "7.5"},
            materialUsage={"batchReportNo": "BR-77", "quantityWasted": "1.5"},
        )
        cells = ("F22", "N22", "X22", "AF22", "O25", "T25", "X25", "B41", "H43", "L44", "W41", "AL44", "I28",
                 "G8", "AD8")
        first, second = ([book[name][c].value for c in cells] for name in ("Conc Mix", "Conc Mix 2"))
        assert first == second
        assert first[:7] == ["X", None, None, "X", None, "X", "Mobile"]
        assert book["Conc Mix 2"]["F22"].alignment.horizontal == "centerContinuous"

    def test_the_inspectors_remarks_print_on_the_first_sheet_only(self):
        _, book = conc_mix_sheets(trucks=[truck(i) for i in range(1, 13)], remarks=words(200))
        first = remarks_lines(book["Conc Mix"])
        assert first[0].startswith("word0") and first[4].endswith("… (continued in ICID)")
        assert first[5] == "Continued on next page"
        assert not any(line and "word" in line for line in remarks_lines(book["Conc Mix 2"]))

    def test_each_sheet_takes_the_next_page_number(self):
        _, book = conc_mix_sheets(page_number=3, trucks=[truck(i) for i in range(1, 24)])
        numbers = [(book[n]["AD10"].value, book[n]["AJ10"].value) for n in ("Conc Mix", "Conc Mix 2", "Conc Mix 3")]
        assert numbers == [(3, 3), (4, 3), (5, 3)]  # OF is the IDR's total_pages, which the dispatcher raises

    def test_clones_follow_their_sheet_in_tab_order(self):
        names = conc_mix_sheets(trucks=[truck(i) for i in range(1, 24)])[1].sheetnames
        start = names.index("Conc Mix")
        assert names[start:start + 3] == ["Conc Mix", "Conc Mix 2", "Conc Mix 3"]


# ---------------------------------------------------------------------------
# generate_idr_export with a CONC_MIX report (the dispatcher)
# ---------------------------------------------------------------------------

GENERAL_ROW = {"report_id": UUID("4e5f6071-8293-4a41-b5c6-d7e8f9a0b1c2"), "report_type": "GEN", "is_addendum": False,
               "parent_report_id": None, "page_number": 1, "report_data": {}}
SWCB_PARENT = object()  # conc_mix_report's default parent: the SWCB report


def conc_mix_report(page_number: Optional[int] = 3, is_addendum: bool = True, parent=SWCB_PARENT,
                    **report_data) -> dict:
    """
    Build the CONC_MIX addendum row list_reports_for_idr returns.
    Takes its page number, its addendum flag, its parent report's id (the SWCB report's unless given; None for none)
    and report_data fields as keyword arguments.
    Returns the row.
    """
    parent_id = swcb_report()["report_id"] if parent is SWCB_PARENT else parent
    return {"report_id": UUID("2c3d4e5f-6071-4829-93a4-b5c6d7e8f9a0"), "report_type": "CONC_MIX",
            "is_addendum": is_addendum, "parent_report_id": parent_id, "page_number": page_number,
            "report_data": report_data}


def print_order_page_numbers(content: bytes) -> list:
    """
    Read each visible sheet's PAGE number in print order.
    Takes the .xlsx bytes.
    Returns one value per visible sheet: AH8 on Gen Fr / Conc Fr / AC Fr (and copies), AD10 on Conc Mix sheets, None on
    the others.
    """
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)

    def page(name: str):
        if name.startswith(("Gen Fr", "Conc Fr", "AC Fr")):
            return book[name]["AH8"].value
        return book[name]["AD10"].value if name.startswith("Conc Mix") else None
    return [page(name) for name in visible_sheets(content)]


def tab_order(content: bytes) -> tuple[list[str], str]:
    """
    Read the visible sheets in tab (print) order, and which sheet AC Fr's print area is scoped to.
    Takes the .xlsx bytes.
    Returns (the visible sheets, the name of the sheet AC Fr's print area is scoped to).
    """
    names = openpyxl.load_workbook(io.BytesIO(content), read_only=True).sheetnames
    xml = zipfile.ZipFile(io.BytesIO(content)).read("xl/workbook.xml").decode()
    return visible_sheets(content), names[int(re.search(r'localSheetId="(\d+)"', xml).group(1))]


class TestConcMixExport:
    def test_a_conc_mix_report_is_exported_on_conc_mix(self):
        content = export_bytes(reports=[conc_mix_report(trucks=[truck(1)], remarks="Pour went well.")])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Mix"]
        assert (sheet["B28"].value, sheet["H49"].value) == ("T-1", "Pour went well.")

    def test_the_header_gets_the_inspector_contractor_and_the_reports_own_page_number(self):
        sheet = openpyxl.load_workbook(io.BytesIO(export_bytes(reports=[conc_mix_report(page_number=3)])),
                                       read_only=True)["Conc Mix"]
        assert (sheet["H17"].value, sheet["F14"].value) == ("Genghis Khan", "Benny Bowers Contracting Co.")
        assert (sheet["AD10"].value, sheet["AJ10"].value) == (3, 3)

    def test_conc_bk_ticks_see_attached_mixing_info_with_an_swcb_report(self):
        content = export_bytes(reports=[swcb_report(description="Poured curb."), conc_mix_report()])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix"]
        box = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Bk"]["Z37"]
        assert (box.value, box.font.sz, box.alignment.horizontal) == ("X", 6, "center")

    def test_without_a_conc_mix_report_the_box_stays_empty(self):
        content = export_bytes(reports=[swcb_report(description="Poured curb.")])
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Bk"]["Z37"].value is None
        assert "Conc Mix" not in visible_sheets(content)

    def test_without_an_swcb_report_conc_bk_stays_hidden_and_unticked(self):
        content = export_bytes(reports=[conc_mix_report()])
        assert "Conc Bk" not in visible_sheets(content)
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Bk"]["Z37"].value is None

    def test_every_conc_mix_report_is_exported(self):
        second = {**conc_mix_report(page_number=4, remarks="Second pour."),
                  "report_id": UUID("3d4e5f60-7182-4930-a4b5-c6d7e8f9a0b1")}
        content = export_bytes(reports=[conc_mix_report(remarks="First pour."), second])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Mix 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [(book[n]["H49"].value, book[n]["AD10"].value) for n in ("Conc Mix", "Conc Mix 2")] == [
            ("First pour.", 3), ("Second pour.", 4)]

    def test_a_conc_mix_report_not_flagged_as_an_addendum_still_exports(self):
        content = export_bytes(reports=[conc_mix_report(is_addendum=False, remarks="Direct API row.")])
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Mix"]["H49"].value == "Direct API row."

    def test_conc_mix_prints_last_and_print_areas_stay_on_their_sheets(self):
        assert tab_order(export_bytes(reports=[swcb_report(description="x"), conc_mix_report()])) == (
            ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix"], "AC Fr")
        assert tab_order(export_bytes(general=general_with(description=words(600)), reports=[conc_mix_report()])) == (
            ["Gen Fr", "Gen Bk", "Report Cont", "Conc Mix"], "AC Fr")

    def test_conc_mix_follows_report_cont_whichever_report_continues_there(self):
        swcb_owns = export_bytes(reports=[swcb_report(description=words(600)), conc_mix_report()])
        assert tab_order(swcb_owns) == (["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Report Cont", "Conc Mix"], "AC Fr")
        general_owns = export_bytes(general=general_with(description=words(600)),
                                    reports=[swcb_report(description="x"), conc_mix_report()])
        assert tab_order(general_owns) == (
            ["Gen Fr", "Gen Bk", "Report Cont", "Conc Fr", "Conc Bk", "Conc Mix"], "AC Fr")
        assert openpyxl.load_workbook(io.BytesIO(general_owns), read_only=True).active.title == "Gen Fr"

    def test_clones_print_after_conc_bk_and_count_in_the_page_numbers(self):
        content = export_bytes(reports=[swcb_report(description="x"),
                                        conc_mix_report(page_number=3, trucks=[truck(i) for i in range(1, 24)])])
        assert tab_order(content) == (
            ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix", "Conc Mix 2", "Conc Mix 3"], "AC Fr")
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        # total_pages 3 (General, SWCB, CONC_MIX) plus two clones
        assert (book["Gen Fr"]["AH8"].value, book["Gen Fr"]["AM8"].value) == (1, 5)
        assert (book["Conc Fr"]["AH8"].value, book["Conc Fr"]["AM8"].value) == (2, 5)
        numbers = [(book[n]["AD10"].value, book[n]["AJ10"].value) for n in ("Conc Mix", "Conc Mix 2", "Conc Mix 3")]
        assert numbers == [(3, 5), (4, 5), (5, 5)]
        assert book["Conc Bk"]["Z37"].value == "X"

    def test_a_report_numbered_after_the_conc_mix_moves_down_by_its_clones(self):
        # The CONC_MIX hangs off the General (page 2), so the SWCB comes after it (page 3) and prints after it too
        content = export_bytes(reports=[GENERAL_ROW, swcb_report(page_number=3, description="x"),
                                        conc_mix_report(page_number=2, parent=GENERAL_ROW["report_id"],
                                                        trucks=[truck(i) for i in range(1, 13)])])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Mix 2", "Conc Fr", "Conc Bk"]
        assert print_order_page_numbers(content) == [1, None, 2, 3, 4, None]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Gen Fr"]["AM8"].value, book["Conc Mix 2"]["AJ10"].value, book["Conc Fr"]["AM8"].value) == (4,) * 3

    def test_a_generals_conc_mix_follows_its_report_cont(self):
        content = export_bytes(general=general_with(description=words(600)),
                               reports=[GENERAL_ROW, swcb_report(page_number=3, description="x"),
                                        conc_mix_report(page_number=2, parent=GENERAL_ROW["report_id"])])
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "Report Cont", "Conc Mix", "Conc Fr", "Conc Bk"], "AC Fr")

    def test_a_generals_conc_mix_clones_stay_together_before_the_swcb(self):
        content = export_bytes(reports=[GENERAL_ROW, swcb_report(page_number=3, description="x"),
                                        conc_mix_report(page_number=2, parent=GENERAL_ROW["report_id"],
                                                        trucks=[truck(i) for i in range(1, 24)])])
        assert tab_order(content) == (
            ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Mix 2", "Conc Mix 3", "Conc Fr", "Conc Bk"], "AC Fr")
        assert print_order_page_numbers(content) == [1, None, 2, 3, 4, 5, None]
        # The mixing information belongs to the General, so the SWCB's Conc Bk doesn't point to it
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Bk"]["Z37"].value is None

    def test_a_conc_mix_without_an_exported_parent_prints_at_the_end(self):
        sewer = {"report_id": UUID("5f607182-93a4-4b52-c6d7-e8f9a0b1c2d3"), "report_type": "SWR",
                 "is_addendum": False, "parent_report_id": None, "page_number": 2, "report_data": {}}
        for parent in (None, sewer["report_id"]):  # a direct-API row with no parent; a Sewer report's addendum
            content = export_bytes(reports=[GENERAL_ROW, sewer, swcb_report(page_number=4, description="x"),
                                            conc_mix_report(page_number=3, parent=parent)])
            assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix"], parent

    def test_a_draft_marks_a_generals_conc_mix_sheets_where_they_print(self):
        content = export_bytes(idr=DRAFT_IDR, reports=[
            {**GENERAL_ROW, "page_number": None}, swcb_report(page_number=None, description="x"),
            conc_mix_report(page_number=None, parent=GENERAL_ROW["report_id"], trucks=[truck(i) for i in range(1, 13)]),
        ])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Mix 2", "Conc Fr", "Conc Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in ("Conc Mix", "Conc Mix 2")] == ["DRAFT - Not for Submission"] * 2

    def test_a_draft_marks_every_conc_mix_sheet_and_leaves_them_unnumbered(self):
        trucks = [truck(i) for i in range(1, 13)]
        content = export_bytes(idr=DRAFT_IDR, reports=[conc_mix_report(page_number=None, trucks=trucks)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Mix 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        for name in ("Conc Mix", "Conc Mix 2"):
            assert book[name]["B1"].value == "DRAFT - Not for Submission", name
            assert (book[name]["AD10"].value, book[name]["AJ10"].value) == (None, None), name

    def test_a_draft_marks_and_fits_the_conc_mix_page(self):
        content = export_bytes(idr=DRAFT_IDR, reports=[conc_mix_report(page_number=None)])
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        sheet = workbook["Conc Mix"]
        assert sheet["B1"].value == "DRAFT - Not for Submission"
        assert (sheet["AD10"].value, sheet["AJ10"].value) == (None, None)
        assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)


# ---------------------------------------------------------------------------
# Several SWCB and CONC_MIX reports in one IDR
# ---------------------------------------------------------------------------

def swcb_row(number: int, page_number: Optional[int], **report_data) -> dict:
    """
    Build the IDR's nth SWCB report row, with an id of its own.
    Takes its number (1 for the first), its page number and report_data fields as keyword arguments.
    Returns the row.
    """
    return {**swcb_report(page_number, **report_data), "report_id": UUID(int=0x5C00 + number)}


def conc_mix_row(number: int, parent: Optional[UUID], page_number: Optional[int], trucks: int = 0,
                 **report_data) -> dict:
    """
    Build the IDR's nth CONC_MIX report row, with an id of its own.
    Takes its number, its parent's id, its page number, how many trucks it carries and report_data fields.
    Returns the row.
    """
    row = conc_mix_report(page_number, parent=parent, trucks=[truck(i) for i in range(1, trucks + 1)], **report_data)
    return {**row, "report_id": UUID(int=0xC300 + number)}


def z37_ticks(content: bytes, backs: tuple[str, ...]) -> list:
    """
    Read the "See attached Concrete Truck and Mixing Information" box on each Conc Bk.
    Takes the .xlsx bytes and the back pages to read.
    Returns each one's Z37 value.
    """
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    return [book[name]["Z37"].value for name in backs]


SWCB_1, SWCB_2 = swcb_row(1, 2)["report_id"], swcb_row(2, 3)["report_id"]
THREE_PAGES_PLUS = {**SUBMITTED_IDR, "total_pages": 5}


class TestSeveralReportsExport:
    def test_two_swcb_reports_each_get_a_conc_fr_and_conc_bk(self):
        content = export_bytes(reports=[swcb_row(1, 2, description="First curb."),
                                        swcb_row(2, 3, description="Second curb.", structural=True)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Fr 2", "Conc Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Fr"]["B23"].value, book["Conc Fr 2"]["B23"].value) == ("First curb.", "Second curb.")
        assert (book["Conc Fr"]["Z29"].value, book["Conc Fr 2"]["Z29"].value) == (None, "X")
        assert print_order_page_numbers(content) == [1, None, 2, None, 3, None]

    def test_three_swcb_reports_print_in_page_order(self):
        content = export_bytes(idr=THREE_PAGES_PLUS, reports=[swcb_row(n, n + 1, description=f"Curb {n}.")
                                                              for n in (1, 2, 3)])
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Fr 2", "Conc Bk 2",
                                       "Conc Fr 3", "Conc Bk 3"], "AC Fr")
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Fr 3"]["B23"].value == "Curb 3."

    def test_conc_mix_reports_follow_their_own_parents(self):
        content = export_bytes(idr=THREE_PAGES_PLUS, reports=[
            GENERAL_ROW, conc_mix_row(1, GENERAL_ROW["report_id"], 2, remarks="General's."),
            swcb_row(1, 3), swcb_row(2, 4), conc_mix_row(2, SWCB_2, 5, remarks="Second SWCB's."),
        ])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Mix", "Conc Fr", "Conc Bk", "Conc Fr 2",
                                           "Conc Bk 2", "Conc Mix 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Mix"]["H49"].value, book["Conc Mix 2"]["H49"].value) == ("General's.", "Second SWCB's.")
        assert print_order_page_numbers(content) == [1, None, 2, 3, None, 4, None, 5]

    def test_a_second_conc_mix_on_the_same_parent_follows_the_first(self):
        content = export_bytes(reports=[swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3, trucks=12, remarks="First."),
                                        conc_mix_row(2, SWCB_1, 4, remarks="Second.")])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix", "Conc Mix 2",
                                           "Conc Mix 3"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["H49"].value for n in ("Conc Mix", "Conc Mix 2", "Conc Mix 3")] == [
            "First.", "(Continued from previous page)", "Second."]

    def test_truck_overflow_sheets_are_numbered_across_the_idr(self):
        content = export_bytes(reports=[swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3, trucks=12),
                                        conc_mix_row(2, SWCB_1, 4, trucks=12)])
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        # Each report's trucks restart at T-1: its 12th truck goes on its own overflow sheet
        assert [book[n]["B28"].value for n in ("Conc Mix", "Conc Mix 2", "Conc Mix 3", "Conc Mix 4")] == [
            "T-1", "T-12", "T-1", "T-12"]
        assert print_order_page_numbers(content)[4:] == [3, 4, 5, 6]

    def test_z37_ticks_only_on_the_conc_bk_whose_swcb_has_a_conc_mix(self):
        first_only = export_bytes(reports=[swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3), swcb_row(2, 4)])
        assert z37_ticks(first_only, ("Conc Bk", "Conc Bk 2")) == ["X", None]
        second_only = export_bytes(reports=[swcb_row(1, 2), swcb_row(2, 3), conc_mix_row(1, SWCB_2, 4)])
        assert z37_ticks(second_only, ("Conc Bk", "Conc Bk 2")) == [None, "X"]

    def test_z37_stays_empty_when_the_conc_mix_belongs_to_the_general(self):
        content = export_bytes(reports=[GENERAL_ROW, conc_mix_row(1, GENERAL_ROW["report_id"], 2), swcb_row(1, 3)])
        assert z37_ticks(content, ("Conc Bk",)) == [None]

    def test_page_numbers_run_in_print_order_across_several_reports_and_overflow(self):
        # General 1, SWCB 2 + its CONC_MIX 3 (23 trucks: 2 extra), SWCB 4 + its CONC_MIX 5 (12 trucks: 1 extra)
        content = export_bytes(idr=THREE_PAGES_PLUS, reports=[
            swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3, trucks=23),
            swcb_row(2, 4), conc_mix_row(2, SWCB_2, 5, trucks=12),
        ])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Mix", "Conc Mix 2",
                                           "Conc Mix 3", "Conc Fr 2", "Conc Bk 2", "Conc Mix 4", "Conc Mix 5"]
        assert print_order_page_numbers(content) == [1, None, 2, None, 3, 4, 5, 6, None, 7, 8]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert {book["Gen Fr"]["AM8"].value, book["Conc Fr 2"]["AM8"].value, book["Conc Mix 5"]["AJ10"].value} == {8}

    def test_report_cont_goes_to_the_first_report_that_needs_it(self):
        content = export_bytes(reports=[swcb_row(1, 2, description="Short."),
                                        swcb_row(2, 3, description=words(600)), swcb_row(3, 4, description=words(600))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Fr 2", "Conc Bk 2",
                                           "Report Cont", "Conc Fr 3", "Conc Bk 3"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Bk 2"]["C52"].value, book["Conc Bk 3"]["C52"].value) == ("X", None)
        assert book["Conc Bk 3"]["C34"].value.endswith("… (continued in ICID)")

    def test_a_draft_marks_every_sheet_of_every_report(self):
        content = export_bytes(idr=DRAFT_IDR, general={**GENERAL, "page_number": None}, reports=[
            swcb_row(1, None), swcb_row(2, None), conc_mix_row(1, SWCB_2, None, trucks=12)])
        shown = visible_sheets(content)
        assert shown == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Conc Fr 2", "Conc Bk 2", "Conc Mix", "Conc Mix 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in shown] == ["DRAFT - Not for Submission"] * len(shown)
        assert print_order_page_numbers(content) == [None] * len(shown)


class TestSeveralReportsRender:
    def test_swcb_render_stamps_the_pair_it_is_given(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.clone_sheet("Conc Fr", "Conc Fr 2")
        workbook.clone_sheet("Conc Bk", "Conc Bk 2")
        pages = export_swcb.render(workbook, SUBMITTED_IDR, PROJECT, None, page_number=3, fronts=["Conc Fr 2"],
                                   back="Conc Bk 2", report_data={"description": "Clone.", "comments": "Back."})
        assert pages == ["Conc Fr 2", "Conc Bk 2"]
        book = written(workbook)
        assert (book["Conc Fr 2"]["B23"].value, book["Conc Fr 2"]["AH8"].value) == ("Clone.", 3)
        assert book["Conc Bk 2"]["C20"].value == "Back."
        # The template's own pair is untouched
        assert (book["Conc Fr"]["B23"].value, book["Conc Fr"]["G8"].value) == (None, "='Contract Info'!C2")

    def test_mark_conc_mix_attached_ticks_the_back_it_is_given(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.clone_sheet("Conc Bk", "Conc Bk 2")
        export_swcb.mark_conc_mix_attached(workbook, "Conc Bk 2")
        book = written(workbook)
        assert (book["Conc Bk"]["Z37"].value, book["Conc Bk 2"]["Z37"].value) == (None, "X")

    def test_conc_mix_render_uses_the_sheets_it_is_given(self):
        workbook = WorkbookTemplate(TEMPLATE)
        sheets = export_conc_mix.allocate_sheets(workbook, {"trucks": [truck(1)]}, first_index=2)
        assert sheets == ["Conc Mix 3"]
        pages = export_conc_mix.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data={"trucks": [truck(1)]},
                                       sheets=sheets)
        assert pages == ["Conc Mix 3"]
        book = written(workbook)
        assert "Conc Mix 2" not in book.sheetnames  # nothing cloned beyond what it was given
        assert (book["Conc Mix 3"]["B28"].value, book["Conc Mix"]["B28"].value) == ("T-1", None)
        with pytest.raises(ValueError, match="need 2 Conc Mix sheets"):
            export_conc_mix.render(workbook, SUBMITTED_IDR, PROJECT, None, sheets=["Conc Mix"],
                                   report_data={"trucks": [truck(i) for i in range(12)]})


# ---------------------------------------------------------------------------
# Pay items past a front page's table continue on copies of it
# ---------------------------------------------------------------------------

def pay_items(count: int) -> list[dict]:
    """
    Build a report's pay items, numbered from 1.
    Takes how many.
    Returns the list.
    """
    return [pay_item(n) for n in range(1, count + 1)]


def item_numbers(sheet, rows: range) -> list:
    """
    Read the Item No. column of a pay-items table.
    Takes the worksheet and the table's rows.
    Returns each row's Item No.
    """
    return [sheet[f"B{row}"].value for row in rows]


GEN_PAY_ROWS, CONC_PAY_ROWS = range(39, 51), range(49, 61)
CONTINUED_FROM_PREVIOUS = "Pay items continued from previous page"


class TestPayItemOverflow:
    def test_items_split_eleven_a_page_and_the_last_page_uses_every_row(self):
        sizes = {count: [len(s) for s in pay_item_slices(pay_items(count), 12)] for count in (0, 11, 12, 13, 23, 25)}
        assert sizes == {0: [0], 11: [11], 12: [12], 13: [11, 2], 23: [11, 12], 25: [11, 11, 3]}

    def test_up_to_twelve_items_need_no_extra_page(self):
        for count in (0, 11, 12):
            content = export_bytes(general=general_with(payItems=pay_items(count)))
            assert visible_sheets(content) == ["Gen Fr", "Gen Bk"], count
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]
        assert sheet["B50"].value == "4.12 AAS"  # the twelfth item takes the last row: no note needed

    def test_a_general_with_25_items_prints_three_fronts(self):
        content = export_bytes(general=general_with(payItems=pay_items(25)))
        assert visible_sheets(content) == ["Gen Fr", "Gen Fr 2", "Gen Fr 3", "Gen Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert item_numbers(book["Gen Fr"], GEN_PAY_ROWS)[:11] == [f"4.{n:02d} AAS" for n in range(1, 12)]
        assert item_numbers(book["Gen Fr 2"], GEN_PAY_ROWS)[:11] == [f"4.{n:02d} AAS" for n in range(12, 23)]
        assert item_numbers(book["Gen Fr 3"], GEN_PAY_ROWS) == ["4.23 AAS", "4.24 AAS", "4.25 AAS"] + [None] * 9
        assert [book[n]["X50"].value for n in ("Gen Fr", "Gen Fr 2", "Gen Fr 3")] == [
            "Pay items continued on next page", "Pay items continued on next page", None]
        assert [book[n]["B22"].value for n in ("Gen Fr 2", "Gen Fr 3")] == [CONTINUED_FROM_PREVIOUS] * 2

    def test_an_overflow_front_has_its_header_and_pay_items_only(self):
        general = general_with(payItems=pay_items(13), description="Poured curb.")
        book = openpyxl.load_workbook(io.BytesIO(export_bytes(general=general)), read_only=True)
        overflow = book["Gen Fr 2"]
        assert [overflow[c].value for c in ("G8", "F14", "H17", "AH8", "AM8")] == [
            "HWS0023", "Benny Bowers Contracting Co.", "Genghis Khan", 2, 4]
        assert book["Gen Fr"]["B22"].value == "Poured curb."
        assert [overflow[f"B{row}"].value for row in range(23, 35)] == [None] * 12  # only the note on B22
        assert overflow["AC36"].value is None

    def test_an_swcb_with_20_items_continues_on_conc_fr_2_before_its_back(self):
        content = export_bytes(reports=[swcb_row(1, 2, description="Curb.", structural=True, payItems=pay_items(20),
                                                 inspectionMatrix={"subgradeCompacted": {"base": "Y"}})])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Fr 2", "Conc Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert item_numbers(book["Conc Fr 2"], CONC_PAY_ROWS)[:10] == [f"4.{n:02d} AAS" for n in range(12, 21)] + [None]
        assert book["Conc Fr"]["U60"].value == "Pay items continued on next page"
        overflow = book["Conc Fr 2"]
        assert (overflow["B23"].value, overflow["AH8"].value, overflow["AM8"].value) == (CONTINUED_FROM_PREVIOUS, 3, 4)
        assert (book["Conc Fr"]["Z29"].value, book["Conc Fr"]["X40"].value) == ("X", "X")
        assert (overflow["Z29"].value, overflow["X40"].value) == (None, None)  # operation and matrix stay blank

    def test_overflow_fronts_number_across_swcb_reports(self):
        content = export_bytes(reports=[swcb_row(1, 2, description="First.", payItems=pay_items(20)),
                                        swcb_row(2, 3, description="Second.", payItems=pay_items(20))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Fr 2", "Conc Bk",
                                           "Conc Fr 3", "Conc Fr 4", "Conc Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B23"].value for n in ("Conc Fr", "Conc Fr 2", "Conc Fr 3", "Conc Fr 4")] == [
            "First.", CONTINUED_FROM_PREVIOUS, "Second.", CONTINUED_FROM_PREVIOUS]
        assert print_order_page_numbers(content) == [1, None, 2, 3, None, 4, 5, None]

    def test_the_generals_overflow_and_report_cont_keep_their_order(self):
        content = export_bytes(general=general_with(payItems=pay_items(13), description=words(600)))
        assert visible_sheets(content) == ["Gen Fr", "Gen Fr 2", "Gen Bk", "Report Cont"]

    def test_general_and_swcb_overflow_both_count_in_the_page_numbers(self):
        content = export_bytes(general=general_with(payItems=pay_items(13)),
                               reports=[swcb_row(1, 2, payItems=pay_items(13))])
        assert print_order_page_numbers(content) == [1, 2, None, 3, 4, None]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert {book[n]["AM8"].value for n in ("Gen Fr", "Gen Fr 2", "Conc Fr", "Conc Fr 2")} == {5}  # 3 + 2 extra

    def test_a_generals_conc_mix_follows_its_whole_group(self):
        content = export_bytes(general=general_with(payItems=pay_items(13)), reports=[
            GENERAL_ROW, conc_mix_row(1, GENERAL_ROW["report_id"], 2), swcb_row(1, 3)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Fr 2", "Gen Bk", "Conc Mix", "Conc Fr", "Conc Bk"]
        assert print_order_page_numbers(content) == [1, 2, None, 3, 4, None]

    def test_z37_still_ticks_on_an_overflowing_swcb(self):
        content = export_bytes(reports=[swcb_row(1, 2, payItems=pay_items(13)), conc_mix_row(1, SWCB_1, 3)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Fr 2", "Conc Bk", "Conc Mix"]
        assert z37_ticks(content, ("Conc Bk",)) == ["X"]

    def test_a_draft_marks_every_overflow_front(self):
        content = export_bytes(idr=DRAFT_IDR, general={**general_with(payItems=pay_items(13)), "page_number": None},
                               reports=[swcb_row(1, None, payItems=pay_items(13))])
        shown = visible_sheets(content)
        assert shown == ["Gen Fr", "Gen Fr 2", "Gen Bk", "Conc Fr", "Conc Fr 2", "Conc Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in shown] == ["DRAFT - Not for Submission"] * len(shown)
        assert print_order_page_numbers(content) == [None] * len(shown)

    def test_a_composed_generals_overflow_is_unnumbered_and_not_counted(self):
        swcb = swcb_row(1, 1, payItems=pay_items(13))
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 1}, general=None, main_reports=[swcb],
                               reports=[swcb])
        assert visible_sheets(content) == ["Gen Fr", "Gen Fr 2", "Gen Bk", "Conc Fr", "Conc Fr 2", "Conc Bk"]
        assert print_order_page_numbers(content) == [None, None, None, 1, 2, None]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Conc Fr 2"]["AM8"].value == 2  # the SWCB's extra page counts; the composed General's don't

    def test_render_refuses_fronts_that_dont_fit_the_items(self):
        with pytest.raises(ValueError, match="need 2 front pages, got 1"):
            export_swcb.render(WorkbookTemplate(TEMPLATE), SUBMITTED_IDR, PROJECT, None, fronts=["Conc Fr"],
                               report_data={"payItems": pay_items(13)})


# ---------------------------------------------------------------------------
# Report attachments: one page each, after their report's last page
# ---------------------------------------------------------------------------

def image_bytes(image_format: str, size: tuple[int, int] = (40, 30), orientation: Optional[int] = None) -> bytes:
    """
    Make an image file in memory.
    Takes the Pillow format name (JPEG, PNG, WEBP, HEIF), its size and an optional EXIF orientation.
    Returns the file's bytes.
    """
    from PIL import Image
    image = Image.new("RGB", size, "red")
    buffer = io.BytesIO()
    if orientation is None:
        image.save(buffer, image_format)
    else:
        exif = Image.Exif()
        exif[0x0112] = orientation
        image.save(buffer, image_format, exif=exif)
    return buffer.getvalue()


def attachment(number: int, report_id: UUID, file_type: str = "image/jpeg", **fields) -> dict:
    """
    Build an uploaded attachment row as the export query returns it.
    Takes its number (for distinct ids, names and paths), its report's id, its MIME type and field overrides.
    Returns the row.
    """
    return {"attachment_id": UUID(int=0xA700 + number), "report_id": report_id, "file_name": f"photo_{number}.jpg",
            "file_type": file_type, "file_size_bytes": 1000, "storage_path": f"{report_id}/{number}_photo.jpg",
            "uploaded_by": REPORTER, "uploaded_at": datetime(2026, 9, 30, 9, number % 60),
            "attachment_name": f"Photo {number}", "attachment_description": f"What photo {number} shows.",
            "is_uploaded": True, **fields}


def added_anchors(content: bytes, sheet: str) -> tuple[zipfile.ZipFile, str, list[str]]:
    """
    Find the one-cell anchors (pictures, text boxes) this export added to a sheet's drawing.
    Takes the .xlsx bytes and the sheet name.
    Returns (the package, the drawing part's name, each anchor's XML).
    """
    package = zipfile.ZipFile(io.BytesIO(content))
    workbook = package.read("xl/workbook.xml").decode()
    rel_id = re.search(rf'<sheet name="{sheet}"[^>]*r:id="(rId\d+)"', workbook).group(1)
    target = re.search(rf'Id="{rel_id}"[^>]*Target="([^"]+)"', package.read("xl/_rels/workbook.xml.rels").decode())
    part = "xl/" + target.group(1)
    sheet_rels = package.read(part.replace("worksheets/", "worksheets/_rels/") + ".rels").decode()
    drawing = "xl/drawings/" + re.search(r'Target="\.\./drawings/([^"]+)"', sheet_rels).group(1)
    return package, drawing, re.findall(r"<xdr:oneCellAnchor>.*?</xdr:oneCellAnchor>", package.read(drawing).decode())


def anchor_box(anchor: str) -> dict:
    """
    Read where a one-cell anchor sits and how big it is.
    Takes the anchor's XML.
    Returns its zero-based column and row, its offsets into that cell and its size, in pixels.
    """
    numbers = [int(n) for n in re.findall(r"<xdr:(?:col|colOff|row|rowOff)>(\d+)<", anchor)]
    cx, cy = (int(n) for n in re.search(r'<xdr:ext cx="(\d+)" cy="(\d+)"', anchor).groups())
    return {"col": numbers[0], "col_off": numbers[1] // 9525, "row": numbers[2], "row_off": numbers[3] // 9525,
            "width": cx // 9525, "height": cy // 9525}


def pictures(content: bytes, sheet: str) -> list[dict]:
    """
    Read the pictures this export added to a sheet, through the package's relationships.
    Takes the .xlsx bytes and the sheet name.
    Returns each added picture's anchor (see anchor_box) and its image's format and pixel size.
    """
    from PIL import Image
    package, drawing, anchors = added_anchors(content, sheet)
    media = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"',
                            package.read(drawing.replace("drawings/", "drawings/_rels/") + ".rels").decode()))
    found = []
    for anchor in (a for a in anchors if "<xdr:pic>" in a):
        image = Image.open(io.BytesIO(package.read("xl/media/" + media[re.search(r'r:embed="(rId\d+)"', anchor)
                                                                         .group(1)])))
        found.append({**anchor_box(anchor), "format": image.format, "size": image.size})
    return found


def text_boxes(content: bytes, sheet: str) -> list[dict]:
    """
    Read the text boxes this export added to a sheet.
    Takes the .xlsx bytes and the sheet name.
    Returns each one's anchor (see anchor_box), text, font size, text colour, fill, outline width and colour.
    """
    found = []
    for anchor in (a for a in added_anchors(content, sheet)[2] if 'txBox="1"' in a):
        line = re.search(r'<a:ln w="(\d+)"><a:solidFill><a:srgbClr val="(\w+)"', anchor)
        found.append({**anchor_box(anchor), "text": re.search(r"<a:t>(.*?)</a:t>", anchor).group(1),
                      "size_pt": int(re.search(r'<a:rPr [^>]*sz="(\d+)"', anchor).group(1)) / 100,
                      "color": re.search(r'<a:rPr .*?<a:srgbClr val="(\w+)"', anchor).group(1),
                      "fill": re.search(r'</a:prstGeom><a:solidFill><a:srgbClr val="(\w+)"', anchor).group(1),
                      "outline": (int(line.group(1)), line.group(2)),
                      "centred": 'anchor="ctr"' in anchor and 'algn="ctr"' in anchor})
    return found


def frame_box(box: dict) -> tuple[int, int, int, int]:
    """
    Place an anchor against the photo frame (B25:AI58, 19 px columns and 17 px rows).
    Takes an anchor as anchor_box reads it.
    Returns (its left and top edges' distance from the frame's top-left corner, its width, its height), in pixels.
    """
    return ((box["col"] - 1) * 19 + box["col_off"], (box["row"] - 24) * 17 + box["row_off"], box["width"],
            box["height"])


PDF_TYPE = "application/pdf"
FRAME_WIDTH, FRAME_HEIGHT = 646, 578  # B25:AI58, in pixels


class TestAttachmentsExport:
    def test_a_photo_gets_a_page_after_its_report(self):
        photo = attachment(1, GENERAL_ROW["report_id"])
        content = export_bytes(reports=[GENERAL_ROW], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG")})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Attachments 1"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Attachments 1"]
        assert (sheet["B21"].value, sheet["B22"].value) == ("Photo 1", "What photo 1 shows.")
        assert sheet["B21"].font.b is True
        # 40 x 30 fills the 646 x 578 px frame's width (646 x 484), centred down it: 47 px = 2 rows + 13 px from B25
        assert pictures(content, "Attachments 1") == [{"col": 1, "col_off": 0, "row": 26, "row_off": 13,
                                                       "width": 646, "height": 484, "format": "JPEG",
                                                       "size": (40, 30)}]

    def test_each_photo_gets_its_own_page_in_upload_order(self):
        rows = [attachment(n, SWCB_1) for n in (1, 2, 3)]
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows,
                               files={r["storage_path"]: image_bytes("PNG") for r in rows})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk",
                                           "Attachments 1", "Attachments 2", "Attachments 3"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[f"Attachments {n}"]["B21"].value for n in (1, 2, 3)] == ["Photo 1", "Photo 2", "Photo 3"]

    def test_attachments_follow_their_own_reports(self):
        rows = [attachment(1, GENERAL_ROW["report_id"]), attachment(2, SWCB_1), attachment(3, UUID(int=0xC301))]
        reports = [GENERAL_ROW, swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3), swcb_row(2, 4)]
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 4}, reports=reports, attachments=rows,
                               files={r["storage_path"]: image_bytes("JPEG") for r in rows})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Attachments 1", "Conc Fr", "Conc Bk",
                                           "Attachments 2", "Conc Mix", "Attachments 3", "Conc Fr 2", "Conc Bk 2"]
        # Attachment pages carry no page number, and the reports' numbering is unchanged
        assert print_order_page_numbers(content) == [1, None, None, 2, None, None, 3, None, 4, None]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Attachments 1"]["AA19"].value == "Sheet No.:"
        assert (book["Gen Fr"]["AH8"].value, book["Gen Fr"]["AM8"].value) == (1, 4)

    def test_a_pdf_is_named_on_its_page_and_never_fetched(self):
        pdf = attachment(1, SWCB_1, "application/pdf", file_name="mix_ticket.pdf", attachment_name="Batch ticket")
        with patched_export(reports=[swcb_row(1, 2)], attachments=[pdf]) as mocks:
            content = generate_idr_export(IDR_ID).content
        mocks["download"].assert_not_called()
        sheet = openpyxl.load_workbook(io.BytesIO(content))["Attachments 1"]
        assert (sheet["B21"].value, sheet["B21"].font.sz, sheet["B21"].font.b) == ("PDF: Batch ticket", 16, True)
        assert (sheet["B22"].value, sheet["B22"].font.sz) == ("What photo 1 shows.", 12)
        assert (sheet.row_dimensions[21].height, sheet.row_dimensions[22].height) == (21, 15.75)  # room for them
        assert (sheet["B59"].value, sheet["B59"].font.sz, sheet["B59"].font.color.rgb) == (
            "File: mix_ticket.pdf", 10, "FF808080")
        assert sheet["B26"].value is None
        assert pictures(content, "Attachments 1") == []

    def test_a_pdfs_page_has_a_no_preview_box_over_the_photo_frame(self):
        pdf = attachment(1, SWCB_1, "application/pdf", file_name="mix_ticket.pdf")
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[pdf])
        assert text_boxes(content, "Attachments 1") == [{
            "col": 1, "col_off": 0, "row": 24, "row_off": 0, "width": 646, "height": 578,  # B25:AI58
            "text": "No preview available in this export — see ICID for the full file", "size_pt": 14,
            "color": "808080", "fill": "FFFFFF", "outline": (9525, "808080"), "centred": True,
        }]
        # A photo's page has its photo there, no box, and keeps the template's 10 pt caption
        photo = attachment(2, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG")})
        assert text_boxes(content, "Attachments 1") == [] and len(pictures(content, "Attachments 1")) == 1
        sheet = openpyxl.load_workbook(io.BytesIO(content))["Attachments 1"]
        assert (sheet["B21"].font.sz, sheet["B22"].font.sz) == (10, 10)

    def test_heic_is_converted_to_jpeg_and_webp_to_png(self):
        rows = [attachment(1, SWCB_1, "image/heic"), attachment(2, SWCB_1, "image/webp")]
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows,
                               files={rows[0]["storage_path"]: image_bytes("HEIF"),
                                      rows[1]["storage_path"]: image_bytes("WEBP")})
        assert [p["format"] for n in (1, 2) for p in pictures(content, f"Attachments {n}")] == ["JPEG", "PNG"]

    def test_a_large_photo_is_shrunk_to_1600_px(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (4000, 3000))})
        assert pictures(content, "Attachments 1")[0]["size"] == (1600, 1200)

    def test_a_sideways_phone_photo_is_stood_upright(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (40, 30), orientation=6)})
        placed = pictures(content, "Attachments 1")[0]
        assert placed["size"] == (30, 40) and (placed["width"], placed["height"]) == (434, 578)  # fits the height

    def test_a_photo_that_cant_be_fetched_gets_a_placeholder_and_frees_its_place(self, monkeypatch):
        monkeypatch.setattr(export_attachments, "MAX_PHOTOS", 2)
        rows = [attachment(n, SWCB_1) for n in (1, 2, 3)]
        files = {rows[0]["storage_path"]: TimeoutError("read timed out"),
                 rows[1]["storage_path"]: image_bytes("JPEG"), rows[2]["storage_path"]: image_bytes("JPEG")}
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows, files=files)
        assert visible_sheets(content)[-3:] == ["Attachments 1", "Attachments 2", "Attachments 3"]  # no closing page
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Attachments 1"]["B21"].value == "Attachment unavailable: Photo 1"
        assert [box["text"] for box in text_boxes(content, "Attachments 1")] == [
            "File: photo_1.jpg (couldn't be fetched for this export)"]
        assert [len(pictures(content, f"Attachments {n}")) for n in (1, 2, 3)] == [0, 1, 1]

    def test_photos_past_the_cap_are_counted_on_a_closing_page(self):
        rows = [attachment(n, SWCB_1) for n in range(1, 52)]
        photo = image_bytes("PNG", (8, 6))
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows + [attachment(52, SWCB_1, PDF_TYPE)],
                               files={r["storage_path"]: photo for r in rows})
        shown = visible_sheets(content)
        # 50 photos, the PDF (not capped) on its own page, then the count
        assert shown[4:] == [f"Attachments {n}" for n in range(1, 53)]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Attachments 51"]["B21"].value == "PDF: Photo 52"
        assert book["Attachments 52"]["B21"].value == "1 more attachment in ICID"
        assert export_attachments.more_attachments_note(3) == "3 more attachments in ICID"

    def test_without_attachments_nothing_is_added_or_fetched(self):
        with patched_export(reports=[swcb_row(1, 2)]) as mocks:
            content = generate_idr_export(IDR_ID).content
        mocks["download"].assert_not_called()
        assert not any(name.startswith("Attachments") for name in openpyxl.load_workbook(
            io.BytesIO(content), read_only=True).sheetnames)
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]

    def test_a_draft_marks_every_attachment_page(self):
        rows = [attachment(1, SWCB_1), attachment(2, SWCB_1, PDF_TYPE)]
        content = export_bytes(idr=DRAFT_IDR, reports=[swcb_row(1, None)], attachments=rows,
                               files={rows[0]["storage_path"]: image_bytes("JPEG")})
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[f"Attachments {n}"]["B1"].value for n in (1, 2)] == ["DRAFT - Not for Submission"] * 2

    def test_the_caption_header_and_a_long_description(self):
        photo = attachment(1, SWCB_1, attachment_description=words(200), attachment_name="N" * 120)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG")})
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Attachments 1"]
        assert sheet["B21"].value == "N" * 82 + "..."
        lines = [sheet[f"B{row}"].value for row in range(22, 25)]
        assert lines[0].startswith("word0") and lines[2].endswith("… (continued in ICID)")
        assert sheet["B25"].value is None  # the frame's first row
        assert [sheet[c].value for c in ("G12", "I13", "F15", "H17", "I19", "U19")] == [
            "HWS0023", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens", "Genghis Khan", "9/30/26", None]

    def test_mixed_types_in_one_report(self):
        rows = [attachment(1, SWCB_1), attachment(2, SWCB_1, PDF_TYPE), attachment(3, SWCB_1, "image/heic")]
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows,
                               files={rows[0]["storage_path"]: image_bytes("JPEG"),
                                      rows[2]["storage_path"]: image_bytes("HEIF")})
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[f"Attachments {n}"]["B21"].value for n in (1, 2, 3)] == ["Photo 1", "PDF: Photo 2", "Photo 3"]
        assert [len(pictures(content, f"Attachments {n}")) for n in (1, 2, 3)] == [1, 0, 1]

    def test_the_frame_is_b25_to_ai58(self):
        assert (export_attachments.PHOTO_FRAME_COL, export_attachments.PHOTO_FRAME_ROW) == ("B", 25)
        assert (export_attachments.PHOTO_FRAME_CX, export_attachments.PHOTO_FRAME_CY) == (
            FRAME_WIDTH * 9525, FRAME_HEIGHT * 9525) == (6153150, 5505450)
        # The template's own sizes: 34 columns of 2.71 characters (19 px), 34 rows of 12.75 pt (17 px)
        sheet = openpyxl.load_workbook(export.TEMPLATE_PATH)["Sketch Cont"]
        assert sheet.column_dimensions["A"].width == 2.7109375 and sheet.column_dimensions["A"].max >= 35
        assert {sheet.row_dimensions[row].height for row in range(25, 59)} == {12.75}

    def test_a_landscape_photo_fills_the_frames_width(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (1200, 900))})
        left, top, width, height = frame_box(pictures(content, "Attachments 1")[0])
        assert (width, height) == (FRAME_WIDTH, 484)  # 646 x 3/4 = 484.5
        assert left == 0 and top == (FRAME_HEIGHT - 484) // 2 == FRAME_HEIGHT - 484 - top  # equal strips: 47 px

    def test_a_portrait_photo_fills_the_frames_height(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (900, 1200))})
        left, top, width, height = frame_box(pictures(content, "Attachments 1")[0])
        assert (width, height) == (434, FRAME_HEIGHT)  # 578 x 3/4 = 433.5
        assert top == 0 and left == (FRAME_WIDTH - 434) // 2 == FRAME_WIDTH - 434 - left  # equal strips: 106 px

    def test_a_square_photo_fills_the_frames_shorter_side(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (1000, 1000))})
        left, top, width, height = frame_box(pictures(content, "Attachments 1")[0])
        assert (width, height) == (FRAME_HEIGHT, FRAME_HEIGHT)
        assert top == 0 and left == (FRAME_WIDTH - FRAME_HEIGHT) // 2 == FRAME_WIDTH - FRAME_HEIGHT - left  # 34 px

    def test_a_sideways_photo_is_measured_after_it_is_stood_upright(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG", (1200, 900), orientation=6)})
        assert frame_box(pictures(content, "Attachments 1")[0]) == (106, 0, 434, FRAME_HEIGHT)  # as a portrait

    def test_a_pdfs_box_fills_the_frame(self):
        pdf = attachment(1, SWCB_1, PDF_TYPE)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[pdf])
        assert [frame_box(box) for box in text_boxes(content, "Attachments 1")] == [
            (0, 0, FRAME_WIDTH, FRAME_HEIGHT)]

    def test_an_unavailable_photos_box_fills_the_frame(self):
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=[photo],
                               files={photo["storage_path"]: TimeoutError("read timed out")})
        box, = text_boxes(content, "Attachments 1")
        assert frame_box(box) == (0, 0, FRAME_WIDTH, FRAME_HEIGHT)
        assert (box["size_pt"], box["color"], box["fill"], box["outline"], box["centred"]) == (
            10, "000000", "FFFFFF", (9525, "808080"), True)
        assert pictures(content, "Attachments 1") == []
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Attachments 1"]["B26"].value is None

    def test_a_photo_a_pdf_and_a_photo_each_fill_their_own_page(self):
        rows = [attachment(1, SWCB_1), attachment(2, SWCB_1, PDF_TYPE), attachment(3, SWCB_1)]
        content = export_bytes(reports=[swcb_row(1, 2)], attachments=rows,
                               files={rows[0]["storage_path"]: image_bytes("JPEG", (1200, 900)),
                                      rows[2]["storage_path"]: image_bytes("JPEG", (900, 1200))})
        assert visible_sheets(content)[-3:] == ["Attachments 1", "Attachments 2", "Attachments 3"]
        assert [[frame_box(p) for p in pictures(content, f"Attachments {n}")] for n in (1, 2, 3)] == [
            [(0, 47, FRAME_WIDTH, 484)], [], [(106, 0, 434, FRAME_HEIGHT)]]
        assert [[frame_box(b) for b in text_boxes(content, f"Attachments {n}")] for n in (1, 2, 3)] == [
            [], [(0, 0, FRAME_WIDTH, FRAME_HEIGHT)], []]

    def test_a_full_idr_keeps_its_tab_order_with_attachments(self):
        ac = ac_row(page_number=2)
        rows = [attachment(1, GENERAL_ROW["report_id"]), attachment(2, ac["report_id"], PDF_TYPE),
                attachment(3, SWCB_1)]
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 3}, reports=[GENERAL_ROW, ac, swcb_row(1, 3)],
                               attachments=rows, files={rows[0]["storage_path"]: image_bytes("JPEG"),
                                                        rows[2]["storage_path"]: image_bytes("PNG")})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Attachments 1", "AC Fr", "AC Bk", "Attachments 2",
                                           "Conc Fr", "Conc Bk", "Attachments 3"]
        assert print_order_page_numbers(content) == [1, None, None, 2, None, None, 3, None, None]
        assert zipfile.ZipFile(io.BytesIO(content)).testzip() is None
        book = openpyxl.load_workbook(io.BytesIO(content))
        assert [book[f"Attachments {n}"]["B21"].value for n in (1, 2, 3)] == ["Photo 1", "PDF: Photo 2", "Photo 3"]
        assert [len(pictures(content, f"Attachments {n}")) for n in (1, 2, 3)] == [1, 0, 1]
        assert [frame_box(b) for b in text_boxes(content, "Attachments 2")] == [(0, 0, FRAME_WIDTH, FRAME_HEIGHT)]

    def test_an_unprinted_reports_attachments_come_last(self):
        sewer = {"report_id": UUID(int=0x5E01), "report_type": "SWR", "is_addendum": False,
                 "parent_report_id": None, "page_number": 3, "report_data": {}}
        rows = [attachment(1, SWCB_1), attachment(2, sewer["report_id"])]
        content = export_bytes(reports=[swcb_row(1, 2), sewer], attachments=rows,
                               files={r["storage_path"]: image_bytes("JPEG") for r in rows})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Attachments 1",
                                           "Attachments 2"]

    def test_the_export_query_leaves_pending_uploads_out(self):
        with patch("api.queries.report_attachments.run_query", return_value=[]) as run:
            assert list_uploaded_attachments_for_reports([SWCB_1, SWCB_2]) == []
        sql, params = run.call_args.args
        assert "WHERE report_id = ANY(%s) AND is_uploaded" in sql
        assert "ORDER BY report_id, uploaded_at, attachment_id" in sql
        assert params == ([SWCB_1, SWCB_2],)


# ---------------------------------------------------------------------------
# export_ac.render and the dispatcher: an AC report on AC Fr / AC Bk (header only so far)
# ---------------------------------------------------------------------------

def ac_row(number: int = 1, page_number: Optional[int] = 2, **report_data) -> dict:
    """
    Build an AC report row as list_reports_for_idr returns it.
    Takes its number (for a distinct id), its page number and report_data fields as keyword arguments.
    Returns the row.
    """
    return {"report_id": UUID(int=0xAC00 + number), "report_type": "AC", "is_addendum": False,
            "parent_report_id": None, "page_number": page_number, "report_data": report_data}


def ac_render(idr: dict = SUBMITTED_IDR, page_number: Optional[int] = 2) -> tuple[list[str], bytes]:
    """
    Run export_ac.render on a fresh template and serialize the result.
    Takes the IDR row and the report's page number.
    Returns (the pages render returned, the .xlsx bytes).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_ac.render(workbook, idr, PROJECT, "Benny Bowers Contracting Co.", inspector="Genghis Khan",
                             page_number=page_number, report_data={})
    return pages, workbook.to_bytes()


class TestAcHeader:
    def test_project_details_and_inspector_land_on_ac_fr(self):
        sheet = openpyxl.load_workbook(io.BytesIO(ac_render()[1]), read_only=True)["AC Fr"]
        assert [sheet[c].value for c in ("G8", "P8", "I10", "F12", "F14", "H17")] == [
            "HWS0023", "2024123457", "Installation of Curb, Sidewalk & Ped-Ramp <Queens>", "Queens",
            "Benny Bowers Contracting Co.", "Genghis Khan",
        ]

    def test_date_day_and_sheet_number(self):
        sheet = openpyxl.load_workbook(io.BytesIO(ac_render()[1]), read_only=True)["AC Fr"]
        assert sheet["AI4"].value == "9/30/26"  # General-formatted, so text
        assert sheet["AL5"].fill.fill_type == "solid"  # Wednesday
        assert sheet["AH6"].value is None  # no I.R. No. in the data model yet
        assert (sheet["AH8"].value, sheet["AM8"].value) == (2, 3)

    def test_times_temperatures_and_weather(self):
        sheet = openpyxl.load_workbook(io.BytesIO(ac_render()[1]), read_only=True)["AC Fr"]
        assert sheet["AG10"].value == "( Start 07:00 End 15:30 )"
        assert sheet["AG12"].value == "( Start 06:45 End ________ )"
        assert (sheet["AD13"].value, sheet["AK13"].value) == ("Low  45", "High  62.5")
        assert (sheet["AD15"].value, sheet["AK15"].value) == ("AM\nCloudy", "PM\nRain")
        assert sheet["AD15"].alignment.wrap_text is True

    def test_no_contract_info_formulas_left_on_ac_fr(self):
        part = f"xl/worksheets/{SHEET_PARTS['AC Fr']}"
        assert zipfile.ZipFile(TEMPLATE).read(part).decode().count("<f>") == 5  # G8, P8, I10, F12, F14
        assert "<f>" not in zipfile.ZipFile(io.BytesIO(ac_render()[1])).read(part).decode()

    def test_a_draft_leaves_the_sheet_number_blank(self):
        sheet = openpyxl.load_workbook(io.BytesIO(ac_render(DRAFT_IDR, None)[1]), read_only=True)["AC Fr"]
        assert (sheet["AH8"].value, sheet["AM8"].value) == (None, None)

    def test_render_returns_both_pages_set_to_one_letter_page(self):
        pages, content = ac_render()
        assert pages == ["AC Fr", "AC Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content))
        for name in pages:
            setup = book[name].page_setup
            assert (setup.paperSize, setup.orientation, setup.fitToWidth, setup.fitToHeight) == (
                1, "portrait", 1, 1), name

    def test_an_empty_report_leaves_ac_bk_blank_but_for_its_labels(self):
        sheet = openpyxl.load_workbook(io.BytesIO(ac_render()[1]), read_only=True)["AC Bk"]
        assert [sheet[c].value for c in ("C5", "C17", "G21", "N21", "L36", "N36", "P36", "AD36", "AG36", "C48")] == [
            None] * 10
        assert [sheet[c].value for c in ("C3", "B21", "I22", "B36", "P35", "D48")] == [
            "Remarks: ", "Superintendent", "Backhoe", "Plastic Barrels", "LOCATION",
            "Attached Pages for Additional Remarks and / or Sketches"]

    def test_render_leaves_every_other_sheet_as_the_template_has_it(self):
        rendered = zipfile.ZipFile(io.BytesIO(ac_render()[1]))
        template = zipfile.ZipFile(TEMPLATE)
        ours = {f"xl/worksheets/{SHEET_PARTS[name]}" for name in ("AC Fr", "AC Bk")}
        others = [n for n in template.namelist() if n.startswith("xl/worksheets/sheet") and n not in ours]
        assert [n for n in others if rendered.read(n) != template.read(n)] == []


class TestAcExport:
    def test_an_ac_report_prints_ac_fr_and_ac_bk(self):
        content = export_bytes(reports=[ac_row(page_number=2)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk"]
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]
        assert (sheet["H17"].value, sheet["AH8"].value, sheet["AM8"].value) == ("Genghis Khan", 2, 3)

    def test_without_an_ac_report_both_pages_stay_hidden(self):
        visible = visible_sheets(export_bytes(reports=[swcb_row(1, 2)]))
        assert "AC Fr" not in visible and "AC Bk" not in visible

    def test_an_ac_addendum_row_isnt_printed_as_an_ac_report(self):
        assert "AC Fr" not in visible_sheets(export_bytes(reports=[{**ac_row(), "is_addendum": True}]))

    def test_a_draft_marks_both_pages(self):
        content = export_bytes(idr=DRAFT_IDR, reports=[ac_row(page_number=None)])
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in ("AC Fr", "AC Bk")] == ["DRAFT - Not for Submission"] * 2
        assert (book["AC Fr"]["AH8"].value, book["AC Fr"]["AM8"].value) == (None, None)

    def test_ac_prints_between_the_general_and_an_swcb_numbered_after_it(self):
        content = export_bytes(reports=[GENERAL_ROW, ac_row(page_number=2), swcb_row(1, 3)])
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Conc Fr", "Conc Bk"], "AC Fr")
        assert print_order_page_numbers(content) == [1, None, 2, None, 3, None]

    def test_ac_prints_after_an_swcb_numbered_before_it(self):
        content = export_bytes(reports=[swcb_row(1, 2), ac_row(page_number=3)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "AC Fr", "AC Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Fr"]["AH8"].value, book["AC Fr"]["AH8"].value) == (2, 3)

    def test_general_ac_swcb_and_conc_mix_print_in_their_slots(self):
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 4}, reports=[
            GENERAL_ROW, ac_row(page_number=2), swcb_row(1, 3), conc_mix_row(1, SWCB_1, 4, trucks=12)])
        assert tab_order(content) == (
            ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Conc Fr", "Conc Bk", "Conc Mix", "Conc Mix 2"], "AC Fr")
        # The CONC_MIX's second sheet takes page 5 and counts in OF; AC keeps its own number
        assert print_order_page_numbers(content) == [1, None, 2, None, 3, None, 4, 5]
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]["AM8"].value == 5

    def test_an_ac_reports_conc_mix_and_attachments_follow_it(self):
        ac = ac_row(page_number=2)
        photo = attachment(1, ac["report_id"])
        content = export_bytes(reports=[ac, conc_mix_row(1, ac["report_id"], 3), swcb_row(1, 4)],
                               attachments=[photo], files={photo["storage_path"]: image_bytes("JPEG")})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Attachments 1", "Conc Mix",
                                           "Conc Fr", "Conc Bk"]

    def test_every_ac_report_prints(self):
        content = export_bytes(reports=[ac_row(1, 2), ac_row(2, 3)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "AC Fr 2", "AC Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Fr"]["AH8"].value, book["AC Fr 2"]["AH8"].value) == (2, 3)

    def test_a_composed_general_leaves_ac_out_of_its_description_and_pay_items(self):
        swcb = swcb_row(1, 1, description="Poured curb.", payItems=[pay_item(1)])
        ac = ac_row(page_number=2, payItems=[pay_item(2)])
        content = export_bytes(general=None, main_reports=[swcb, ac], reports=[swcb, ac])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]
        assert sheet["B22"].value == "Sidewalk, Curb, Concrete Base: Poured curb."
        assert "Asphaltic" not in " ".join(str(sheet[f"B{row}"].value or "") for row in range(22, 35))
        assert [sheet[f"B{row}"].value for row in (39, 40)] == ["4.01 AAS", None]


FULL_SITE_CONDITIONS = {
    "pavingContractor": {"pavingContractorName": "Tri-State Paving", "subcontractor": "Ace Milling",
                         "riceNo": "2.456"},
    "temperature": {"surfaceStart": "52", "surfaceFinish": "61", "ambientStart": "48", "ambientFinish": "58"},
    "maxDensity": {"top": "150.2", "binder": "148.7"},
}


def ac_front(**report_data):
    """
    Render an AC report with the given report_data and open its AC Fr sheet.
    Takes report_data fields as keyword arguments.
    Returns the read-only AC Fr worksheet.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data=report_data)
    return openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)["AC Fr"]


class TestAcSiteConditions:
    def test_paving_contractor_name_is_centred_across_the_subcontractor_span(self):
        sheet = ac_front(pavingContractor={"pavingContractorName": "  Tri-State Paving "})
        assert sheet["L20"].value == "Tri-State Paving"
        # Row 20 has no box or line of its own: centred across L20:X20, the Subcontractor box's span below
        assert [sheet[c].alignment.horizontal for c in ("L20", "X20")] == ["centerContinuous"] * 2
        assert sheet["B20"].value == "Name of Paving Contractor:"

    def test_subcontractor_and_rice_no(self):
        sheet = ac_front(pavingContractor={"subcontractor": "Ace Milling", "riceNo": "2.456"})
        assert (sheet["L21"].value, sheet["L22"].value) == ("Ace Milling", "2.456")
        assert sheet["L23"].value == "(Contractor to Supply)"  # the note under Rice No. stays

    def test_surface_and_ambient_temperatures_fill_their_boxes(self):
        sheet = ac_front(temperature=FULL_SITE_CONDITIONS["temperature"])
        assert [sheet[c].value for c in ("AA23", "AE23", "AI23", "AM23")] == ["52", "61", "48", "58"]
        assert sheet["AA22"].value == "START"

    def test_max_density_is_centred_after_each_label(self):
        sheet = ac_front(maxDensity={"top": "150.2", "binder": "148.7"})
        assert (sheet["Y25"].value, sheet["AH25"].value) == ("150.2", "148.7")
        assert [sheet[c].alignment.horizontal for c in ("Y25", "AD25", "AH25", "AP25")] == ["centerContinuous"] * 4
        assert (sheet["W25"].value, sheet["AE25"].value) == ("TOP:", "BINDER:")

    def test_numbers_saved_as_numbers_stay_numbers(self):
        sheet = ac_front(pavingContractor={"riceNo": 2.456}, temperature={"surfaceStart": 285},
                         maxDensity={"top": 150.2, "binder": 148})
        assert [sheet[c].value for c in ("L22", "AA23", "Y25", "AH25")] == [2.456, 285, 150.2, 148]
        assert all(isinstance(sheet[c].value, (int, float)) for c in ("L22", "AA23", "Y25", "AH25"))

    def test_blank_missing_or_malformed_fields_leave_the_cells_empty(self):
        cells = ("L20", "L21", "L22", "AA23", "AE23", "AI23", "AM23", "Y25", "AH25")
        for data in ({}, {"pavingContractor": {"pavingContractorName": "  ", "riceNo": ""},
                          "temperature": {"surfaceStart": None}, "maxDensity": {"top": "", "binder": "   "}},
                     {"pavingContractor": "x", "temperature": [], "maxDensity": 3}):
            sheet = ac_front(**data)
            assert [sheet[c].value for c in cells] == [None] * 9, data
        assert ac_front()["L20"].alignment.horizontal == "center"  # a blank name leaves the template's style

    def test_an_exported_ac_report_carries_all_nine_fields(self):
        content = export_bytes(reports=[ac_row(page_number=2, **FULL_SITE_CONDITIONS)])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]
        assert [sheet[c].value for c in ("L20", "L21", "L22", "AA23", "AE23", "AI23", "AM23", "Y25", "AH25")] == [
            "Tri-State Paving", "Ace Milling", "2.456", "52", "61", "48", "58", "150.2", "148.7"]

    def test_a_draft_still_gets_its_marker(self):
        content = export_bytes(idr=DRAFT_IDR, reports=[ac_row(page_number=None, **FULL_SITE_CONDITIONS)])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]
        assert (sheet["B1"].value, sheet["L20"].value) == ("DRAFT - Not for Submission", "Tri-State Paving")

    def test_typed_value_keeps_numbers_and_trims_text(self):
        assert [typed_value(v) for v in (3, 2.5, Decimal("1.25"), " 7 ", "", "  ", None, True)] == [
            3, 2.5, Decimal("1.25"), "7", None, None, None, "True"]


PAVEMENT_COLUMNS = ("B", "G", "K", "N", "Q", "T", "W", "AA", "AE", "AI", "AM")


def course(number: int, **fields) -> dict:
    """
    Build one pavement course row as the frontend saves it, every field filled.
    Takes the course's number (used in its values) and field overrides.
    Returns the row.
    """
    return {"itemNo": f"6.{number:02d}", "mixType": "Top", "stationFrom": f"{number}+00", "stationTo": f"{number}+50",
            "lane": "N1", "length": "50", "width": "12", "course": "Top", "designDepth": "1.5", "area": "66.7",
            "weight": f"{number}.5", **fields}


def course_row(sheet, row: int) -> list:
    """
    Read one row of AC Fr's pavement course table.
    Takes the worksheet and the row.
    Returns the eleven values, left to right.
    """
    return [sheet[f"{column}{row}"].value for column in PAVEMENT_COLUMNS]


def ac_front_full(**report_data) -> openpyxl.Workbook:
    """
    Render an AC report and fully load the result, for checks read-only mode can't make (row heights).
    Takes report_data fields as keyword arguments.
    Returns the workbook.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data=report_data)
    return openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))


WORDS = lambda count: " ".join(["asphalt"] * count)  # 8 characters a word with its space


class TestAcPavementCourses:
    def test_one_course_fills_row_28(self):
        sheet = ac_front(pavementCourses=[course(1)])
        assert course_row(sheet, 28) == ["6.01", "Top", "1+00", "1+50", "N1", "50", "12", "Top", "1.5", "66.7", "1.5"]
        assert course_row(sheet, 29) == [None] * 11

    def test_two_courses_fill_rows_28_and_29(self):
        sheet = ac_front(pavementCourses=[course(1), course(2, lane="S1")])
        assert [course_row(sheet, row)[0] for row in (28, 29, 30, 31)] == ["6.01", "6.02", None, None]
        assert course_row(sheet, 29)[4] == "S1"

    def test_four_courses_fill_the_table(self):
        sheet = ac_front(pavementCourses=[course(n) for n in (1, 2, 3, 4)])
        assert [course_row(sheet, row)[0] for row in (28, 29, 30, 31)] == ["6.01", "6.02", "6.03", "6.04"]
        assert [course_row(sheet, row)[10] for row in (28, 29, 30, 31)] == ["1.5", "2.5", "3.5", "4.5"]

    def test_a_fifth_course_moves_to_the_next_sheet(self):
        sheet = ac_front(pavementCourses=[course(n) for n in (1, 2, 3, 4, 5)])
        assert [course_row(sheet, row)[0] for row in (28, 29, 30, 31)] == ["6.01", "6.02", "6.03", "6.04"]
        assert sheet["B32"].value == "Pavement courses continued on next page"

    def test_numbers_stay_numbers_and_the_header_stays(self):
        sheet = ac_front(pavementCourses=[course(1, length=120, area=None, itemNo=" 6.01 "), "not a row"])
        assert course_row(sheet, 28)[:6] == ["6.01", "Top", "1+00", "1+50", "N1", 120]
        assert course_row(sheet, 28)[9] is None
        assert course_row(sheet, 29) == [None] * 11  # a malformed row is skipped
        assert (sheet["K27"].value, sheet["N27"].value) == ("From", "To")


class TestAcMaterialUsage:
    def test_top_and_binder_values_sit_right_of_their_equals_signs(self):
        top = {"noOfTickets": "12", "firstTicketNo": "4401", "lastTicketNo": "4412", "qtyReceived": "240.5",
               "qtyUsed": "236", "qtyWasted": "4.5"}
        binder = {"noOfTickets": 6, "firstTicketNo": "5101", "lastTicketNo": "5106", "qtyReceived": 120,
                  "qtyUsed": 118.5, "qtyWasted": 1.5}
        sheet = ac_front(materialUsageTop=top, materialUsageBinder=binder)
        assert [sheet[c].value for c in ("H34", "H35", "H36", "S34", "S35", "S36")] == [
            "12", "4401", "4412", "240.5", "236", "4.5"]
        assert [sheet[c].value for c in ("AB34", "AB35", "AB36", "AM34", "AM35", "AM36")] == [
            6, "5101", "5106", 120, 118.5, 1.5]

    def test_blank_or_malformed_usage_leaves_the_cells_empty(self):
        cells = ("H34", "H35", "H36", "S34", "S35", "S36", "AB34", "AB35", "AB36", "AM34", "AM35", "AM36")
        for data in ({}, {"materialUsageTop": {"noOfTickets": " "}, "materialUsageBinder": "x"}):
            sheet = ac_front(**data)
            assert [sheet[c].value for c in cells] == [None] * 12, data
        assert ac_front()["G34"].value == "="


class TestAcPayItems:
    def test_items_use_conc_frs_columns_with_the_unit_in_pay_quantity(self):
        sheet = ac_front(payItems=[pay_item(1, quantityChk="RM"), pay_item(2)])
        assert [sheet[f"{c}39"].value for c in ("B", "F", "K", "P", "U")] == [
            "4.01 AAS", "12345", "312.50 S.F.", None, "Sidewalk 1"]
        assert (sheet["B40"].value, sheet["B41"].value) == ("4.02 AAS", None)

    def test_a_short_description_keeps_one_10_pt_line(self):
        book = ac_front_full(payItems=[pay_item(1, description=WORDS(5))])  # 39 characters
        cell = book["AC Fr"]["U39"]
        assert (cell.value, cell.font.sz, book["AC Fr"].row_dimensions[39].height) == (WORDS(5), 10, 15)

    def test_a_longer_description_takes_two_10_pt_lines(self):
        book = ac_front_full(payItems=[pay_item(1, description=WORDS(7))])  # 55 characters: two lines of 46
        cell = book["AC Fr"]["U39"]
        assert (cell.value, cell.font.sz, cell.alignment.wrap_text) == (WORDS(7), 10, True)
        assert book["AC Fr"].row_dimensions[39].height == 25.5

    def test_a_description_too_long_for_10_pt_shrinks_to_8(self):
        book = ac_front_full(payItems=[pay_item(1, description=WORDS(12))])  # 3 lines at 46, 2 at 55
        cell = book["AC Fr"]["U39"]
        assert (cell.value, cell.font.sz, book["AC Fr"].row_dimensions[39].height) == (WORDS(12), 8, 22.5)

    def test_a_description_too_long_for_8_pt_is_cut(self):
        book = ac_front_full(payItems=[pay_item(1, description=WORDS(30))])
        cell = book["AC Fr"]["U39"]
        assert cell.value.endswith("...") and cell.font.sz == 8 and len(cell.value) <= 2 * 55

    def test_ten_items_fill_the_table_and_more_continue(self):
        sheet = ac_front(payItems=pay_items(10))
        assert item_numbers(sheet, range(39, 49)) == [f"4.{n:02d} AAS" for n in range(1, 11)]
        sheet = ac_front(payItems=pay_items(13))
        assert item_numbers(sheet, range(39, 48)) == [f"4.{n:02d} AAS" for n in range(1, 10)]
        assert (sheet["B48"].value, sheet["U48"].value) == (None, "Pay items continued on next page")

    def test_an_exported_ac_report_carries_its_body_and_draft_marker(self):
        content = export_bytes(idr=DRAFT_IDR, reports=[ac_row(
            page_number=None, pavementCourses=[course(1)], materialUsageTop={"noOfTickets": "12"},
            payItems=[pay_item(1)])])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]
        assert (sheet["B1"].value, sheet["B28"].value, sheet["H34"].value, sheet["B39"].value) == (
            "DRAFT - Not for Submission", "6.01", "12", "4.01 AAS")


REQUIREMENT_KEYS = ["subgradeCompacted", "roadwayCleanDry", "acRollerPerSpec", "densityTestsTaken",
                    "spotCheckAcDepth", "tackCoatPerSpec", "tackCoatOnEdges"]


def requirement_row(sheet, row: int) -> tuple:
    """
    Read one requirement row's answer boxes on AC Fr.
    Takes the worksheet and the row.
    Returns (the Y box W, the N box Y, the Remarks box AA).
    """
    return sheet[f"W{row}"].value, sheet[f"Y{row}"].value, sheet[f"AA{row}"].value


def one_requirement(value, remarks="", key="subgradeCompacted") -> tuple:
    """
    Render an AC report answering one requirement and read its row.
    Takes the answer, its remarks and the requirement's key.
    Returns the row's (Y, N, Remarks) values.
    """
    sheet = ac_front(acRequirements={key: {"value": value, "remarks": remarks}})
    return requirement_row(sheet, 51 + REQUIREMENT_KEYS.index(key))


class TestAcRequirements:
    def test_the_answer_boxes_are_laid_out_beside_the_labels(self):
        sheet = ac_front_full()["AC Fr"]
        assert [sheet[c].value for c in ("W50", "Y50", "AA50", "G50")] == ["Y", "N", "REMARKS", "REQUIREMENTS:"]
        merged = {str(r) for r in sheet.merged_cells.ranges}
        for row in range(50, 58):
            assert {f"W{row}:X{row}", f"Y{row}:Z{row}", f"AA{row}:AP{row}"} <= merged, row
        assert "B51:V51" in merged  # the labels keep their own boxes

    def test_y_and_n_put_an_x_in_their_box(self):
        assert one_requirement("Y") == ("X", None, None)
        assert one_requirement("N") == (None, "X", None)

    def test_n_a_opens_the_remarks(self):
        assert one_requirement("NA", "Not part of today's work") == (None, None, "N/A — Not part of today's work")
        assert one_requirement("NA") == (None, None, "N/A")

    def test_remarks_print_with_or_without_an_answer(self):
        assert one_requirement("", "Rain delayed testing") == (None, None, "Rain delayed testing")
        assert one_requirement("Y", " 3 cores taken ") == ("X", None, "3 cores taken")
        assert one_requirement("") == (None, None, None)

    def test_each_requirement_has_its_own_row(self):
        answers = {key: {"value": "Y" if index % 2 else "N", "remarks": key}
                   for index, key in enumerate(REQUIREMENT_KEYS)}
        sheet = ac_front(acRequirements=answers)
        rows = [requirement_row(sheet, row) for row in range(51, 58)]
        assert [remark for _, _, remark in rows] == REQUIREMENT_KEYS
        assert [(y, n) for y, n, _ in rows] == [(None, "X"), ("X", None)] * 3 + [(None, "X")]
        assert sheet["B57"].value == "TACK COAT APPLIED ON ALL EDGES OF HARDWARE"

    def test_remarks_are_left_aligned_and_shrink_to_fit(self):
        cell = ac_front(acRequirements={"densityTestsTaken": {"value": "Y", "remarks": WORDS(20)}})["AA54"]
        assert (cell.value, cell.alignment.horizontal, cell.alignment.shrink_to_fit) == (WORDS(20), "left", True)

    def test_malformed_answers_leave_the_boxes_empty(self):
        for requirements in ("x", {"subgradeCompacted": "Y"}, {"subgradeCompacted": {"value": "maybe"}}):
            sheet = ac_front(acRequirements=requirements)
            assert requirement_row(sheet, 51) == (None, None, None), requirements
            assert sheet["W50"].value == "Y"  # the layout is there all the same

    def test_an_older_boolean_answer_reads_as_y_or_n(self):
        assert one_requirement(True) == ("X", None, None)
        assert one_requirement(False) == (None, "X", None)


class TestAcTackCoat:
    def test_gallons_rate_and_method_fill_their_cells(self):
        sheet = ac_front(tackCoat={"noOfGallons": "40", "gallonsPerSy": 0.05, "applicationMethod": "Spray bar, SS-1h"})
        assert [sheet[c].value for c in ("AA58", "AL58", "Q61")] == ["40", 0.05, "Spray bar, SS-1h"]
        assert (sheet["B58"].value, sheet["L58"].value) == ("QUANTITY OF TACK COAT:", None)  # no field for it

    def test_blank_tack_coat_leaves_the_cells_empty(self):
        for data in ({}, {"tackCoat": {"noOfGallons": " ", "gallonsPerSy": "", "applicationMethod": None}},
                     {"tackCoat": "x"}):
            assert [ac_front(**data)[c].value for c in ("AA58", "AL58", "Q61")] == [None] * 3, data


class TestMergeCells:
    def test_merging_adds_the_range_once(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.merge_cells("Gen Fr", "B60:F60")
        assert "B60:F60" in {str(r) for r in openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))["Gen Fr"]
                             .merged_cells.ranges}
        with pytest.raises(ValueError, match="already merged"):
            workbook.merge_cells("Gen Fr", "B60:F60")

    def test_the_merge_count_stays_right(self):
        workbook = WorkbookTemplate(TEMPLATE)
        before = int(re.search(r'<mergeCells count="(\d+)">', workbook._sheet("AC Fr")).group(1))
        workbook.merge_cells("AC Fr", "W60:X60")
        xml = workbook._sheet("AC Fr")
        assert int(re.search(r'<mergeCells count="(\d+)">', xml).group(1)) == before + 1 == xml.count("<mergeCell ")


class TestAcFrontComplete:
    def test_a_full_ac_report_exports_every_front_section(self):
        report = ac_row(page_number=None, **FULL_SITE_CONDITIONS, pavementCourses=[course(1)],
                        materialUsageTop={"noOfTickets": "12"}, payItems=[pay_item(1)],
                        acRequirements={"subgradeCompacted": {"value": "Y"}, "tackCoatOnEdges": {"value": "NA"}},
                        tackCoat={"noOfGallons": "40"})
        content = export_bytes(idr=DRAFT_IDR, reports=[report])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr"]
        assert [sheet[c].value for c in ("B1", "L20", "AA23", "Y25", "B28", "H34", "B39", "W51", "AA57", "AA58")] == [
            "DRAFT - Not for Submission", "Tri-State Paving", "52", "150.2", "6.01", "12", "4.01 AAS", "X", "N/A", "40"]


def ac_back(**report_data):
    """
    Render an AC report with the given report_data and open its AC Bk sheet.
    Takes report_data fields as keyword arguments.
    Returns the read-only AC Bk worksheet.
    """
    workbook = WorkbookTemplate(TEMPLATE)
    export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data=report_data)
    return openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)["AC Bk"]


def ticket(number: int, **fields) -> dict:
    """
    Build one delivery ticket row as the frontend saves it.
    Takes the ticket's number (used in its values) and field overrides.
    Returns the row.
    """
    return {"location": f"Main St lane {number}", "ticketNo": f"T{number:03d}", "temperature": "290", **fields}


def ac_remarks(sheet) -> list:
    """
    Read AC Bk's thirteen remarks lines.
    Takes the worksheet.
    Returns C5 ... C17.
    """
    return [sheet[f"C{row}"].value for row in range(5, 18)]


class TestAcBackRemarks:
    def test_comments_fill_the_remarks_lines_left_aligned(self):
        sheet = ac_back(comments="Paving started at 7:30.\n\nInspector from DOT visited.")
        assert ac_remarks(sheet)[:3] == ["Paving started at 7:30.", "Inspector from DOT visited.", None]
        assert sheet["C5"].alignment.horizontal == "left"
        assert (sheet["C3"].value, sheet["C4"].value) == ("Remarks: ", "Comments, Visitors, Other Work, Etc.")

    def test_long_comments_are_cut_when_report_cont_is_taken(self):
        workbook = WorkbookTemplate(TEMPLATE)
        pages = export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data={"comments": words(400)},
                                 report_cont_available=False)
        lines = ac_remarks(written(workbook)["AC Bk"])
        assert pages == ["AC Fr", "AC Bk"]
        assert all(lines) and all(len(line) <= 75 for line in lines)
        assert lines[12].endswith("… (continued in ICID)")


class TestAcBackWorkforceAndEquipment:
    def test_the_five_roles_land_on_rows_21_to_25(self):
        sheet = ac_back(workforce={"superintendent": "1", "foremen": "2", "operators": "3", "laborers": "6",
                                   "flaggers": "2"})
        assert [sheet[f"G{row}"].value for row in range(21, 26)] == [1, 2, 3, 6, 2]

    def test_a_singular_older_key_still_counts(self):
        assert ac_back(workforce={"foreman": "2"})["G22"].value == 2

    def test_added_trades_use_their_preprinted_rows_then_the_blank_ones(self):
        trades = [{"label": "Teamsters", "count": "2"}, {"label": "Masons", "count": "1"},
                  {"label": "Electricians", "count": "3"}]
        sheet = ac_back(additionalWorkforce=trades)
        assert (sheet["G26"].value, sheet["G28"].value) == (2, 1)
        assert (sheet["B29"].value, sheet["G29"].value) == ("Electricians", 3)

    def test_too_many_added_trades_end_with_a_count(self):
        trades = [{"label": f"Trade {n}", "count": "1"} for n in range(1, 8)]  # 5 blank rows for 7 trades
        sheet = ac_back(additionalWorkforce=trades)
        assert [sheet[f"B{row}"].value for row in range(29, 34)] == [
            "Trade 1", "Trade 2", "Trade 3", "Trade 4", "+3 more (see ICID)"]

    def test_standard_equipment_lands_on_its_rows(self):
        equipment = {"frontEndLoader": {"model": "CAT 950", "number": "1"},
                     "backhoe": {"model": "JD 310", "number": "1"},
                     "truckDump": {"model": "Mack", "number": "4"},
                     "compressor": {"model": "185 CFM", "number": "1"}}
        sheet = ac_back(equipment=equipment)
        assert [(sheet[f"N{row}"].value, sheet[f"V{row}"].value) for row in (21, 22, 28, 31)] == [
            ("CAT 950", 1), ("JD 310", 1), ("Mack", 4), ("185 CFM", 1)]

    def test_excavator_has_no_row_so_it_takes_the_blank_one(self):
        sheet = ac_back(equipment={"excavator": {"model": "PC200", "number": "1"}})
        assert (sheet["I33"].value, sheet["N33"].value, sheet["V33"].value) == ("Excavator", "PC200", 1)

    def test_added_paving_equipment_uses_its_preprinted_rows(self):
        added = [{"label": "Paving Machine", "model": "Vogele 1900", "number": "1"},
                 {"label": "Roller – Static", "model": "Hamm HD+", "number": "2"},
                 {"label": "Sweepers", "model": "Elgin", "number": "1"}]
        sheet = ac_back(additionalEquipment=added)
        assert [(sheet[f"N{row}"].value, sheet[f"V{row}"].value) for row in (24, 29, 26)] == [
            ("Vogele 1900", 1), ("Hamm HD+", 2), ("Elgin", 1)]

    def test_too_much_added_equipment_ends_with_a_count(self):
        added = [{"label": f"Thing {n}", "model": "M", "number": "1"} for n in range(1, 4)]  # one blank row
        assert ac_back(additionalEquipment=added)["I33"].value == "+3 more (see ICID)"


class TestAcBackSafety:
    def test_y_and_n_use_ac_bks_own_columns(self):
        sheet = ac_back(safetyChecks={"plasticBarrels": "Y", "siteCleaned": "N"})
        assert (sheet["L36"].value, sheet["N36"].value) == ("X", None)
        assert (sheet["L45"].value, sheet["N45"].value) == (None, "X")
        assert sheet["B45"].value == "Site Cleaned and Secured"

    def test_n_a_and_safety_remarks_have_nowhere_to_go(self):
        sheet = ac_back(safetyChecks={"fencing": "NA", "plates": "Y"},
                        safetyRemarks={"fencing": "Not needed", "plates": "Two plates on Main"})
        assert (sheet["L42"].value, sheet["N42"].value) == (None, None)
        assert (sheet["L43"].value, sheet["N43"].value) == ("X", None)
        # The cells right of Y / N belong to the ticket log: no remark lands there
        assert (sheet["P42"].value, sheet["P43"].value) == (None, None)


class TestAcBackDeliveryTickets:
    def test_tickets_fill_the_log_beside_the_safety_list(self):
        sheet = ac_back(deliveryTickets=[ticket(1), ticket(2, temperature=285)])
        assert [(sheet[f"P{r}"].value, sheet[f"AD{r}"].value, sheet[f"AG{r}"].value) for r in (36, 37, 38)] == [
            ("Main St lane 1", "T001", "290"), ("Main St lane 2", "T002", 285), (None, None, None)]

    def test_ten_tickets_fill_every_row_without_a_note(self):
        sheet = ac_back(deliveryTickets=[ticket(n) for n in range(1, 11)], comments="Done.")
        assert [sheet[f"AD{row}"].value for row in range(36, 46)] == [f"T{n:03d}" for n in range(1, 11)]
        assert ac_remarks(sheet)[:2] == ["Done.", None]

    def test_more_than_ten_tickets_are_counted_after_the_comments(self):
        sheet = ac_back(deliveryTickets=[ticket(n) for n in range(1, 14)], comments="Done.")
        assert sheet["AD45"].value == "T010"
        assert ac_remarks(sheet)[:3] == ["Done.", "[Note] 3 more delivery tickets — see ICID", None]
        assert ac_remarks(ac_back(deliveryTickets=[ticket(n) for n in range(1, 12)]))[0] == (
            "[Note] 1 more delivery ticket — see ICID")

    def test_malformed_tickets_are_skipped(self):
        sheet = ac_back(deliveryTickets=["x", ticket(1), None])
        assert (sheet["AD36"].value, sheet["AD37"].value) == ("T001", None)


class TestAcAttachedPagesBox:
    def test_ticked_when_the_ac_report_has_attachments(self):
        ac = ac_row(page_number=2)
        photo = attachment(1, ac["report_id"])
        content = export_bytes(reports=[ac], attachments=[photo], files={photo["storage_path"]: image_bytes("JPEG")})
        box = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Bk"]["C48"]
        assert (box.value, box.font.sz, box.alignment.horizontal) == ("X", 6, "center")

    def test_empty_without_attachments_or_with_only_another_reports(self):
        assert openpyxl.load_workbook(io.BytesIO(export_bytes(reports=[ac_row()])), read_only=True)[
            "AC Bk"]["C48"].value is None
        photo = attachment(1, SWCB_1)
        content = export_bytes(reports=[ac_row(page_number=2), swcb_row(1, 3)], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG")})
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Bk"]["C48"].value is None

    def test_mark_attachments_ticks_the_box(self):
        workbook = WorkbookTemplate(TEMPLATE)
        export_ac.mark_attachments(workbook)
        assert written(workbook)["AC Bk"]["C48"].value == "X"


class TestAcComplete:
    def test_a_full_ac_report_exports_both_pages(self):
        report = ac_row(page_number=None, **FULL_SITE_CONDITIONS, pavementCourses=[course(1)], payItems=[pay_item(1)],
                        acRequirements={"subgradeCompacted": {"value": "Y"}}, tackCoat={"noOfGallons": "40"},
                        comments="Paved Main St.", workforce={"foremen": "1"},
                        equipment={"compressor": {"model": "185", "number": "1"}},
                        safetyChecks={"plates": "Y"}, deliveryTickets=[ticket(1)])
        content = export_bytes(idr=DRAFT_IDR, reports=[report])
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        back = book["AC Bk"]
        assert [back[c].value for c in ("B1", "C5", "G22", "N31", "L43", "AD36", "C55", "S55")] == [
            "DRAFT - Not for Submission", "Paved Main St.", 1, "185", "X", "T001", None, None]  # signatures blank
        assert (book["AC Fr"]["L20"].value, book["AC Fr"]["W51"].value) == ("Tri-State Paving", "X")
        assert back["B51"].value.startswith("The above described work was incorporated")


def ac_sheets(report_cont_available: bool = True, **report_data) -> tuple[list[str], openpyxl.Workbook]:
    """
    Render an AC report (cloning its AC Fr sheets itself) and open the result.
    Takes whether Report Cont is free and report_data fields as keyword arguments.
    Returns (the pages render used, the read-only workbook).
    """
    workbook = WorkbookTemplate(TEMPLATE)
    pages = export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, page_number=2, report_data=report_data,
                             report_cont_available=report_cont_available)
    return pages, openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()), read_only=True)


def first_courses(book, sheets: tuple[str, ...]) -> list[list]:
    """
    Read each AC Fr sheet's four pavement-course Item Nos.
    Takes the workbook and the sheet names.
    Returns one list of four values per sheet.
    """
    return [[book[name][f"B{row}"].value for row in range(28, 32)] for name in sheets]


LONG = words(600)  # past AC Bk's 13 lines, and past Report Cont's too


class TestAcOverflow:
    def test_front_count_covers_courses_and_pay_items(self):
        courses = {n: export_ac.front_count({"pavementCourses": [course(i) for i in range(n)]})
                   for n in (0, 1, 4, 5, 8, 9, 20)}
        assert courses == {0: 1, 1: 1, 4: 1, 5: 2, 8: 2, 9: 3, 20: 5}
        items = {n: export_ac.front_count({"payItems": pay_items(n)}) for n in (10, 11, 20, 21)}
        assert items == {10: 1, 11: 2, 20: 3, 21: 3}
        assert export_ac.front_count({"pavementCourses": [course(i) for i in range(9)], "payItems": pay_items(11)}) == 3

    def test_five_courses_take_a_second_sheet(self):
        pages, book = ac_sheets(pavementCourses=[course(n) for n in range(1, 6)])
        assert pages == ["AC Fr", "AC Fr 2", "AC Bk"]
        assert first_courses(book, ("AC Fr", "AC Fr 2")) == [
            ["6.01", "6.02", "6.03", "6.04"], ["6.05", None, None, None]]
        assert (book["AC Fr"]["B24"].value, book["AC Fr"]["B32"].value) == (
            None, "Pavement courses continued on next page")
        assert (book["AC Fr 2"]["B24"].value, book["AC Fr 2"]["B32"].value) == ("Continued from previous page", None)
        assert book["AC Fr 2"]["B24"].font.sz == 7

    def test_a_middle_sheet_points_both_ways(self):
        pages, book = ac_sheets(pavementCourses=[course(n) for n in range(1, 10)])
        assert pages == ["AC Fr", "AC Fr 2", "AC Fr 3", "AC Bk"]
        assert [(book[n]["B24"].value, book[n]["B32"].value) for n in ("AC Fr", "AC Fr 2", "AC Fr 3")] == [
            (None, "Pavement courses continued on next page"),
            ("Continued from previous page", "Pavement courses continued on next page"),
            ("Continued from previous page", None)]
        assert first_courses(book, ("AC Fr 3",)) == [["6.09", None, None, None]]

    def test_twenty_courses_take_five_sheets(self):
        pages, book = ac_sheets(pavementCourses=[course(n) for n in range(1, 21)])
        assert pages == ["AC Fr", "AC Fr 2", "AC Fr 3", "AC Fr 4", "AC Fr 5", "AC Bk"]
        assert first_courses(book, ("AC Fr 5",)) == [["6.17", "6.18", "6.19", "6.20"]]

    def test_every_sheet_with_courses_repeats_the_whole_front(self):
        _, book = ac_sheets(**FULL_SITE_CONDITIONS, pavementCourses=[course(n) for n in range(1, 6)],
                            materialUsageTop={"noOfTickets": "12"}, tackCoat={"noOfGallons": "40"},
                            acRequirements={"subgradeCompacted": {"value": "Y"}})
        second = book["AC Fr 2"]
        assert [second[c].value for c in ("G8", "AH8", "AM8", "L20", "AA23", "Y25", "H34", "W50", "W51", "AA58")] == [
            "HWS0023", 3, 3, "Tri-State Paving", "52", "150.2", "12", "Y", "X", "40"]
        assert book["AC Fr"]["AH8"].value == 2  # the report's own number; its next sheet takes the next one

    def test_pay_items_continue_on_a_header_only_sheet(self):
        pages, book = ac_sheets(**FULL_SITE_CONDITIONS, payItems=pay_items(13),
                                acRequirements={"subgradeCompacted": {"value": "Y"}})
        assert pages == ["AC Fr", "AC Fr 2", "AC Bk"]
        second = book["AC Fr 2"]
        assert item_numbers(second, range(39, 49)) == ["4.10 AAS", "4.11 AAS", "4.12 AAS", "4.13 AAS"] + [None] * 6
        assert (second["G8"].value, second["AH8"].value, second["B24"].value) == (
            "HWS0023", 3, "Continued from previous page")
        # Past the last course: header and pay items only
        assert [second[c].value for c in ("L20", "AA23", "B28", "W50", "W51", "AA58")] == [None] * 6
        assert second["B32"].value is None

    def test_twenty_one_items_take_three_sheets(self):
        pages, book = ac_sheets(payItems=pay_items(21))
        assert pages == ["AC Fr", "AC Fr 2", "AC Fr 3", "AC Bk"]
        assert [book[n]["U48"].value for n in ("AC Fr", "AC Fr 2")] == ["Pay items continued on next page"] * 2
        assert item_numbers(book["AC Fr 3"], range(39, 42)) == ["4.19 AAS", "4.20 AAS", "4.21 AAS"]

    def test_the_description_cascade_works_on_a_continuation_sheet(self):
        items = pay_items(10) + [pay_item(11, description=WORDS(7))]
        workbook = WorkbookTemplate(TEMPLATE)
        export_ac.render(workbook, SUBMITTED_IDR, PROJECT, None, report_data={"payItems": items})
        book = openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))
        cell = book["AC Fr 2"]["U40"]  # item 10 on row 39, item 11 on row 40
        assert (cell.value, cell.font.sz, book["AC Fr 2"].row_dimensions[40].height) == (WORDS(7), 10, 25.5)

    def test_courses_and_pay_items_share_the_sheets(self):
        pages, book = ac_sheets(pavementCourses=[course(n) for n in range(1, 10)], payItems=pay_items(13))
        assert pages == ["AC Fr", "AC Fr 2", "AC Fr 3", "AC Bk"]
        assert first_courses(book, ("AC Fr", "AC Fr 2", "AC Fr 3")) == [
            ["6.01", "6.02", "6.03", "6.04"], ["6.05", "6.06", "6.07", "6.08"], ["6.09", None, None, None]]
        assert [book[n]["B39"].value for n in ("AC Fr", "AC Fr 2", "AC Fr 3")] == ["4.01 AAS", "4.10 AAS", None]
        assert [book[n]["U48"].value for n in ("AC Fr", "AC Fr 2")] == ["Pay items continued on next page", None]

    def test_render_refuses_fronts_that_dont_fit(self):
        with pytest.raises(ValueError, match="needs 2 AC Fr sheets, got 1"):
            export_ac.render(WorkbookTemplate(TEMPLATE), SUBMITTED_IDR, PROJECT, None, fronts=["AC Fr"],
                             report_data={"pavementCourses": [course(n) for n in range(5)]})

    def test_the_dispatcher_numbers_the_ac_sheets_and_shifts_what_follows(self):
        ac = ac_row(page_number=2, pavementCourses=[course(n) for n in range(9)])
        content = export_bytes(reports=[GENERAL_ROW, ac, swcb_row(1, 3)])
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "AC Fr", "AC Fr 2", "AC Fr 3", "AC Bk", "Conc Fr",
                                       "Conc Bk"], "AC Fr")
        assert print_order_page_numbers(content) == [1, None, 2, 3, 4, None, 5, None]
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Conc Fr"]["AM8"].value == 5  # 3 + 2


class TestAcRemarksCascade:
    def test_long_remarks_continue_on_report_cont(self):
        content = export_bytes(reports=[ac_row(page_number=2, comments=LONG)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Report Cont"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert ac_remarks(book["AC Bk"])[12] is not None and "continued" not in ac_remarks(book["AC Bk"])[12]
        assert book["Report Cont"]["B21"].value.startswith("word")
        assert book["AC Bk"]["C48"].value == "X"  # its remarks continue on another page

    def test_the_ticket_note_and_safety_remarks_follow_the_comments(self):
        sheet = ac_back(comments="Paved Main St.", deliveryTickets=[ticket(n) for n in range(1, 12)],
                        safetyChecks={"fencing": "NA", "plates": "Y"},
                        safetyRemarks={"fencing": "Not needed", "plates": "Two plates on Main"})
        assert ac_remarks(sheet)[:6] == [
            "Paved Main St.", "[Note] 1 more delivery ticket — see ICID", "Safety check list remarks:",
            "Fencing: N/A — Not needed", "Plates: Two plates on Main", None]

    def test_an_n_a_answer_without_remarks_is_still_listed(self):
        assert ac_remarks(ac_back(safetyChecks={"arrowBoard": "NA"}))[:3] == [
            "Safety check list remarks:", "Arrow Board: N/A", None]
        assert export_ac.safety_remarks({"safetyChecks": {"plates": "Y"}}) == []

    def test_the_general_keeps_report_cont_and_the_ac_remarks_are_cut(self):
        content = export_bytes(general=general_with(description=words(600)),
                               reports=[GENERAL_ROW, ac_row(page_number=2, comments=LONG)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Report Cont", "AC Fr", "AC Bk"]
        back = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Bk"]
        assert ac_remarks(back)[12].endswith("… (continued in ICID)")
        assert back["C48"].value is None

    def test_an_ac_report_before_an_swcb_takes_report_cont(self):
        content = export_bytes(reports=[ac_row(page_number=2, comments=LONG),
                                        swcb_row(1, 3, description=words(600))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Report Cont", "Conc Fr", "Conc Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Conc Bk"]["C34"].value.endswith("… (continued in ICID)")
        assert (book["AC Bk"]["C48"].value, book["Conc Bk"]["C52"].value) == ("X", None)

    def test_an_swcb_before_the_ac_report_takes_report_cont(self):
        content = export_bytes(reports=[swcb_row(1, 2, description=words(600)),
                                        ac_row(page_number=3, comments=LONG)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Report Cont", "AC Fr", "AC Bk"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert ac_remarks(book["AC Bk"])[12].endswith("… (continued in ICID)")
        assert (book["Conc Bk"]["C52"].value, book["AC Bk"]["C48"].value) == ("X", None)

    def test_a_draft_marks_every_ac_sheet_and_report_cont(self):
        content = export_bytes(idr=DRAFT_IDR, general={**GENERAL, "page_number": None}, reports=[ac_row(
            page_number=None, comments=LONG, pavementCourses=[course(n) for n in range(5)], payItems=pay_items(11))])
        shown = visible_sheets(content)
        assert shown == ["Gen Fr", "Gen Bk", "AC Fr", "AC Fr 2", "AC Bk", "Report Cont"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in shown] == ["DRAFT - Not for Submission"] * len(shown)
        assert print_order_page_numbers(content) == [None] * len(shown)


def named_ac(number: int, page_number: Optional[int], name: str, **report_data) -> dict:
    """
    Build an AC report row whose paving contractor names it, so its sheets can be told apart.
    Takes its number, its page number, the contractor's name and other report_data fields.
    Returns the row.
    """
    paving = {"pavingContractorName": name}
    return ac_row(number, page_number, pavingContractor=paving, **report_data)


class TestMultiAcExport:
    def test_two_ac_reports_each_get_ac_fr_and_ac_bk(self):
        content = export_bytes(reports=[named_ac(1, 2, "First Paving", comments="First."),
                                        named_ac(2, 3, "Second Paving", comments="Second.")])
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "AC Fr 2", "AC Bk 2"], "AC Fr")
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Fr"]["L20"].value, book["AC Fr 2"]["L20"].value) == ("First Paving", "Second Paving")
        assert (book["AC Bk"]["C5"].value, book["AC Bk 2"]["C5"].value) == ("First.", "Second.")
        assert print_order_page_numbers(content) == [1, None, 2, None, 3, None]

    def test_the_first_reports_overflow_comes_before_the_second_report(self):
        content = export_bytes(reports=[named_ac(1, 2, "First", pavementCourses=[course(n) for n in range(9)]),
                                        named_ac(2, 3, "Second")])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Fr 2", "AC Fr 3", "AC Bk", "AC Fr 4",
                                           "AC Bk 2"]
        assert print_order_page_numbers(content) == [1, None, 2, 3, 4, None, 5, None]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Fr 3"]["L20"].value, book["AC Fr 4"]["L20"].value) == ("First", "Second")
        assert {book[n]["AM8"].value for n in ("Gen Fr", "AC Fr", "AC Fr 4")} == {5}  # 3 pages + 2 clones

    def test_the_second_reports_overflow_follows_its_own_front(self):
        content = export_bytes(reports=[named_ac(1, 2, "First"),
                                        named_ac(2, 3, "Second", pavementCourses=[course(n) for n in range(5)])])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "AC Fr 2", "AC Fr 3", "AC Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Fr 2"]["B32"].value, book["AC Fr 3"]["B24"].value) == (
            "Pavement courses continued on next page", "Continued from previous page")
        assert book["AC Fr"]["B32"].value is None

    def test_the_first_ac_report_keeps_report_cont_and_the_second_is_cut(self):
        content = export_bytes(reports=[ac_row(1, 2, comments=LONG), ac_row(2, 3, comments=LONG)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Report Cont", "AC Fr 2", "AC Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert ac_remarks(book["AC Bk 2"])[12].endswith("… (continued in ICID)")
        assert (book["AC Bk"]["C48"].value, book["AC Bk 2"]["C48"].value) == ("X", None)

    def test_the_general_keeps_report_cont_over_both_ac_reports(self):
        content = export_bytes(general=general_with(description=words(600)),
                               reports=[ac_row(1, 2, comments=LONG), ac_row(2, 3, comments=LONG)])
        assert visible_sheets(content)[:3] == ["Gen Fr", "Gen Bk", "Report Cont"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert all(ac_remarks(book[n])[12].endswith("… (continued in ICID)") for n in ("AC Bk", "AC Bk 2"))

    def test_an_swcb_between_two_ac_reports_can_take_report_cont(self):
        content = export_bytes(reports=[ac_row(1, 2, comments="Short."), swcb_row(1, 3, description=words(600)),
                                        ac_row(2, 4, comments=LONG)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Conc Fr", "Conc Bk", "Report Cont",
                                           "AC Fr 2", "AC Bk 2"]
        assert ac_remarks(openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Bk 2"])[12].endswith(
            "… (continued in ICID)")

    def test_each_ac_reports_conc_mix_follows_it(self):
        first, second = ac_row(1, 2), ac_row(2, 4)
        content = export_bytes(reports=[first, conc_mix_row(1, first["report_id"], 3, remarks="First's."),
                                        second, conc_mix_row(2, second["report_id"], 5, remarks="Second's.")])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Conc Mix", "AC Fr 2", "AC Bk 2",
                                           "Conc Mix 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["Conc Mix"]["H49"].value, book["Conc Mix 2"]["H49"].value) == ("First's.", "Second's.")

    def test_each_ac_reports_attachments_follow_its_own_back_page(self):
        first, second = ac_row(1, 2), ac_row(2, 3)
        photos = [attachment(1, first["report_id"]), attachment(2, second["report_id"])]
        content = export_bytes(reports=[first, second], attachments=photos,
                               files={p["storage_path"]: image_bytes("JPEG") for p in photos})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "AC Fr", "AC Bk", "Attachments 1", "AC Fr 2", "AC Bk 2",
                                           "Attachments 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Bk"]["C48"].value, book["AC Bk 2"]["C48"].value) == ("X", "X")

    def test_c48_ticks_only_on_the_ac_bk_whose_report_has_attachments(self):
        first, second = ac_row(1, 2), ac_row(2, 3)
        photo = attachment(1, second["report_id"])
        content = export_bytes(reports=[first, second], attachments=[photo],
                               files={photo["storage_path"]: image_bytes("JPEG")})
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert (book["AC Bk"]["C48"].value, book["AC Bk 2"]["C48"].value) == (None, "X")

    def test_a_mixed_idr_prints_and_numbers_in_page_order(self):
        reports = [GENERAL_ROW, named_ac(1, 2, "First", pavementCourses=[course(n) for n in range(5)]),
                   swcb_row(1, 3, payItems=pay_items(13)), conc_mix_row(1, SWCB_1, 4),
                   named_ac(2, 5, "Second")]
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 5}, reports=reports)
        assert tab_order(content) == (["Gen Fr", "Gen Bk", "AC Fr", "AC Fr 2", "AC Bk", "Conc Fr", "Conc Fr 2",
                                       "Conc Bk", "Conc Mix", "AC Fr 3", "AC Bk 2"], "AC Fr")
        assert print_order_page_numbers(content) == [1, None, 2, 3, None, 4, 5, None, 6, 7, None]
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["AC Fr 3"]["AM8"].value == 7  # 5 + 2

    def test_a_draft_marks_every_ac_sheet(self):
        content = export_bytes(idr=DRAFT_IDR, general={**GENERAL, "page_number": None}, reports=[
            ac_row(1, None, pavementCourses=[course(n) for n in range(5)], comments=LONG), ac_row(2, None)])
        shown = visible_sheets(content)
        assert shown == ["Gen Fr", "Gen Bk", "AC Fr", "AC Fr 2", "AC Bk", "Report Cont", "AC Fr 3", "AC Bk 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[n]["B1"].value for n in shown] == ["DRAFT - Not for Submission"] * len(shown)

    def test_mark_attachments_ticks_the_back_it_is_given(self):
        workbook = WorkbookTemplate(TEMPLATE)
        workbook.clone_sheet("AC Bk", "AC Bk 2")
        export_ac.mark_attachments(workbook, "AC Bk 2")
        book = written(workbook)
        assert (book["AC Bk"]["C48"].value, book["AC Bk 2"]["C48"].value) == (None, "X")


# ---------------------------------------------------------------------------
# Draft exports: "DRAFT - Not for Submission" across the top of every printed page
# ---------------------------------------------------------------------------

EXPORT_SHEETS = ("Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Report Cont")


class TestDraftExport:
    def test_every_visible_sheet_of_a_draft_is_marked(self):
        content = export_bytes(idr=DRAFT_IDR, general=general_with(description=words(600)),
                               reports=[swcb_report(description="Poured curb.")])
        shown = visible_sheets(content)
        assert shown == ["Gen Fr", "Gen Bk", "Report Cont", "Conc Fr", "Conc Bk"]
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        for name in shown:
            cell = workbook[name]["B1"]
            assert cell.value == "DRAFT - Not for Submission", name
            assert (cell.font.sz, cell.font.b, cell.font.color.rgb) == (14, True, "FFFF0000"), name
            assert (cell.alignment.horizontal, cell.alignment.vertical) == ("center", "center"), name
            row = re.search(r'<row r="1"[^>]*?\sht="([\d.]+)"', zipfile.ZipFile(io.BytesIO(content)).read(
                f"xl/worksheets/{SHEET_PARTS[name]}").decode()).group(1)
            assert float(row) == 18.75, name  # raised so 14 pt fits
        # A hidden sheet isn't marked
        assert workbook["AC Fr"]["B1"].value is None

    def test_a_submitted_export_has_no_marker(self):
        content = export_bytes(general=general_with(description=words(600)), reports=[swcb_report(description="x")])
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [workbook[name]["B1"].value for name in visible_sheets(content)] == [None] * 5

    def test_the_marker_cell_is_empty_in_the_template_on_every_sheet_it_uses(self):
        template = openpyxl.load_workbook(TEMPLATE)
        for name in EXPORT_SHEETS:
            sheet = template[name]
            assert sheet["B1"].value is None, name
            assert any(str(r).startswith("B1:") for r in sheet.merged_cells.ranges), name  # a strip across the page
        # and the header below it is untouched in a draft
        workbook = openpyxl.load_workbook(io.BytesIO(export_bytes(idr=DRAFT_IDR)), read_only=True)
        assert (workbook["Gen Fr"]["G8"].value, workbook["Gen Fr"]["H17"].value) == ("HWS0023", "Genghis Khan")


# Worksheet part for each export sheet in the template package
SHEET_PARTS = {"Gen Fr": "sheet4.xml", "Gen Bk": "sheet5.xml", "Report Cont": "sheet3.xml",
               "Conc Fr": "sheet19.xml", "Conc Bk": "sheet20.xml", "Conc Mix": "sheet8.xml",
               "AC Fr": "sheet17.xml", "AC Bk": "sheet18.xml"}


# ---------------------------------------------------------------------------
# get_project_contractor_name (query)
# ---------------------------------------------------------------------------

class TestProjectContractorQuery:
    def test_returns_the_contractor_client_name(self):
        with patch("api.queries.projects.run_query", return_value=[{"client_name": "Benny Bowers Contracting Co."}]) as mock:
            assert get_project_contractor_name("HWS0023") == "Benny Bowers Contracting Co."
        sql, params = mock.call_args.args
        assert "FROM icid.project_clients pc" in sql and "JOIN icid.clients c" in sql
        assert "pc.client_role = 'Contractor'" in sql
        assert params == ("HWS0023",)

    def test_returns_none_without_a_contractor(self):
        with patch("api.queries.projects.run_query", return_value=[]):
            assert get_project_contractor_name("HWS0023") is None


# ---------------------------------------------------------------------------
# GET /v1/idrs/{idr_id}/export
# ---------------------------------------------------------------------------

SIGNED_EXPORT_URL = "https://example.supabase.co/storage/v1/object/sign/idr-exports/x?token=t"
EXPORT_FILENAME = f"IDR_{IDR_ID}_2026-09-30.xlsx"


@contextmanager
def stored_export(upload_error: Optional[Exception] = None, sign_error: Optional[Exception] = None):
    """
    Patch the exports bucket: uploads succeed (or raise upload_error) and signing returns SIGNED_EXPORT_URL (or raises).
    Takes the errors to raise, if any.
    Yields (the upload mock, the signing mock).
    """
    with (
        patch.object(export, "upload_file", side_effect=upload_error) as upload,
        patch.object(export, "create_signed_url", return_value=SIGNED_EXPORT_URL, side_effect=sign_error) as sign,
    ):
        yield upload, sign


class TestExportEndpoint:
    url = f"/v1/idrs/{IDR_ID}/export"

    def test_returns_a_download_url_and_the_file_name(self, admin_client):
        with patched_export(), stored_export():
            response = admin_client.get(self.url)
        assert response.status_code == 200
        assert response.json() == {"download_url": SIGNED_EXPORT_URL, "filename": EXPORT_FILENAME}

    def test_the_xlsx_is_uploaded_to_the_exports_bucket(self, admin_client):
        with patched_export(), stored_export() as (upload, sign):
            admin_client.get(self.url)
        bucket, path, content, content_type = upload.call_args.args
        assert (bucket, content_type) == ("idr-exports", export_media_type())
        assert re.fullmatch(rf"{IDR_ID}/\d{{8}}_\d{{6}}_{re.escape(EXPORT_FILENAME)}", path)
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Contract Info"]["C2"].value == "HWS0023"
        # The URL is signed for that object, for 10 minutes, downloading under the export's name
        assert sign.call_args.args == (path, 600, EXPORT_FILENAME)
        assert sign.call_args.kwargs == {"bucket": "idr-exports"}

    def test_the_path_is_stamped_with_the_export_time(self):
        at = datetime(2026, 10, 2, 14, 5, 9)
        assert export.export_storage_path(IDR_ID, "IDR.xlsx", at) == f"{IDR_ID}/20261002_140509_IDR.xlsx"
        with patched_export(), stored_export() as (upload, _):
            export.publish_idr_export(IDR_ID, now=at)
        assert upload.call_args.args[1] == f"{IDR_ID}/20261002_140509_{EXPORT_FILENAME}"

    def test_a_draft_is_exported_too(self, admin_client):
        with patched_export(idr=DRAFT_IDR), stored_export() as (upload, _):
            assert admin_client.get(self.url).status_code == 200
        content = upload.call_args.args[2]
        assert openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Gen Fr"]["B1"].value == DRAFT_MARKER

    def test_unknown_idr_is_404_and_nothing_is_stored(self, admin_client):
        with patched_export(idr=None), stored_export() as (upload, _):
            assert admin_client.get(self.url).status_code == 404
        upload.assert_not_called()

    def test_missing_project_is_500(self, admin_client):
        with patched_export(project=None), stored_export():
            assert admin_client.get(self.url).status_code == 500

    def test_storage_failing_is_502_with_a_readable_message(self, admin_client):
        for errors in ({"upload_error": RuntimeError("bucket not found")},
                       {"sign_error": RuntimeError("no signed URL")}):
            with patched_export(), stored_export(**errors):
                response = admin_client.get(self.url)
            assert response.status_code == 502, errors
            assert response.json() == {"detail": export.STORAGE_UNAVAILABLE}


def export_media_type() -> str:
    """
    The .xlsx media type the export endpoint sends.
    Takes nothing.
    Returns the MIME type string.
    """
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# The inspector's signature and its date, on every page with a signature line
# ---------------------------------------------------------------------------

SIGNATURE_PATH = f"idrs/{IDR_ID}/inspector_0123456789abcdef0123456789abcdef.png"
# 02:30 UTC on Oct 1 is 10:30 pm on Sep 30 in New York, where the forms are from
SIGNED_IDR = {**SUBMITTED_IDR, "inspector_signature_path": SIGNATURE_PATH,
              "inspector_signed_at": datetime(2026, 10, 1, 2, 30, tzinfo=timezone.utc)}
SIGNED_ON = "9/30/26"
# Each signed page: its signature line's first cell, its Date cell, and the label cells under them
SIGNATURE_LINES = {
    "Gen Bk": ("C59", "AE59", "C60", "AE60"), "Conc Bk": ("C59", "AE59", "C60", "AE60"),
    "AC Bk": ("C55", "AE55", "C56", "AE56"), "Conc Mix": ("C60", "AK60", "C61", "AK61"),
    "Report Cont": ("C49", "AF49", "C50", "AF50"), "Sketch Cont": ("C61", "AF61", "C62", "AF62"),
}


def signature_png(size: tuple[int, int] = (600, 60)) -> bytes:
    """
    Make a signature file in memory: a transparent PNG with a stroke across it.
    Takes its size in pixels.
    Returns the file's bytes.
    """
    from PIL import Image, ImageDraw
    image = Image.new("RGBA", size, (0, 0, 0, 0))
    ImageDraw.Draw(image).line([(2, size[1] - 2), (size[0] - 2, 2)], fill=(0, 0, 80, 255), width=3)
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def stamped(content: bytes, sheet: str) -> list[dict]:
    """
    Read the pictures this export added to a sheet, whether or not the sheet has a drawing at all.
    Takes the .xlsx bytes and the sheet name.
    Returns what pictures() returns, or an empty list for a sheet without a drawing.
    """
    try:
        return pictures(content, sheet)
    except (AttributeError, KeyError):
        return []


def full_signed_export(**overrides) -> bytes:
    """
    Export an IDR with every kind of signed page: a General that runs onto Report Cont, an AC, an SWCB with a
    Conc Mix addendum, and a photo.
    Takes export_bytes overrides (the IDR is SIGNED_IDR and the signature a wide PNG unless given).
    Returns the .xlsx bytes.
    """
    photo = attachment(1, SWCB_1)
    setup = {"idr": {**SIGNED_IDR, "total_pages": 4}, "general": CASCADE, "signature": signature_png(),
             "reports": [GENERAL_ROW, ac_row(page_number=2), swcb_row(1, 3), conc_mix_row(1, SWCB_1, 4)],
             "attachments": [photo], "files": {photo["storage_path"]: image_bytes("JPEG")}}
    return export_bytes(**{**setup, **overrides})


SIGNED_PAGES = ["Gen Bk", "Report Cont", "AC Bk", "Conc Bk", "Conc Mix", "Attachments 1"]
UNSIGNED_PAGES = ["Gen Fr", "AC Fr", "Conc Fr"]


class TestSignatureExport:
    def test_an_idr_without_a_signature_exports_with_blank_signature_lines(self):
        with patched_export(reports=[GENERAL_ROW, swcb_row(1, 2)]) as mocks:
            content = generate_idr_export(IDR_ID).content
        mocks["signature"].assert_not_called()
        book = openpyxl.load_workbook(io.BytesIO(content))
        for sheet in ("Gen Bk", "Conc Bk"):
            line, date_cell, label, date_label = SIGNATURE_LINES[sheet]
            assert (book[sheet][line].value, book[sheet][date_cell].value) == (None, None)
            assert (book[sheet][label].value, book[sheet][date_label].value) == ("Inspector's Signature", "Date")
            assert stamped(content, sheet) == []

    def test_a_signed_idr_is_stamped_on_every_page_with_a_signature_line(self):
        content = full_signed_export()
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Report Cont", "AC Fr", "AC Bk", "Conc Fr", "Conc Bk",
                                           "Attachments 1", "Conc Mix"]
        for sheet in SIGNED_PAGES:
            # an attachment page has its photo too: the signature is the PNG
            signatures = [p for p in stamped(content, sheet) if p["format"] == "PNG"]
            assert len(signatures) == 1, sheet
        for sheet in UNSIGNED_PAGES:
            assert stamped(content, sheet) == [], sheet

    def test_the_date_cell_shows_the_day_it_was_signed_in_new_york(self):
        book = openpyxl.load_workbook(io.BytesIO(full_signed_export()))
        dates = {"Gen Bk": "AE59", "Conc Bk": "AE59", "AC Bk": "AE55", "Conc Mix": "AK60", "Report Cont": "AF49",
                 "Attachments 1": "AF61"}
        assert {sheet: book[sheet][cell].value for sheet, cell in dates.items()} == dict.fromkeys(dates, SIGNED_ON)
        # the labels under the line are untouched, and nothing is written in the signature cell itself
        assert (book["Gen Bk"]["C60"].value, book["Gen Bk"]["AE60"].value) == ("Inspector's Signature", "Date")
        assert book["Gen Bk"]["C59"].value is None and book["Gen Bk"]["S59"].value is None  # the RE's line stays blank

    def test_the_signature_is_fetched_once_from_the_signatures_bucket(self):
        photo = attachment(1, SWCB_1)
        with patched_export(idr=SIGNED_IDR, signature=signature_png(), reports=[GENERAL_ROW, swcb_row(1, 2)],
                            attachments=[photo], files={photo["storage_path"]: image_bytes("JPEG")}) as mocks:
            generate_idr_export(IDR_ID)
        mocks["signature"].assert_called_once_with(SIGNATURE_PATH, "signatures")

    def test_a_wide_signature_fills_the_boxs_width_centred_down_it(self):
        content = export_bytes(idr=SIGNED_IDR, signature=signature_png((600, 60)))
        # 600 x 60 in the 209 x 34 px box C58:M59 -> 209 x 21, 6 px down from the top of row 58
        assert stamped(content, "Gen Bk") == [{"col": 2, "col_off": 0, "row": 57, "row_off": 6, "width": 209,
                                               "height": 21, "format": "PNG", "size": (600, 60)}]

    def test_a_tall_signature_fills_the_boxs_height_centred_across_it(self):
        content = export_bytes(idr=SIGNED_IDR, signature=signature_png((60, 120)))
        # 60 x 120 -> 17 x 34, 96 px in from C: five 19 px columns and 1 px, so in column H
        assert stamped(content, "Gen Bk") == [{"col": 7, "col_off": 1, "row": 57, "row_off": 0, "width": 17,
                                               "height": 34, "format": "PNG", "size": (60, 120)}]

    def test_a_small_signature_is_enlarged_to_the_box(self):
        placed = stamped(export_bytes(idr=SIGNED_IDR, signature=signature_png((61, 10))), "Gen Bk")[0]
        assert (placed["width"], placed["height"]) == (207, 34) and placed["col_off"] == 1

    def test_a_large_signature_is_shrunk_before_it_is_stored(self):
        placed = stamped(export_bytes(idr=SIGNED_IDR, signature=signature_png((2400, 800))), "Gen Bk")[0]
        assert placed["size"] == (408, 136)  # within 836 x 136, proportions kept
        assert (placed["width"], placed["height"]) == (102, 34)

    def test_each_sheet_is_stamped_in_its_own_box(self):
        content = full_signed_export()
        boxes = {sheet: next(p for p in stamped(content, sheet) if p["format"] == "PNG") for sheet in SIGNED_PAGES}
        # the same 209 x 21 image, 6 px down from the top of the row above each sheet's own signature line
        for sheet, row in (("Gen Bk", 57), ("Conc Bk", 57), ("AC Bk", 53), ("Report Cont", 47), ("Attachments 1", 59)):
            assert {k: boxes[sheet][k] for k in ("col", "col_off", "row", "row_off", "width", "height")} == {
                "col": 2, "col_off": 0, "row": row, "row_off": 6, "width": 209, "height": 21}, sheet
        # Conc Mix's box is C59:P60, 224 x 32 px: 224 x 22, 5 px down
        assert {k: boxes["Conc Mix"][k] for k in ("col", "col_off", "row", "row_off", "width", "height")} == {
            "col": 2, "col_off": 0, "row": 58, "row_off": 5, "width": 224, "height": 22}

    def test_copies_of_a_page_are_signed_like_the_original(self):
        photos = [attachment(n, SWCB_1) for n in (1, 2, 3)]
        content = export_bytes(idr={**SIGNED_IDR, "total_pages": 5}, signature=signature_png(),
                               reports=[GENERAL_ROW, swcb_row(1, 2), conc_mix_row(1, SWCB_1, 3, trucks=12),
                                        swcb_row(2, 4)],
                               attachments=photos, files={p["storage_path"]: image_bytes("JPEG") for p in photos})
        copies = ["Conc Bk", "Conc Bk 2", "Conc Mix", "Conc Mix 2", "Attachments 1", "Attachments 2", "Attachments 3"]
        assert set(copies) <= set(visible_sheets(content))
        for sheet in copies:
            assert len([p for p in stamped(content, sheet) if p["format"] == "PNG"]) == 1, sheet
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [book[s]["AE59"].value for s in ("Conc Bk", "Conc Bk 2")] == [SIGNED_ON, SIGNED_ON]
        assert [book[f"Attachments {n}"]["AF61"].value for n in (1, 2, 3)] == [SIGNED_ON] * 3
        assert stamped(content, "Conc Fr 2") == []

    def test_a_pdfs_page_and_the_closing_page_are_signed_too(self, monkeypatch):
        monkeypatch.setattr(export_attachments, "MAX_PHOTOS", 1)
        rows = [attachment(1, SWCB_1, PDF_TYPE), attachment(2, SWCB_1), attachment(3, SWCB_1)]
        content = export_bytes(idr=SIGNED_IDR, signature=signature_png(), reports=[swcb_row(1, 2)], attachments=rows,
                               files={r["storage_path"]: image_bytes("JPEG") for r in rows[1:]})
        assert visible_sheets(content)[-3:] == ["Attachments 1", "Attachments 2", "Attachments 3"]
        assert len(stamped(content, "Attachments 1")) == 1  # the PDF's page: the signature is its only picture
        assert len([p for p in stamped(content, "Attachments 3") if p["format"] == "PNG"]) == 1  # the closing count

    def test_a_signature_that_cant_be_fetched_leaves_the_lines_blank_and_is_logged(self, caplog):
        with patched_export(idr=SIGNED_IDR, signature=TimeoutError("read timed out"),
                            reports=[GENERAL_ROW, swcb_row(1, 2)]) as mocks:
            content = generate_idr_export(IDR_ID).content
        mocks["signature"].assert_called_once()
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]  # the export still completes
        book = openpyxl.load_workbook(io.BytesIO(content))
        for sheet in ("Gen Bk", "Conc Bk"):
            assert stamped(content, sheet) == [] and book[sheet]["AE59"].value is None
        assert SIGNATURE_PATH in caplog.text and "unavailable for the export" in caplog.text

    def test_a_missing_signature_file_is_handled_the_same_way(self, caplog):
        content = export_bytes(idr=SIGNED_IDR, signature=None)
        assert stamped(content, "Gen Bk") == []
        assert openpyxl.load_workbook(io.BytesIO(content))["Gen Bk"]["AE59"].value is None
        assert SIGNATURE_PATH in caplog.text

    def test_a_signature_file_that_isnt_an_image_is_handled_the_same_way(self, caplog):
        content = export_bytes(idr=SIGNED_IDR, signature=b"not a png at all")
        assert stamped(content, "Gen Bk") == []
        assert openpyxl.load_workbook(io.BytesIO(content))["Gen Bk"]["AE59"].value is None
        assert "unavailable for the export" in caplog.text

    def test_a_draft_is_never_signed_even_with_a_path_on_its_row(self):
        draft = {**SIGNED_IDR, "status": "draft", "submitted_at": None, "total_pages": None}
        with patched_export(idr=draft, signature=signature_png(), reports=[GENERAL_ROW, swcb_row(1, None)]) as mocks:
            content = generate_idr_export(IDR_ID).content
        mocks["signature"].assert_not_called()  # not even fetched
        book = openpyxl.load_workbook(io.BytesIO(content))
        for sheet in ("Gen Bk", "Conc Bk"):
            assert stamped(content, sheet) == [] and book[sheet]["AE59"].value is None
            assert book[sheet]["B1"].value == "DRAFT - Not for Submission"

    def test_a_signed_page_carries_no_draft_marker(self):
        book = openpyxl.load_workbook(io.BytesIO(export_bytes(idr=SIGNED_IDR, signature=signature_png())))
        assert book["Gen Bk"]["B1"].value != "DRAFT - Not for Submission"

    def test_gen_bk_which_had_no_drawing_gets_a_valid_one(self):
        content = export_bytes(idr=SIGNED_IDR, signature=signature_png())
        package, drawing, anchors = added_anchors(content, "Gen Bk")
        assert len(anchors) == 1 and 'descr="Inspector\'s signature"' in anchors[0]
        assert f'<Override PartName="/{drawing}" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>' \
            in package.read("[Content_Types].xml").decode()
        assert package.testzip() is None
        sheet_xml = package.read("xl/worksheets/sheet5.xml").decode()
        assert sheet_xml.count("<drawing ") == 1
        assert sheet_xml.index("<pageSetup") < sheet_xml.index("<drawing ") < sheet_xml.index("</worksheet>")
        assert len(openpyxl.load_workbook(io.BytesIO(content))["Gen Bk"]._images) == 1

    def test_the_tab_order_and_page_numbers_are_what_they_were(self):
        signed, unsigned = full_signed_export(), full_signed_export(idr={**SUBMITTED_IDR, "total_pages": 4})
        assert tab_order(signed) == tab_order(unsigned)
        assert print_order_page_numbers(signed) == print_order_page_numbers(unsigned)


class TestSignatureLayouts:
    layouts = {"Gen Bk": export.SIGNATURE_LAYOUTS["Gen Bk"], "Conc Bk": export.SIGNATURE_LAYOUTS["Conc Bk"],
               "AC Bk": export.SIGNATURE_LAYOUTS["AC Bk"], "Conc Mix": export.SIGNATURE_LAYOUTS["Conc Mix"],
               "Report Cont": export.SIGNATURE_LAYOUTS["Report Cont"],
               "Sketch Cont": export.SIGNATURE_LAYOUTS["Attachments"]}

    def test_every_template_sheet_with_a_signature_line_has_a_layout(self):
        assert set(self.layouts) == set(SIGNATURE_LINES)
        assert set(export.SIGNATURE_LAYOUTS) == {"Gen Bk", "Conc Bk", "AC Bk", "Conc Mix", "Report Cont", "Attachments"}

    @pytest.mark.parametrize("sheet", sorted(SIGNATURE_LINES))
    def test_the_layout_matches_the_template(self, sheet):
        from openpyxl.utils import range_boundaries
        layout = self.layouts[sheet]
        template = openpyxl.load_workbook(export.TEMPLATE_PATH)[sheet]
        line, date_cell, label, date_label = SIGNATURE_LINES[sheet]
        merged = {str(m).split(":")[0]: str(m) for m in template.merged_cells.ranges}
        left, top, right, bottom = range_boundaries(layout.signature_cells)
        # the box is the signature line's merged cell plus the one row above it
        assert merged[line] == f"{line}:{openpyxl.utils.get_column_letter(right)}{bottom}"
        assert (template[label].value, template[date_label].value) == ("Inspector's Signature", "Date")
        assert openpyxl.utils.cell.coordinate_from_string(line) == (openpyxl.utils.get_column_letter(left), bottom)
        assert top == bottom - 1 and len(layout.row_px) == 2
        assert layout.date_cell == date_cell and date_cell in merged
        # nothing is in the row above the line for the image to cover
        assert all(template.cell(top, column).value is None for column in range(left, right + 1))
        # and the measured size is the template's own: Excel shows a column 7 px a character, a row 4 px to 3 pt
        widths = {round(template.column_dimensions[key].width * 7) for key, d in template.column_dimensions.items()
                  if d.min <= left and right <= d.max}
        assert widths == {layout.column_px}
        heights = tuple(round((template.row_dimensions[row].height or 12.75) / 0.75) for row in (top, bottom))
        assert heights == layout.row_px
        assert layout.signature_cx_emu == (right - left + 1) * layout.column_px * 9525
        assert layout.signature_cy_emu == sum(layout.row_px) * 9525

    def test_a_copys_name_finds_its_originals_layout(self):
        assert export._signature_layout("Conc Bk 2") is export.SIGNATURE_LAYOUTS["Conc Bk"]
        assert export._signature_layout("Attachments 17") is export.SIGNATURE_LAYOUTS["Attachments"]
        assert export._signature_layout("Conc Mix 3") is export.SIGNATURE_LAYOUTS["Conc Mix"]
        assert export._signature_layout("Report Cont") is export.SIGNATURE_LAYOUTS["Report Cont"]
        for front in ("Gen Fr", "Gen Fr 2", "Conc Fr", "Conc Fr 2", "AC Fr", "Contract Info"):
            assert export._signature_layout(front) is None


class TestSignatureHelpers:
    def test_the_signed_date_is_the_day_in_new_york(self):
        from api.services.export_common import signed_date
        assert signed_date(datetime(2026, 10, 1, 2, 30, tzinfo=timezone.utc)) == date(2026, 9, 30)  # 10:30 pm EDT
        assert signed_date(datetime(2026, 10, 1, 16, 0, tzinfo=timezone.utc)) == date(2026, 10, 1)
        assert signed_date(datetime(2026, 1, 15, 4, 59, tzinfo=timezone.utc)) == date(2026, 1, 14)  # 11:59 pm EST
        assert signed_date(datetime(2026, 1, 15, 5, 0, tzinfo=timezone.utc)) == date(2026, 1, 15)
        assert signed_date(datetime(2026, 10, 1, 2, 30)) == date(2026, 9, 30)  # no zone: taken as UTC

    def test_stamping_no_signature_changes_nothing(self):
        from api.services.export_common import stamp_signature
        workbook = WorkbookTemplate(export.TEMPLATE_PATH)
        before = workbook.to_bytes()
        stamp_signature(workbook, "Gen Bk", export.SIGNATURE_LAYOUTS["Gen Bk"], None,
                        datetime(2026, 10, 1, tzinfo=timezone.utc))
        after = openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))
        assert after["Gen Bk"]["AE59"].value is None  # no date without a signature
        assert len(workbook.to_bytes()) == len(before)

    def test_a_signature_with_no_signing_time_leaves_the_date_blank(self):
        from api.services.export_common import prepare_signature, stamp_signature
        workbook = WorkbookTemplate(export.TEMPLATE_PATH)
        stamp_signature(workbook, "Conc Bk", export.SIGNATURE_LAYOUTS["Conc Bk"],
                        prepare_signature(signature_png()), None)
        content = workbook.to_bytes()
        assert len(stamped(content, "Conc Bk")) == 1
        assert openpyxl.load_workbook(io.BytesIO(content))["Conc Bk"]["AE59"].value is None

    def test_prepare_signature_keeps_png_and_transparency(self):
        from PIL import Image
        from api.services.export_common import prepare_signature
        prepared = prepare_signature(signature_png((300, 100)))
        image = Image.open(io.BytesIO(prepared.data))
        assert (image.format, image.mode, image.size) == ("PNG", "RGBA", (300, 100))
        assert (prepared.width, prepared.height) == (300, 100)
        assert image.getpixel((150, 5))[3] == 0  # still see-through off the stroke
        with pytest.raises(Exception):
            prepare_signature(b"not an image")

    def test_a_sheets_existing_drawing_is_reused_not_replaced(self):
        workbook = WorkbookTemplate(export.TEMPLATE_PATH)
        existing = workbook._drawing_part("Conc Bk")
        assert workbook._ensure_drawing("Conc Bk") == existing
        created = workbook._ensure_drawing("Gen Bk")
        assert workbook._ensure_drawing("Gen Bk") == created and created != existing  # made once, then found
