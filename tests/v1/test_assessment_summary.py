import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal
from unittest.mock import patch
from uuid import UUID

import openpyxl
import pytest

from tests.conftest import ADMIN_USER_ROW, signed_in

# ---------------------------------------------------------------------------
# GET /v1/projects/{project_id}/assessment-summary.xlsx
# ---------------------------------------------------------------------------

SUMMARY = "api.queries.assessment_summary.run_query"
PROJECTS = "api.queries.projects.run_query"
URL = "/v1/projects/SE384/assessment-summary.xlsx"
SITES_URL = "/v1/projects/SE384/assessment-summary"
TAILORED_URL = "/v1/projects/SE384/assessment-summary/tailored.xlsx"
PROJECT_ROW = {"project_id": "SE384", "project_name": "Sidewalks Queens", "project_description": None,
               "registration_code": "2024123456", "borough": "Queens", "status": None,
               "cost_per_sf_4in": Decimal("11.54"), "cost_per_sf_7in": Decimal("12.64")}
INSPECTOR_ROW = {**ADMIN_USER_ROW, "uuid": UUID("b0000000-0000-4000-8000-000000000002"), "role": None}
FLAG = {"idr_id": UUID("1d000000-0000-4000-8000-000000000001"),
        "re_signed_at": datetime(2026, 10, 15, 18, 0, tzinfo=timezone.utc),
        "property_id": UUID("20000000-0000-4000-8000-000000000001"), "property_kind": "sidewalk",
        "address": "138-22 227 STREET", "intersection_street_1": None, "intersection_street_2": None,
        "quadrant": None, "facing_street": None, "date_constructed": None, "surveyed_thickness": "4in",
        "surveyed_is_assessable": True, "computed_quantity": Decimal("160.69"), "quantity": Decimal("113.19")}


def projects_for(on_project: bool = True):
    """
    Answer the project statements: the project itself, and whether the user is on it.
    Takes whether they are.
    Returns the side effect.
    """
    def answer(sql, params=None):
        if "FROM icid.project_users" in sql:
            return [{"?column?": 1}] if on_project else []
        return [PROJECT_ROW]
    return answer


def book(response):
    """
    Open a response's workbook.
    Takes the response.
    Returns the openpyxl workbook.
    """
    return openpyxl.load_workbook(io.BytesIO(response.content))


class TestAssessmentSummary:
    def test_it_is_an_xlsx_named_for_the_project_and_today(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]) as run:
            response = admin_client.get(URL)
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        disposition = response.headers["content-disposition"]
        assert disposition.startswith('attachment; filename="SE384_assessment_summary_') and disposition.endswith('.xlsx"')
        assert book(response).sheetnames == ["Oct 2026", "Project Total"]
        sql, params = run.call_args.args
        assert params == ("SE384",)
        assert "i.status = 'approved' AND i.deleted_at IS NULL" in sql
        assert "pg1.page_number = 1" in sql

    def test_a_range_keeps_its_months(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            assert book(admin_client.get(URL, params={"from": "2026-10", "to": "2026-10"})).sheetnames == ["Oct 2026", "Project Total"]
            later = book(admin_client.get(URL, params={"from": "2026-11"}))
            earlier = book(admin_client.get(URL, params={"to": "2026-09"}))
        # Every month of the range gets a tab (ML1a): a month with nothing is an empty one
        assert later.sheetnames == ["Nov 2026", "Project Total"] and earlier.sheetnames == ["Sep 2026", "Project Total"]
        assert later["Nov 2026"]["A4"].value == "Totals"
        assert later["Project Total"]["A4"].value == "No approved IDRs in the selected range."

    def test_a_project_with_nothing_approved_is_a_no_activity_workbook(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[]):
            response = admin_client.get(URL)
        assert response.status_code == 200
        assert book(response)["No activity"]["A4"].value == "No approved IDRs in this project yet."

    def test_from_after_to_is_400(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY) as run:
            response = admin_client.get(URL, params={"from": "2026-11", "to": "2026-10"})
        assert response.status_code == 400 and response.json() == {"detail": "from after to"}
        run.assert_not_called()

    @pytest.mark.parametrize("name, value", [("from", "10-2026"), ("to", "2026-13")])
    def test_a_month_that_isnt_yyyy_mm_is_400(self, admin_client, name, value):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY) as run:
            response = admin_client.get(URL, params={name: value})
        assert response.status_code == 400 and response.json() == {"detail": f"{name} must be YYYY-MM"}
        run.assert_not_called()

    def test_a_failed_read_is_500(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=None):
            assert admin_client.get(URL).status_code == 500

    def test_an_unknown_project_is_404(self, admin_client):
        with patch(PROJECTS, return_value=[]), patch(SUMMARY) as run:
            assert admin_client.get(URL).status_code == 404
        run.assert_not_called()

    def test_a_user_on_the_project_downloads_it(self):
        with signed_in(INSPECTOR_ROW) as client, patch(PROJECTS, side_effect=projects_for()), \
                patch(SUMMARY, return_value=[FLAG]):
            assert client.get(URL).status_code == 200

    def test_a_user_who_isnt_on_the_project_is_forbidden(self):
        with signed_in(INSPECTOR_ROW) as client, patch(PROJECTS, side_effect=projects_for(on_project=False)), \
                patch(SUMMARY) as run:
            assert client.get(URL).status_code == 403
        run.assert_not_called()

    def test_an_admin_downloads_any_projects(self, admin_client):
        # admin_client is on no project: project_member lets an admin through
        with patch(PROJECTS, side_effect=projects_for(on_project=False)), patch(SUMMARY, return_value=[FLAG]):
            assert admin_client.get(URL).status_code == 200

    def test_a_demo_user_downloads_their_own_projects(self, demo_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[]):
            assert demo_client.get("/v1/projects/DEMO01/assessment-summary.xlsx").status_code == 200

    def test_a_demo_user_gets_404_for_another_project(self, demo_client):
        with patch(PROJECTS, side_effect=projects_for(on_project=False)), patch(SUMMARY) as run:
            assert demo_client.get(URL).status_code == 404
        run.assert_not_called()


# ---------------------------------------------------------------------------
# GET /v1/projects/{project_id}/assessment-summary (the page's sites)
# ---------------------------------------------------------------------------

ELM_FLAG = {**FLAG, "property_id": UUID("20000000-0000-4000-8000-000000000002"), "address": "350 Main St, Brooklyn",
            "community_board": "BK10", "block": "1234", "lot": "51", "violation_no": "V-77",
            "date_constructed": date(2026, 11, 3)}


class TestAssessmentSites:
    def test_a_row_per_site_with_its_figures_and_the_totals(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            response = admin_client.get(SITES_URL)
        assert response.status_code == 200
        data = response.json()["data"]
        assert data["range"] == {"from": None, "to": None} and data["cost_per_sf_4in"] == 11.54
        assert [site["display_address"] for site in data["sites"]] == ["138-22 227 STREET", "350 Main St, Brooklyn"]
        elm = data["sites"][1]
        assert (elm["community_board"], elm["block"], elm["lot"], elm["violation_no"], elm["kind"]) == (
            "BK10", "1234", "51", "V-77", "sidewalk")
        assert elm["four_assessable"] == 113.19 and elm["four_amount"] == 1306.21
        assert data["totals"]["four_assessable"] == 226.38 and data["totals"]["seven_amount"] == 0.0

    def test_the_range_is_of_days(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            data = admin_client.get(SITES_URL, params={"from": "2026-11-01", "to": "2026-11-30"}).json()["data"]
        assert data["range"] == {"from": "2026-11-01", "to": "2026-11-30"}
        assert [site["display_address"] for site in data["sites"]] == ["350 Main St, Brooklyn"]

    @pytest.mark.parametrize("params, detail", [
        ({"from": "2026-11"}, "from must be a date, YYYY-MM-DD"),
        ({"from": "2026-11-02", "to": "2026-11-01"}, "from cannot be after to")])
    def test_a_bad_range_is_400(self, admin_client, params, detail):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY) as run:
            response = admin_client.get(SITES_URL, params=params)
        assert response.status_code == 400 and response.json() == {"detail": detail}
        run.assert_not_called()

    def test_a_user_who_isnt_on_the_project_is_forbidden(self):
        with signed_in(INSPECTOR_ROW) as inspector, patch(PROJECTS, side_effect=projects_for(on_project=False)), \
                patch(SUMMARY, return_value=[FLAG]):
            assert inspector.get(SITES_URL).status_code == 403

    def test_a_failed_read_is_500(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=None):
            assert admin_client.get(SITES_URL).status_code == 500


class TestTailoredDownload:
    def test_it_is_one_tab_of_the_sites_asked_for(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            response = admin_client.get(TAILORED_URL, params={
                "from": "2026-10-01", "to": "2026-11-30", "property_id": str(ELM_FLAG["property_id"])})
        assert response.status_code == 200
        assert 'filename="SE384_assessment_summary_tailored_' in response.headers["content-disposition"]
        workbook = book(response)
        assert workbook.sheetnames == ["2026-10-01 to 2026-11-30"]
        sheet = workbook["2026-10-01 to 2026-11-30"]
        assert sheet["A4"].value == "350 Main St, Brooklyn" and sheet["B4"].value == "BK10" and sheet["A5"].value == "Totals"

    def test_without_property_ids_it_holds_every_site(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            sheet = book(admin_client.get(TAILORED_URL))["All dates"]
        assert sheet["A4"].value == "138-22 227 STREET" and sheet["A5"].value == "350 Main St, Brooklyn"

    def test_a_property_id_that_isnt_a_uuid_is_422(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            assert admin_client.get(TAILORED_URL, params={"property_id": "nope"}).status_code == 422

    def test_a_bad_range_is_400(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY) as run:
            assert admin_client.get(TAILORED_URL, params={"to": "31-10-2026"}).status_code == 400
        run.assert_not_called()


# ---------------------------------------------------------------------------
# The Assessment Summary polish (ML1a, ML1b, ML1c, ML4b)
# ---------------------------------------------------------------------------

TASK_ORDERS = "api.queries.task_orders.run_query"
TASK_ORDER_ID = UUID("70000000-0000-4000-8000-000000000001")
CSV_URL = "/v1/projects/SE384/assessment-summary.csv"
TAILORED_CSV_URL = "/v1/projects/SE384/assessment-summary/tailored.csv"
RE_ROLES = [{"user_uuid": UUID("b0000000-0000-4000-8000-000000000009"), "email": "jd@x.com", "first_name": "Jane",
             "last_name": "Doe", "role": "re"},
            {"user_uuid": UUID("b0000000-0000-4000-8000-000000000009"), "email": "jd@x.com", "first_name": "Jane",
             "last_name": "Doe", "role": "oe"},
            {"user_uuid": UUID("b0000000-0000-4000-8000-000000000008"), "email": "sam@x.com", "first_name": "Sam",
             "last_name": None, "role": "re"},
            {"user_uuid": UUID("b0000000-0000-4000-8000-000000000007"), "email": "i@x.com", "first_name": "Ina",
             "last_name": "Spector", "role": "inspector"}]


def projects_with_roles(roles):
    """
    Answer the project statements as projects_for does, the project's roles as given.
    Takes the roles list_project_roles returns.
    Returns the side effect.
    """
    def answer(sql, params=None):
        if "pu.role" in sql and "pu.assigned_at" in sql:
            return roles
        if "FROM icid.project_users" in sql:
            return [{"?column?": 1}]
        return [PROJECT_ROW]
    return answer


def task_order(properties=(), **changes):
    """A task order of SE384 as get_task_order returns it, its bundled properties' ids given."""
    return [{"id": TASK_ORDER_ID, "project_id": "SE384", "task_order_number": "TO-1", "deleted_at": None,
             "properties": [{"id": str(found)} for found in properties], **changes}]


def column_a(sheet):
    """Column A of a sheet, top to bottom, without its empty cells."""
    return [cell.value for cell in sheet["A"] if cell.value is not None]


def csv_lines(response):
    """A CSV response's lines, its byte-order mark dropped."""
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))


class TestEveryMonth:
    @pytest.mark.parametrize("start, end, names", [
        ("2026-10", "2026-12", ["Oct 2026", "Nov 2026", "Dec 2026"]),
        ("2026-01", "2026-12", ["Jan 2026", "Feb 2026", "Mar 2026", "Apr 2026", "May 2026", "Jun 2026", "Jul 2026",
                                "Aug 2026", "Sep 2026", "Oct 2026", "Nov 2026", "Dec 2026"]),
        ("2025-12", "2026-03", ["Dec 2025", "Jan 2026", "Feb 2026", "Mar 2026"]),
    ], ids=["oct-dec", "a-year", "across-years"])
    def test_a_range_has_a_tab_for_every_month(self, admin_client, start, end, names):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            workbook = book(admin_client.get(URL, params={"from": start, "to": end}))
        assert workbook.sheetnames == names + ["Project Total"]

    def test_an_empty_month_is_its_titles_and_a_totals_row_of_0(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            sheet = book(admin_client.get(URL, params={"from": "2026-10", "to": "2026-11"}))["Nov 2026"]
        assert sheet["A2"].value == "November 2026" and sheet["A3"].value == "Property Address"
        assert [cell.value for cell in sheet[4]][:12] == ["Totals", None, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]

    def test_without_a_range_it_runs_from_the_first_month_with_rows_to_the_last(self, admin_client):
        december = {**ELM_FLAG, "date_constructed": date(2026, 12, 2)}
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, december]):
            workbook = book(admin_client.get(URL))
        assert workbook.sheetnames == ["Oct 2026", "Nov 2026", "Dec 2026", "Project Total"]
        assert workbook["Nov 2026"]["A4"].value == "Totals"


class TestTaskOrderFilter:
    def test_it_keeps_the_task_orders_properties(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]), \
                patch(TASK_ORDERS, return_value=task_order([ELM_FLAG["property_id"]])):
            sites = admin_client.get(SITES_URL, params={"task_order_id": str(TASK_ORDER_ID)}).json()["data"]["sites"]
            tailored = book(admin_client.get(TAILORED_URL, params={"task_order_id": str(TASK_ORDER_ID)}))["All dates"]
            monthly = book(admin_client.get(URL, params={"task_order_id": str(TASK_ORDER_ID)}))
        assert [site["display_address"] for site in sites] == ["350 Main St, Brooklyn"]
        assert tailored["A4"].value == "350 Main St, Brooklyn" and tailored["A5"].value == "Totals"
        assert monthly.sheetnames == ["Nov 2026", "Project Total"]

    def test_without_one_every_property_stays(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]), \
                patch(TASK_ORDERS) as orders:
            sites = admin_client.get(SITES_URL).json()["data"]["sites"]
        assert len(sites) == 2
        orders.assert_not_called()

    @pytest.mark.parametrize("found", [[], task_order(project_id="HWS0023"), task_order(deleted_at=datetime(2026, 10, 1))],
                             ids=["unknown", "another-project", "deleted"])
    def test_a_task_order_that_isnt_the_projects_is_404(self, admin_client, found):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]), \
                patch(TASK_ORDERS, return_value=found):
            for url in (SITES_URL, URL, CSV_URL, TAILORED_URL, TAILORED_CSV_URL):
                response = admin_client.get(url, params={"task_order_id": str(TASK_ORDER_ID)})
                assert response.status_code == 404 and response.json()["detail"] == "Task order not found", url

    def test_a_task_order_id_that_isnt_a_uuid_is_422(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            assert admin_client.get(URL, params={"task_order_id": "TO-1"}).status_code == 422


class TestCsv:
    def test_the_monthly_csv_is_a_row_per_month_and_site(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            response = admin_client.get(CSV_URL, params={"from": "2026-09", "to": "2026-11"})
        assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
        disposition = response.headers["content-disposition"]
        assert disposition.startswith('attachment; filename="SE384_assessment_summary_') and disposition.endswith('.csv"')
        assert response.content.startswith(b"\xef\xbb\xbf")
        header, *rows = csv_lines(response)
        assert header[0] == "Month" and header[1] == "Property Address" and header[-1] == "Remeasure Pending"
        assert len(header) == 15
        assert [(row[0], row[1]) for row in rows] == [("2026-10", "138-22 227 STREET"), ("2026-11", "350 Main St, Brooklyn")]

    def test_its_figures_are_to_the_cent_and_match_the_workbook(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG]):
            header, row = csv_lines(admin_client.get(CSV_URL))
            sheet = book(admin_client.get(URL))["Oct 2026"]
        figures = dict(zip(header, row))
        assert figures['4" Assessable (SF)'] == "113.19" and figures['4" Total (SF)'] == "160.69"
        assert figures['4" Amount ($)'] == f"{sheet['F4'].value:.2f}" == "1306.21"
        assert figures['7" Total (SF)'] == "0.00" and figures["Remeasure Pending"] == ""

    def test_the_tailored_csv_is_a_row_per_site_without_a_month(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[FLAG, ELM_FLAG]):
            response = admin_client.get(TAILORED_CSV_URL, params={"property_id": str(FLAG["property_id"])})
        assert 'filename="SE384_assessment_summary_tailored_' in response.headers["content-disposition"]
        header, *rows = csv_lines(response)
        assert header[0] == "Property Address" and len(header) == 14
        assert [row[0] for row in rows] == ["138-22 227 STREET"]

    def test_a_csv_with_nothing_is_its_header(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY, return_value=[]):
            assert len(csv_lines(admin_client.get(CSV_URL))) == 1
            assert len(csv_lines(admin_client.get(TAILORED_CSV_URL))) == 1

    def test_a_bad_range_is_400_and_a_stranger_is_forbidden(self, admin_client):
        with patch(PROJECTS, side_effect=projects_for()), patch(SUMMARY) as run:
            assert admin_client.get(CSV_URL, params={"from": "2026-12", "to": "2026-10"}).status_code == 400
            assert admin_client.get(TAILORED_CSV_URL, params={"from": "10/1/2026"}).status_code == 400
        run.assert_not_called()
        with signed_in(INSPECTOR_ROW) as client, patch(PROJECTS, side_effect=projects_for(on_project=False)):
            assert client.get(CSV_URL).status_code == 403
            assert client.get(TAILORED_CSV_URL).status_code == 403


class TestFooter:
    def test_every_tab_ends_with_who_compiled_it_the_project_engineers_and_approved_only(self, admin_client):
        with patch(PROJECTS, side_effect=projects_with_roles(RE_ROLES)), patch(SUMMARY, return_value=[FLAG]):
            workbook = book(admin_client.get(URL, params={"from": "2026-10", "to": "2026-11"}))
        compiled = f"Compiled by: {ADMIN_USER_ROW['first_name']} {ADMIN_USER_ROW['last_name']}"
        for name in workbook.sheetnames:
            footer = column_a(workbook[name])[-4:]
            assert footer[0] == compiled, name
            assert footer[1].startswith("Date: 20") and len(footer[1]) == len("Date: 2026-10-28")
            assert footer[2] == "Project Engineer: Jane Doe, Sam"
            assert footer[3] == "Approved IDRs only"

    def test_the_footer_follows_the_notes_after_a_blank_row(self, admin_client):
        with patch(PROJECTS, side_effect=projects_with_roles(RE_ROLES)), patch(SUMMARY, return_value=[FLAG]):
            sheet = book(admin_client.get(URL))["Oct 2026"]
        # Row 4 the site, 5 the totals, 7 the approval-date note, 9-12 the footer
        assert sheet["A5"].value == "Totals" and sheet["A7"].value.startswith("1 row used IDR approval date")
        assert sheet["A8"].value is None and sheet["A9"].value.startswith("Compiled by: ")
        assert sheet["A12"].value == "Approved IDRs only"

    def test_without_an_re_the_project_engineer_is_not_assigned(self, admin_client):
        with patch(PROJECTS, side_effect=projects_with_roles(RE_ROLES[3:])), patch(SUMMARY, return_value=[]):
            workbook = book(admin_client.get(URL))
            tailored = book(admin_client.get(TAILORED_URL))
        assert workbook.sheetnames == ["No activity", "Project Total"]
        for sheet in (workbook["No activity"], workbook["Project Total"], tailored["All dates"]):
            assert "Project Engineer: not assigned" in column_a(sheet)
        assert workbook["No activity"]["A6"].value.startswith("Compiled by: ")

    def test_the_csvs_have_no_footer(self, admin_client):
        with patch(PROJECTS, side_effect=projects_with_roles(RE_ROLES)), patch(SUMMARY, return_value=[FLAG]):
            text = (admin_client.get(CSV_URL).content.decode("utf-8-sig")
                    + admin_client.get(TAILORED_CSV_URL).content.decode("utf-8-sig"))
        assert "Compiled by" not in text and "Approved IDRs only" not in text
