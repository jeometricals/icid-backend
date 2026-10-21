"""
A Daily Site Patrol (DSP) report on the export: export_dsp.render on DSP Fr, and the dispatcher placing, numbering and
signing it.
"""

import io
import re
import zipfile
from typing import Optional
from uuid import UUID

import openpyxl
import pytest

from api.schemas.dsp_report import MOS_ITEMS, MPT_ITEMS, blank_dsp_report
from api.services import export, export_dsp
from api.services.export_common import DRAFT_MARKER
from api.services.xlsx_template import WorkbookTemplate
from tests.v1.test_exports import (
    APPROVED_IDR, APPROVED_ON, DRAFT_IDR, GENERAL, PROJECT, RE_USER, RE_UUID, SIGNED_IDR, SUBMITTED_IDR, attachment,
    export_bytes, image_bytes, signature_png, signatures_on, stamped, swcb_row, visible_sheets,
)

CONTRACTOR = "Benny Bowers Contracting Co."
MOS_FIRST_ROWS = [19, 22, 25, 28, 31, 34]
MPT_FIRST_ROWS = [38 + 3 * n for n in range(16)]
ANSWER_COLUMNS = {"yes": "W", "no": "Z", "na": "AC"}


def dsp_data(mos: Optional[list] = None, mpt: Optional[list] = None, **fields) -> dict:
    """
    Build a DSP report_data: the blank form with the given answers and fields.
    Takes the MOS and MPT rows' (answer, details) pairs, in the form's order (missing ones stay blank), and top-level
    fields as keyword arguments.
    Returns the report_data.
    """
    data = blank_dsp_report()
    for key, pairs in (("mosChecklist", mos or []), ("mptChecklist", mpt or [])):
        for row, (answer, details) in zip(data[key], pairs):
            row.update(answer=answer, details=details)
    return {**data, **fields}


def dsp_page(data: Optional[dict] = None, idr: dict = SUBMITTED_IDR, page_number: Optional[int] = 2) -> bytes:
    """
    Render a DSP report on a fresh template.
    Takes the report_data, the IDR row and the page number.
    Returns the workbook's bytes.
    """
    workbook = WorkbookTemplate(export.TEMPLATE_PATH)
    assert export_dsp.render(workbook, idr, PROJECT, CONTRACTOR, page_number=page_number,
                             report_data=data) == ["DSP Fr"]
    return workbook.to_bytes()


def dsp_sheet(content: bytes):
    """
    Open the DSP Fr sheet of a workbook's bytes.
    Takes the bytes.
    Returns the openpyxl sheet.
    """
    return openpyxl.load_workbook(io.BytesIO(content))["DSP Fr"]


def dsp_row(number: int = 1, page_number: Optional[int] = 2, parent: Optional[UUID] = None,
            data: Optional[dict] = None) -> dict:
    """
    Build the IDR's nth DSP report row as list_reports_for_idr returns it.
    Takes its number (for a distinct id), its page number, its parent's id (None for a main report) and report_data.
    Returns the row.
    """
    return {"report_id": UUID(int=0xD500 + number), "report_type": "DSP", "is_addendum": parent is not None,
            "parent_report_id": parent, "page_number": page_number,
            "report_data": data if data is not None else dsp_data()}


class TestDspHeader:
    def test_the_project_details_are_values_not_contract_info_formulas(self):
        sheet = dsp_sheet(dsp_page(dsp_data()))
        assert [sheet[cell].value for cell in ("G8", "P8", "I10", "F12", "F14")] == [
            "HWS0023", "2024123457", PROJECT["project_description"], "Queens", CONTRACTOR]

    def test_the_date_is_the_work_date_on_one_merged_line_without_slashes(self):
        sheet = dsp_sheet(dsp_page(dsp_data()))
        merged = {str(area) for area in sheet.merged_cells.ranges}
        assert "AK8:AQ8" in merged and "AK8:AL8" not in merged
        assert sheet["AK8"].value == "9/30/26"
        assert all(sheet[cell].border.diagonal is None or sheet[cell].border.diagonal.style is None
                   for cell in ("AM8", "AO8"))

    def test_the_day_of_week_is_shaded_from_the_work_date(self):
        sheet = dsp_sheet(dsp_page(dsp_data()))
        shaded = [cell for cell in export_dsp.DAY_OF_WEEK if sheet[cell].fill.fill_type == "solid"]
        assert shaded == ["AN10"]  # Wednesday, Sept 30 2026

    def test_ir_no_sheet_and_time_of_patrol(self):
        idr = {**SUBMITTED_IDR, "idr_number": "005", "total_pages": 4}
        sheet = dsp_sheet(dsp_page(dsp_data(timeOfPatrol="14:30"), idr=idr, page_number=3))
        assert [sheet[cell].value for cell in ("AJ12", "AJ13", "AO13", "AK15")] == ["005", 3, 4, "14:30"]

    def test_a_draft_has_no_number_sheet_or_time(self):
        sheet = dsp_sheet(dsp_page(dsp_data(), idr=DRAFT_IDR, page_number=None))
        assert [sheet[cell].value for cell in ("AJ12", "AJ13", "AO13", "AK15")] == [None, None, None, None]


class TestDspChecklists:
    def test_each_answer_puts_an_x_in_its_own_column(self):
        mos = [("yes", None), ("no", "Trench open"), ("na", None), (None, None)]
        sheet = dsp_sheet(dsp_page(dsp_data(mos=mos)))
        marks = [[sheet[f"{column}{row + 1}"].value for column in ANSWER_COLUMNS.values()]
                 for row in MOS_FIRST_ROWS[:4]]
        assert marks == [["X", None, None], [None, "X", None], [None, None, "X"], [None, None, None]]

    def test_the_traffic_checklist_takes_its_sixteen_rows(self):
        mpt = [(answer, None) for answer in ["no"] * 15 + ["yes"]]
        sheet = dsp_sheet(dsp_page(dsp_data(mpt=mpt)))
        assert [sheet[f"Z{row + 1}"].value for row in MPT_FIRST_ROWS] == ["X"] * 15 + [None]
        assert sheet[f"W{MPT_FIRST_ROWS[-1] + 1}"].value == "X"

    def test_details_print_beside_their_question_on_two_lines(self):
        sheet = dsp_sheet(dsp_page(dsp_data(mpt=[(None, None), ("no", "Advance warning sign missing on Jamaica Ave.")])))
        assert sheet["AE41"].value == "Advance warning sign missing on\nJamaica Ave."
        assert sheet["AE41"].font.sz == export_dsp.DETAILS_FONT_PT
        assert sheet["AE38"].value is None

    def test_details_too_long_for_two_lines_are_cut(self):
        long = " ".join(f"word{n}" for n in range(40))
        text = export_dsp.details_text(long)
        assert text.endswith("(continued in ICID)") and len(text.split("\n")) == 2

    def test_rows_are_found_by_item_not_position(self):
        data = dsp_data(mos=[("yes", None)])
        data["mosChecklist"] = list(reversed(data["mosChecklist"]))  # the first item now saved last
        sheet = dsp_sheet(dsp_page(data))
        assert sheet["W20"].value == "X"

    def test_an_empty_report_prints_a_blank_form(self):
        sheet = dsp_sheet(dsp_page({}))
        assert all(sheet[f"{column}{row + 1}"].value is None for row in MOS_FIRST_ROWS + MPT_FIRST_ROWS
                   for column in ANSWER_COLUMNS.values())
        assert (sheet["AI87"].value, sheet["AK87"].value) == (None, None)

    def test_the_question_text_is_the_templates(self):
        sheet = dsp_sheet(dsp_page(dsp_data()))
        assert len(MOS_ITEMS) == len(MOS_FIRST_ROWS) and len(MPT_ITEMS) == len(MPT_FIRST_ROWS)
        assert sheet["B19"].value == "Has waste, rubbish and debris been removed from site ?"
        assert sheet["B83"].value == "Are adequate flaggers / temporary signs utilized when required ?"


class TestDspFieldOrder:
    def test_issued_ticks_y_and_names_the_field_order(self):
        sheet = dsp_sheet(dsp_page(dsp_data(fieldOrderIssued=True, fieldOrderNumber="FO-12")))
        assert (sheet["AI87"].value, sheet["AK87"].value) == ("X", None)
        assert sheet["AA87"].value == "Field Order Issued? (No. FO-12)"

    def test_not_issued_ticks_n(self):
        sheet = dsp_sheet(dsp_page(dsp_data(fieldOrderIssued=False, fieldOrderNumber="FO-12")))
        assert (sheet["AI87"].value, sheet["AK87"].value) == (None, "X")
        assert sheet["AA87"].value == "Field Order Issued?"

    def test_unanswered_ticks_neither(self):
        sheet = dsp_sheet(dsp_page(dsp_data(fieldOrderIssued=None)))
        assert (sheet["AI87"].value, sheet["AK87"].value) == (None, None)

    def test_the_boxes_are_see_through_so_the_x_shows(self):
        content = dsp_page(dsp_data(fieldOrderIssued=True, fieldOrderNumber="FO-12"))
        package = zipfile.ZipFile(io.BytesIO(content))
        drawings = [package.read(name).decode() for name in package.namelist()
                    if name.startswith("xl/drawings/drawing") and 'name="Rectangle 1"' in package.read(name).decode()
                    and "<a:off x=\"4991100\"" in package.read(name).decode()]
        assert len(drawings) == 1
        for shape in re.findall(r"<xdr:sp\b.*?</xdr:sp>", drawings[0], re.DOTALL):
            assert "<a:solidFill><a:srgbClr val=\"FFFFFF\"/></a:solidFill><a:ln" not in shape

    def test_the_template_file_keeps_its_filled_boxes(self):
        package = zipfile.ZipFile(export.TEMPLATE_PATH)
        assert any('<a:srgbClr val="FFFFFF"/></a:solidFill><a:ln' in package.read(name).decode()
                   for name in package.namelist() if name.startswith("xl/drawings/drawing"))


class TestDspSignatureLayout:
    def test_dsp_fr_has_an_re_layout_and_no_inspector_layout(self):
        assert export.RE_SIGNATURE_LAYOUTS["DSP Fr"] is export_dsp.RE_SIGNATURE_LAYOUT
        assert "DSP Fr" not in export.SIGNATURE_LAYOUTS

    def test_the_layout_matches_the_template(self):
        template = openpyxl.load_workbook(export.TEMPLATE_PATH)["DSP Fr"]
        merged = {str(area) for area in template.merged_cells.ranges}
        layout = export_dsp.RE_SIGNATURE_LAYOUT
        assert {"Q87:Z87", "Q88:Z88", "AN87:AQ87"} <= merged
        assert template["Q88"].value == "Resident Engineer" and template["AN88"].value == "Date"
        assert all(template[f"{column}87"].border.bottom.style for column in "QRSTUVWXYZ")
        assert (layout.caption_cells, layout.date_cell) == ("Q88:Z88", "AN87")
        assert sum(layout.column_widths_px) == layout.signature_cx_emu // 9525 == 145


class TestDspExport:
    def test_a_dsp_report_prints_on_dsp_fr_and_dsp_bk_stays_hidden(self):
        content = export_bytes(reports=[dsp_row()])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "DSP Fr"]

    def test_without_a_dsp_report_both_pages_stay_hidden(self):
        sheets = visible_sheets(export_bytes(reports=[swcb_row(1, 2)]))
        assert "DSP Fr" not in sheets and "DSP Bk" not in sheets

    def test_it_prints_among_the_main_reports_in_page_order(self):
        content = export_bytes(reports=[swcb_row(1, 2), dsp_row(1, 3), swcb_row(2, 4)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "DSP Fr", "Conc Fr 2", "Conc Bk 2"]

    def test_one_with_a_parent_follows_it(self):
        swcb, other = swcb_row(1, 2), swcb_row(2, 4)
        content = export_bytes(reports=[swcb, dsp_row(1, 3, parent=swcb["report_id"]), other])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "DSP Fr", "Conc Fr 2", "Conc Bk 2"]

    def test_its_attachments_follow_dsp_fr_before_the_next_report(self):
        dsp = dsp_row(1, 3)
        rows = [attachment(1, dsp["report_id"]), attachment(2, dsp["report_id"])]
        content = export_bytes(reports=[swcb_row(1, 2), dsp, swcb_row(2, 4)], attachments=rows,
                               files={r["storage_path"]: image_bytes("JPEG") for r in rows})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "DSP Fr", "Attachments 1",
                                           "Attachments 2", "Conc Fr 2", "Conc Bk 2"]

    def test_every_dsp_report_gets_a_sheet_of_its_own_numbered_as_a_page(self):
        content = export_bytes(idr={**SUBMITTED_IDR, "total_pages": 3},
                               reports=[dsp_row(1, 2, data=dsp_data(timeOfPatrol="09:00")),
                                        dsp_row(2, 3, data=dsp_data(timeOfPatrol="15:00"))])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "DSP Fr", "DSP Fr 2"]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert [(book[name]["AK15"].value, book[name]["AJ13"].value, book[name]["AO13"].value)
                for name in ("DSP Fr", "DSP Fr 2")] == [("09:00", 2, 3), ("15:00", 3, 3)]

    def test_a_draft_marks_the_page_and_signs_nothing(self):
        content = export_bytes(idr=DRAFT_IDR, general={**GENERAL, "page_number": None},
                               reports=[dsp_row(page_number=None)])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["DSP Fr"]
        assert DRAFT_MARKER in [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert [p for p in stamped(content, "DSP Fr") if p["format"] == "PNG"] == []

    def test_a_submitted_idr_prints_no_signature_and_no_date(self):
        content = export_bytes(idr=SIGNED_IDR, signature=signature_png(), reports=[dsp_row()])
        assert signatures_on(content, "DSP Fr") == []
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["DSP Fr"]
        assert sheet["AN87"].value is None and sheet["Q88"].value == "Resident Engineer"

    def test_an_approved_idr_prints_only_the_res_signature_dated_the_day_it_was_approved(self):
        content = export_bytes(idr={**APPROVED_IDR, "total_pages": 2}, signature=signature_png(),
                               re_signature=signature_png((300, 90)), users={RE_UUID: RE_USER}, reports=[dsp_row()])
        (resident,) = signatures_on(content, "DSP Fr")
        assert resident["size"] == (300, 90) and 16 <= resident["col"] <= 25  # within Q:Z
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["DSP Fr"]
        assert sheet["AN87"].value == APPROVED_ON
        assert sheet["Q88"].value == "RE: Rex Resident"
        assert sheet["AJ12"].value == "005"

    def test_an_approved_idr_whose_re_file_is_missing_leaves_the_date_blank(self):
        content = export_bytes(idr={**APPROVED_IDR, "total_pages": 2}, signature=signature_png(),
                               re_signature=None, reports=[dsp_row()])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["DSP Fr"]
        assert sheet["AN87"].value is None

    @pytest.mark.parametrize("answer, column", list(ANSWER_COLUMNS.items()))
    def test_the_dispatcher_hands_render_the_reports_answers(self, answer, column):
        content = export_bytes(reports=[dsp_row(data=dsp_data(mos=[(answer, "x")]))])
        sheet = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["DSP Fr"]
        assert sheet[f"{column}20"].value == "X"
