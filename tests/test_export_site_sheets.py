"""
An IDR's imported sites on its export (W5): a sidewalk's PIR pages and a ped ramp corner's inspection, printed as
unnumbered pages after their report's attachments (export_site_sheets, through export_attachments).
"""

import io
import zipfile
from contextlib import contextmanager
from typing import Any, Optional
from unittest.mock import patch
from uuid import UUID

import openpyxl
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen.canvas import Canvas

from api.queries.idr_report_sites import idr_site_sheets
from api.services import export_site_sheets
from api.services.export import generate_idr_export
from tests.v1.test_exports import (
    GENERAL_ROW, IDR_ID, SIGNED_IDR, SUBMITTED_IDR, attachment, image_bytes, patched_export, pictures,
    print_order_page_numbers, signature_png, signatures_on, swcb_row, visible_sheets,
)

SWCB_1, SWCB_2 = swcb_row(1, 2)["report_id"], swcb_row(2, 3)["report_id"]
CORNER = UUID(int=0xC0C0)
PIR_TITLE = "PIR — 120 Main St — page {} of {}"
DOT_TITLE = "DOT ped ramp inspection — 3rd Ave & 38 Street NE — page {} of {}"
ICID_TITLE = "Ped ramp inspection (ICID) — 3rd Ave & 38 Street NE — page {} of {}"


def sidewalk(number: int, report_id: UUID, pages: int = 2, **fields) -> dict:
    """
    Build a sidewalk site row as idr_site_sheets returns it, its PIR's pages at pir/{number}/page-{n}.png.
    Takes its number, its report's id, how many PIR pages it has and field overrides.
    Returns the row.
    """
    return {"id": UUID(int=0x5100 + number), "report_id": report_id, "property_id": UUID(int=0x9900 + number),
            "kind": "sidewalk", "boro": "Queens", "address": "120 Main St", "block": "123", "lot": "45",
            "intersection_street_1": None, "intersection_street_2": None, "quadrant": None, "facing_street": None,
            "corner_id": None, "north_south_street": None, "east_west_street": None, "corner_direction": None,
            "pir_pages": [{"page_number": n, "image_path": f"pir/{number}/page-{n}.png"} for n in range(1, pages + 1)],
            "dot_pdf_path": None, **fields}


def ped_ramp(number: int, report_id: UUID, dot_pdf_path: Optional[str] = None, **fields) -> dict:
    """
    Build a ped ramp site row as idr_site_sheets returns it, at CORNER (3rd Ave & 38 Street, NE).
    Takes its number, its report's id, its corner's filed DOT forms (None for none) and field overrides.
    Returns the row.
    """
    return {"id": UUID(int=0x5200 + number), "report_id": report_id, "property_id": UUID(int=0x9A00 + number),
            "kind": "ped_ramp", "boro": "Queens", "address": None, "block": None, "lot": None,
            "intersection_street_1": "3rd Ave", "intersection_street_2": "38 Street", "quadrant": "NE",
            "facing_street": "street_1", "corner_id": CORNER, "north_south_street": "3rd Ave",
            "east_west_street": "38 Street", "corner_direction": "NE", "pir_pages": [], "dot_pdf_path": dot_pdf_path,
            **fields}


def pdf_bytes(*sizes: tuple[float, float]) -> bytes:
    """
    Make a PDF in memory, a page per size given.
    Takes each page's size in points.
    Returns the PDF's bytes.
    """
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    for size in sizes:
        canvas.setPageSize(size)
        canvas.drawString(72, 72, "page")
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def zip_bytes(members: dict[str, bytes]) -> bytes:
    """
    Make a ZIP in memory.
    Takes {member name: bytes}.
    Returns the ZIP's bytes.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for name, data in members.items():
            package.writestr(name, data)
    return buffer.getvalue()


PORTRAIT, LANDSCAPE, SQUARE = letter, landscape(letter), (400, 400)
ICID_REPORT = pdf_bytes(PORTRAIT, PORTRAIT)  # the corner page and one ramp page


@contextmanager
def site_export(sites: Optional[list[dict]], site_files: Optional[dict[tuple[str, str], Any]] = None,
                icid_report: Any = ICID_REPORT, **overrides):
    """
    Patch an export's imported sites and their files on top of patched_export.
    Takes the site rows (None: the read fails), {(bucket, path): bytes or an exception} for Storage (anything else is
    missing), ICID's inspection report PDF (or an exception drawing it raises) and patched_export's overrides.
    Yields a dict of the mocks: patched_export's, and sites, site_download and get_corner.
    """
    def download(path, bucket):
        result = (site_files or {}).get((bucket, path), FileNotFoundError(path))
        if isinstance(result, Exception):
            raise result
        return result

    def draw(report):
        if isinstance(icid_report, Exception):
            raise icid_report
        return icid_report

    with (
        patched_export(**overrides) as mocks,
        patch.object(export_site_sheets, "idr_site_sheets", return_value=sites) as read,
        patch.object(export_site_sheets, "download_file", side_effect=download) as site_download,
        patch.object(export_site_sheets, "get_corner", return_value={"corner_id": CORNER}) as corner,
        patch.object(export_site_sheets, "inspection_report", return_value=object()),
        patch.object(export_site_sheets, "render_pdf", side_effect=draw),
    ):
        yield {**mocks, "sites": read, "site_download": site_download, "get_corner": corner}


def run(sites: Optional[list[dict]], **kwargs) -> tuple[bytes, dict]:
    """
    Export the IDR with the given sites.
    Takes the site rows and site_export's keyword arguments.
    Returns (the .xlsx bytes, the mocks).
    """
    kwargs.setdefault("reports", [GENERAL_ROW, swcb_row(1, 2)])
    with site_export(sites, **kwargs) as mocks:
        return generate_idr_export(IDR_ID).content, mocks


def titles(content: bytes, sheets: list[str]) -> list:
    """
    Read the titles (B21) of an export's pages.
    Takes the .xlsx bytes and the sheet names.
    Returns each one's title.
    """
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    return [book[name]["B21"].value for name in sheets]


def picture_sizes(content: bytes, sheets: list[str]) -> list[tuple[int, int]]:
    """
    Read the size of the one picture on each page.
    Takes the .xlsx bytes and the sheet names.
    Returns each page's picture's pixel size.
    """
    sizes = []
    for name in sheets:
        found = pictures(content, name)
        assert len(found) == 1, name
        sizes.append(found[0]["size"])
    return sizes


PIR_FILES = {("pir-images", f"pir/1/page-{n}.png"): image_bytes("PNG", (60 + n, 40)) for n in (1, 2)}


class TestSidewalkSites:
    def test_the_pir_pages_follow_the_reports_attachments_in_page_order(self):
        photo = attachment(1, SWCB_1)
        content, mocks = run([sidewalk(1, SWCB_1)], site_files=PIR_FILES, attachments=[photo],
                             files={photo["storage_path"]: image_bytes("JPEG")})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk",
                                           "Attachments 1", "Attachments 2", "Attachments 3"]
        assert titles(content, ["Attachments 2", "Attachments 3"]) == [PIR_TITLE.format(1, 2), PIR_TITLE.format(2, 2)]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Attachments 2"]["B22"].value == "Queens Block 123 Lot 45"
        assert picture_sizes(content, ["Attachments 2", "Attachments 3"]) == [(61, 40), (62, 40)]
        # Site pages, like attachment pages, carry no page number
        assert print_order_page_numbers(content)[-2:] == [None, None]
        mocks["sites"].assert_called_once_with(IDR_ID)

    def test_a_site_without_a_pir_adds_nothing(self):
        content, _ = run([sidewalk(1, SWCB_1, pages=0)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]

    def test_a_page_that_cant_be_fetched_or_read_is_left_out(self):
        files = {("pir-images", "pir/1/page-1.png"): b"not an image", **{k: v for k, v in PIR_FILES.items()
                                                                          if k[1].endswith("page-2.png")}}
        content, _ = run([sidewalk(1, SWCB_1, pages=3)], site_files=files)
        assert visible_sheets(content)[4:] == ["Attachments 1"]
        assert titles(content, ["Attachments 1"]) == ["PIR — 120 Main St — page 2 of 3"]

    def test_each_reports_sites_follow_that_report(self):
        reports = [GENERAL_ROW, swcb_row(1, 2), swcb_row(2, 3)]
        sites = [sidewalk(1, SWCB_2, pages=1), sidewalk(2, SWCB_1, pages=1, address="9 Elm St")]
        files = {("pir-images", "pir/1/page-1.png"): image_bytes("PNG"),
                 ("pir-images", "pir/2/page-1.png"): image_bytes("PNG")}
        content, _ = run(sites, reports=reports, site_files=files, idr={**SUBMITTED_IDR, "total_pages": 3})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Attachments 1",
                                           "Conc Fr 2", "Conc Bk 2", "Attachments 2"]
        assert titles(content, ["Attachments 1", "Attachments 2"]) == ["PIR — 9 Elm St — page 1 of 1",
                                                                       "PIR — 120 Main St — page 1 of 1"]

    def test_an_unprinted_reports_sites_go_at_the_end(self):
        unprinted = {"report_id": UUID(int=0x5E01), "report_type": "SWR", "is_addendum": False, "page_number": 3,
                     "report_data": {}}
        content, _ = run([sidewalk(1, unprinted["report_id"], pages=1)], reports=[GENERAL_ROW, unprinted],
                         site_files={("pir-images", "pir/1/page-1.png"): image_bytes("PNG")})
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Attachments 1"]

    def test_site_pages_carry_the_inspectors_signature(self):
        content, _ = run([sidewalk(1, SWCB_1, pages=1)], idr=SIGNED_IDR, signature=signature_png(),
                         site_files={("pir-images", "pir/1/page-1.png"): image_bytes("JPEG")})
        assert len(signatures_on(content, "Attachments 1")) == 1


class TestPedRampSites:
    def test_an_approved_inspections_dot_pdf_prints_a_page_a_picture(self):
        path = "inspections/1/form.pdf"
        content, mocks = run([ped_ramp(1, SWCB_1, path)],
                             site_files={("pedramp-exports", path): pdf_bytes(PORTRAIT, LANDSCAPE)})
        assert visible_sheets(content)[4:] == ["Attachments 1", "Attachments 2"]
        assert titles(content, ["Attachments 1", "Attachments 2"]) == [DOT_TITLE.format(1, 2), DOT_TITLE.format(2, 2)]
        # 150 DPI, shrunk to 1600 px on the long side
        assert picture_sizes(content, ["Attachments 1", "Attachments 2"]) == [(1236, 1600), (1600, 1236)]
        mocks["get_corner"].assert_not_called()

    def test_a_packages_inspection_copies_print_in_order_and_its_tifs_stay_out(self):
        path = "inspections/1/package.zip"
        package = zip_bytes({"tif_CT-A.pdf": pdf_bytes(SQUARE), "inspection_2.pdf": pdf_bytes(LANDSCAPE),
                             "inspection_1.pdf": pdf_bytes(PORTRAIT)})
        content, _ = run([ped_ramp(1, SWCB_1, path)], site_files={("pedramp-exports", path): package})
        assert titles(content, ["Attachments 1", "Attachments 2"]) == [DOT_TITLE.format(1, 2), DOT_TITLE.format(2, 2)]
        assert picture_sizes(content, ["Attachments 1", "Attachments 2"]) == [(1236, 1600), (1600, 1236)]

    def test_a_package_with_one_copy_names_it_inspection_pdf(self):
        path = "inspections/1/package.zip"
        package = zip_bytes({"inspection.pdf": pdf_bytes(PORTRAIT), "tif_1.pdf": pdf_bytes(SQUARE)})
        content, _ = run([ped_ramp(1, SWCB_1, path)], site_files={("pedramp-exports", path): package})
        assert visible_sheets(content)[4:] == ["Attachments 1"]

    def test_without_an_approved_form_icids_report_stands_in(self):
        content, mocks = run([ped_ramp(1, SWCB_1)])
        assert titles(content, ["Attachments 1", "Attachments 2"]) == [ICID_TITLE.format(1, 2),
                                                                       ICID_TITLE.format(2, 2)]
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        assert book["Attachments 1"]["B22"].value.startswith("ICID's inspection report of the corner")
        mocks["get_corner"].assert_called_once_with(CORNER)

    def test_a_dot_form_that_cant_be_fetched_or_rendered_falls_back_to_icids_report(self):
        path = "inspections/1/form.pdf"
        for files in ({}, {("pedramp-exports", path): b"%PDF-1.4 broken"},
                      {("pedramp-exports", path): zip_bytes({"tif_1.pdf": pdf_bytes(SQUARE)})}):
            content, _ = run([ped_ramp(1, SWCB_1, path)], site_files=files)
            assert titles(content, ["Attachments 1"]) == [ICID_TITLE.format(1, 2)]

    def test_a_corner_prints_once_for_two_of_its_ramps(self):
        reports = [GENERAL_ROW, swcb_row(1, 2), swcb_row(2, 3)]
        content, mocks = run([ped_ramp(1, SWCB_2), ped_ramp(2, SWCB_1), ped_ramp(3, SWCB_1)], reports=reports,
                             idr={**SUBMITTED_IDR, "total_pages": 3})
        # With the first report, in print order, that brings it
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk", "Attachments 1",
                                           "Attachments 2", "Conc Fr 2", "Conc Bk 2"]
        mocks["get_corner"].assert_called_once_with(CORNER)

    def test_a_ramp_without_a_live_corner_adds_nothing(self):
        content, _ = run([ped_ramp(1, SWCB_1, corner_id=None)])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]

    def test_nothing_prints_when_neither_can_be_had(self):
        content, _ = run([ped_ramp(1, SWCB_1)], icid_report=RuntimeError("no ramps"))
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]


class TestNoSites:
    def test_no_sites_leaves_the_export_as_it_was(self):
        content, _ = run([])
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]

    def test_sites_that_cant_be_read_dont_stop_the_export(self):
        content, mocks = run(None)
        assert visible_sheets(content) == ["Gen Fr", "Gen Bk", "Conc Fr", "Conc Bk"]
        mocks["site_download"].assert_not_called()


class TestIdrSiteSheetsQuery:
    def test_reads_the_idrs_sites_with_their_pir_pages_and_filed_forms(self):
        with patch("api.queries.idr_report_sites.run_query", return_value=[{"id": 1}]) as run_query:
            assert idr_site_sheets(IDR_ID) == [{"id": 1}]
        sql, params = run_query.call_args.args
        assert params == (IDR_ID,)
        for piece in ("icid.idr_report_sites rs", "icid.pir_pages pg", "pg.deleted_at IS NULL",
                      "pir.deleted_at IS NULL", "icid.ped_ramp_inspections i", "i.status = 'approved'",
                      "c.deleted_at IS NULL", "p.deleted_at IS NULL", "ORDER BY pg.page_number"):
            assert piece in sql
