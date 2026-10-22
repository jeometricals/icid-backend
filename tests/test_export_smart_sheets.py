"""
The SMART sheets an IDR's imported sites bring into its export (Slice X.4, export_smart_sheets): a sidewalk's flags,
curb segments and pay items, and a ped ramp corner's inspection, as tables, each before the site's own pages.
"""

import io
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from uuid import UUID

import openpyxl
import pytest

from api.schemas.ped_ramp import (
    CORNER_LABELS, AdaWarning, CornerOut, InspectionRamp, InspectionReport, RampAttributesOut,
)
from api.services.export import TEMPLATE_PATH
from api.services.export_smart_sheets import CORNER_FIELDS, is_smart_sheet, sheet_name, site_pay_item_rows
from api.services.xlsx_template import WorkbookTemplate
from tests.test_export_site_sheets import CORNER, PIR_FILES, SMART, SWCB_1, ped_ramp, run, sidewalk
from tests.v1.test_exports import APPROVED_IDR, DRAFT_IDR, IDR_ID, visible_sheets

SITE = sidewalk(1, SWCB_1)
PROPERTY = SITE["property_id"]
SF_ITEM, LF_ITEM, EXTRA_ITEM = UUID(int=0xA1), UUID(int=0xA2), UUID(int=0xA3)
CREATED = datetime(2026, 10, 1, tzinfo=timezone.utc)


def flag(number: int, amount: str = "90", state: str = "attributed", **fields) -> dict:
    """
    Build a flag row as list_idr_flags returns it, on SITE's property, page 1, mapped to SF_ITEM.
    Takes its number, its amount in SF, its state and field overrides.
    Returns the row.
    """
    return {"flag_id": UUID(int=0xF000 + number), "page_number": 1, "display_order": number, "shape": "rectangle",
            "surveyed_thickness": "4in", "property_id": PROPERTY, "pay_item_id": SF_ITEM, "pay_item_ref": "4.01 AA",
            "budget_code": "SE0944", "pay_item_description": "4in concrete sidewalk", "unit": "SF",
            "computed_quantity": Decimal("100.00"), "reported_quantity": None, "amount": Decimal(amount),
            "overridden": False, "state": state, **fields}


def segment(number: int, length: str = "20", **fields) -> dict:
    """
    Build a curb segment row as list_idr_segments returns it, on SITE's property, page 2, mapped to LF_ITEM.
    Takes its number, its length in feet and field overrides.
    Returns the row.
    """
    return {"segment_id": UUID(int=0x5E00 + number), "page_number": 2, "segment_type": "curb_concrete",
            "length_ft": Decimal(length), "point_count": 4, "property_id": PROPERTY, "pay_item_id": LF_ITEM,
            "pay_item_ref": "4.08 AA", "budget_code": "SE0944", "pay_item_description": "Concrete curb", "unit": "LF",
            "amount": Decimal(length), "overridden": False, "state": "attributed", **fields}


def site_item(item_id: UUID, quantity: str, overrides_auto: bool = False, **fields) -> dict:
    """
    Build a site pay item as idr_site_pay_items returns it, on SITE.
    Takes its contract item, its quantity, whether it overrides the automatic quantity and field overrides.
    Returns the row.
    """
    refs = {SF_ITEM: ("4.01 AA", "4in concrete sidewalk", "S.F."), LF_ITEM: ("4.08 AA", "Concrete curb", "L.F."),
            EXTRA_ITEM: ("6.59 P", "Pedestrian ramp", "EACH")}
    ref, description, unit = refs[item_id]
    return {"id": UUID(int=0x1700 + int(quantity.replace(".", ""))), "report_site_id": SITE["id"],
            "contract_item_id": item_id, "quantity": Decimal(quantity), "note": None, "overrides_auto": overrides_auto,
            "pay_item_ref": ref, "budget_code": "SE0944", "description": description, "unit": unit,
            "report_id": SWCB_1, "property_id": PROPERTY, "override_dropped": False, **fields}


def rows_of(content: bytes, sheet: str) -> list[list]:
    """
    Read a sheet's rows as values, each padded with None to the sheet's eight columns.
    Takes the .xlsx bytes and the sheet's name.
    Returns the rows.
    """
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    return [list(row) + [None] * (8 - len(row)) for row in book[sheet].iter_rows(values_only=True)]


def sheet_xml(content: bytes, sheet: str) -> str:
    """
    Read a sheet's worksheet XML from the package.
    Takes the .xlsx bytes and the sheet's name.
    Returns the XML.
    """
    package = zipfile.ZipFile(io.BytesIO(content))
    workbook = package.read("xl/workbook.xml").decode()
    rel = re.search(rf'<sheet name="{re.escape(sheet)}" [^>]*r:id="([^"]+)"', workbook).group(1)
    rels = package.read("xl/_rels/workbook.xml.rels").decode()
    target = re.search(rf'Id="{rel}" [^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)" [^>]*Id="{rel}"', rels)
    return package.read("xl/" + target.group(1).removeprefix("/xl/")).decode()


def row_starting(rows: list[list], first: str) -> list:
    """
    Find the first row whose first cell is the text given.
    Takes the rows and the text.
    Returns the row.
    """
    return next(row for row in rows if row and row[0] == first)


class TestSheetNames:
    def test_an_address_in_capitals_with_underscores(self):
        assert sheet_name("SMART_PIR_", "152 23 St", set()) == "SMART_PIR_152_23_ST"
        assert sheet_name("SMART_INSP_", "3rd Ave & 38 Street NE", set()) == "SMART_INSP_3RD_AVE_38_STREET_NE"

    def test_cut_to_excels_31_characters(self):
        name = sheet_name("SMART_PIR_", "1234 Very Long Boulevard Name Extension", set())
        assert name == "SMART_PIR_1234_VERY_LONG_BOULEV" and len(name) == 31

    def test_a_taken_name_gets_a_number_within_the_limit(self):
        long = "1234 Very Long Boulevard Name Extension"
        first = sheet_name("SMART_PIR_", long, set())
        second = sheet_name("SMART_PIR_", long, {first})
        assert second == "SMART_PIR_1234_VERY_LONG_BOUL_2" and len(second) == 31
        assert sheet_name("SMART_PIR_", "9 Elm", {"smart_pir_9_elm"}) == "SMART_PIR_9_ELM_2"

    def test_smart_sheets_are_told_by_their_prefix(self):
        assert is_smart_sheet("SMART_PIR_9_ELM") and is_smart_sheet("SMART_INSP_X")
        assert not is_smart_sheet("Attachments 1") and not is_smart_sheet("Gen Fr")


class TestSitePayItemRows:
    def test_the_drawings_come_to_an_auto_row_per_pay_item_and_totals_per_unit(self):
        rows = site_pay_item_rows([flag(1), flag(2, "60.5")], [segment(1), segment(2, "4.25")], [])
        assert [(r["item_no"], r["quantity"], r["source"]) for r in rows] == [
            ("4.01 AA", "150.50 SF", "auto (2 flags)"), ("4.08 AA", "24.25 LF", "auto (2 curb segments)"),
            ("Total", "24.25 LF", ""), ("Total", "150.50 SF", "")]
        assert [r.get("_total", False) for r in rows] == [False, False, True, True]

    def test_an_override_replaces_the_auto_amount_and_says_what_it_replaced(self):
        rows = site_pay_item_rows([flag(1), flag(2)], [], [site_item(SF_ITEM, "185", overrides_auto=True)])
        assert rows[0]["quantity"] == "185.00 SF"
        assert rows[0]["source"] == "auto (2 flags) · edited, originally 180.00 SF"
        assert rows[-1]["quantity"] == "185.00 SF"

    def test_an_override_an_approval_dropped_bills_nothing(self):
        rows = site_pay_item_rows([flag(1)], [], [site_item(SF_ITEM, "185", overrides_auto=True, override_dropped=True)])
        assert rows == [{"item_no": "4.01 AA", "budget_code": "SE0944", "description": "4in concrete sidewalk",
                         "quantity": "185.00 SF", "source": "override dropped at approval (bills nothing)"}]

    def test_entered_items_and_an_override_without_drawings_follow(self):
        rows = site_pay_item_rows([], [], [site_item(EXTRA_ITEM, "2", note="NE corner"),
                                           site_item(SF_ITEM, "0", overrides_auto=True)])
        assert [(r["item_no"], r["quantity"], r["source"]) for r in rows] == [
            ("6.59 P", "2.00 EA", "entered: NE corner"), ("4.01 AA", "0.00 SF", "override (no flags left)"),
            ("Total", "2.00 EA", ""), ("Total", "0.00 SF", "")]

    def test_only_the_drawings_that_count_add_up(self):
        rows = site_pay_item_rows([flag(1), flag(2, state="excluded"), flag(3, state="claimed_elsewhere"),
                                   {**flag(4), "amount": None}], [], [])
        assert rows[0]["quantity"] == "90.00 SF" and rows[0]["source"] == "auto (1 flag)"


class TestSidewalkSheet:
    def export(self, **smart) -> bytes:
        """
        Export the IDR with SITE imported and its PIR's pages, the SMART sheet reading what is given.
        Takes the SMART reads ({flags, segments, site_items}).
        Returns the .xlsx bytes.
        """
        content, _ = run([SITE], site_files=PIR_FILES, smart=smart)
        return content

    def test_it_prints_before_the_sites_pir_pages(self):
        content = self.export(flags=[flag(1)])
        assert visible_sheets(content)[4:] == [SMART, "Attachments 1", "Attachments 2"]

    def test_header_flags_segments_and_pay_items(self):
        flags = [flag(1, reported_quantity=Decimal("95")), flag(2, state="excluded")]
        content = self.export(flags=flags, segments=[segment(1)], site_items=[site_item(EXTRA_ITEM, "1")])
        rows = rows_of(content, SMART)
        assert rows[1][0] == "SMART PIR — 120 Main St"
        assert rows[2][0] == "Queens Block 123 Lot 45"
        assert row_starting(rows, "Inspector")[:6] == ["Inspector", None, "Genghis Khan", None, "IDR No.", "(not yet numbered)"]
        assert row_starting(rows, "Work date")[:6] == ["Work date", None, "9/30/26", None, "Status", "Submitted"]
        assert row_starting(rows, "Page") == ["Page", "#", "Shape", "Thickness", "Pay item", "Area (computed)",
                                              "Reported", "Derived quantity"]
        assert row_starting(rows, 1) == [1, 1, "Rectangle", '4"', "4.01 AA — 4in concrete sidewalk", "100.00 SF",
                                         "95.00 SF", "90.00 SF"]
        assert any(row[0] == "1 more flag on this site is not this IDR's (excluded, or claimed by another IDR)."
                   for row in rows)
        segments_at = rows.index(row_starting(rows, "Curb segments"))
        assert rows[segments_at + 1][:6] == ["Page", "Segment #", "Shape", "Dimensions", "Pay item", "Derived quantity"]
        assert rows[segments_at + 2][:6] == [2, 1, "Concrete curb", "20.00 ft · 3 runs", "4.08 AA — Concrete curb",
                                             "20.00 LF"]
        items_at = rows.index(row_starting(rows, "Pay items for this site"))
        assert [rows[items_at + 1][i] for i in (0, 2, 3, 5, 7)] == ["Item", "Budget code", "Description", "Quantity",
                                                                    "Source"]
        assert [r[0] for r in rows[items_at + 2: items_at + 8]] == ["4.01 AA", "4.08 AA", "6.59 P", "Total", "Total",
                                                                     "Total"]

    def test_an_overridden_flag_is_starred_and_the_override_marked(self):
        content = self.export(flags=[flag(1, overridden=True)],
                              site_items=[site_item(SF_ITEM, "185", overrides_auto=True)])
        rows = rows_of(content, SMART)
        assert row_starting(rows, 1)[7] == "90.00 SF *"
        assert any(row[0] and str(row[0]).startswith("* Its pay item is overridden") for row in rows)
        item = row_starting(rows, "4.01 AA")
        assert item[5] == "185.00 SF" and item[7] == "auto (1 flag) · edited, originally 90.00 SF"

    def test_it_prints_on_letter_landscape_a_page_wide_with_half_inch_margins(self):
        content = self.export(flags=[flag(1)])
        xml = sheet_xml(content, SMART)
        setup = re.search(r"<pageSetup\b[^>]*/>", xml).group(0)
        for attribute in ('paperSize="1"', 'orientation="landscape"', 'fitToWidth="1"', 'fitToHeight="0"'):
            assert attribute in setup
        assert '<pageMargins left="0.5" right="0.5" top="0.5" bottom="0.5"' in xml
        assert '<pageSetUpPr fitToPage="1"/>' in xml
        book = openpyxl.load_workbook(io.BytesIO(content))
        sheet = book[SMART]
        assert sheet.page_setup.paperSize == 1 and sheet.page_setup.orientation == "landscape"
        assert sheet.print_area == f"'{SMART}'!$A$1:$H${sheet.max_row}"
        assert sheet.page_margins.left == 0.5

    def test_a_section_that_wouldnt_start_on_the_page_goes_to_the_next(self):
        # 30 flags fill a landscape page to within a few rows of its foot: Curb segments' title and first rows don't fit
        content = self.export(flags=[flag(n) for n in range(1, 31)], segments=[segment(1)])
        rows = rows_of(content, SMART)
        segments_at = rows.index(row_starting(rows, "Curb segments")) + 1  # its row number
        breaks = [int(b) for b in re.findall(r'<brk id="(\d+)"', sheet_xml(content, SMART))]
        assert segments_at - 1 in breaks
        assert 'manualBreakCount="' in sheet_xml(content, SMART)

    def test_a_short_sheet_has_no_page_breaks(self):
        assert "<rowBreaks" not in sheet_xml(self.export(flags=[flag(1)]), SMART)

    def test_a_drafts_sheet_carries_the_draft_marker_and_no_signature(self):
        content, _ = run([SITE], site_files=PIR_FILES, smart={"flags": [flag(1)]}, idr=DRAFT_IDR)
        assert rows_of(content, SMART)[0][1] == "DRAFT - Not for Submission"
        assert "<drawing" not in sheet_xml(content, SMART)

    def test_an_approved_idr_reads_the_drawings_it_claimed(self):
        content, mocks = run([SITE], site_files=PIR_FILES, smart={"flags": [flag(1, state="claimed")]},
                             idr=APPROVED_IDR)
        mocks["smart_flags"].assert_called_once_with(IDR_ID, True)
        assert row_starting(rows_of(content, SMART), "Inspector")[4:6] == ["IDR No.", "005"]

    def test_a_site_that_cant_be_read_gets_no_sheet_and_its_pir_pages_still_print(self):
        content = self.export(flags=None)
        assert visible_sheets(content)[4:] == ["Attachments 1", "Attachments 2"]

    def test_two_sites_at_one_address_get_two_sheets(self):
        sites = [SITE, sidewalk(2, SWCB_1, pages=0, lot="46")]
        content, _ = run(sites, site_files=PIR_FILES)
        assert [s for s in visible_sheets(content) if is_smart_sheet(s)] == [SMART, SMART + "_2"]

    def test_an_idr_without_sites_has_none(self):
        content, _ = run([])
        assert not any(is_smart_sheet(s) for s in visible_sheets(content))


def corner_report() -> InspectionReport:
    """
    Build a corner's inspection report with two perpendicular ramps, the first with a running slope over 8.3%.
    Takes nothing.
    Returns the InspectionReport.
    """
    corner = CornerOut.model_validate({
        "corner_id": CORNER, "project_id": "HWS0023", "boro": "Queens", "north_south_street": "3rd Ave",
        "east_west_street": "38 Street", "corner_direction": "NE", "contractor_name": "Benny Bowers",
        "created_at": CREATED, "updated_at": CREATED, "created_by": UUID(int=1)})
    ramps = []
    for number, slope in ((1, "9.5"), (2, "7")):
        property_id = UUID(int=0x9A00 + number)
        warnings = [AdaWarning(field="ramp_running_slope_pct", value=Decimal(slope), limit="8.3%",
                               message="Ramp running slope 9.5% is over 8.3%")] if number == 1 else []
        attributes = RampAttributesOut.model_validate({
            "property_id": property_id, "ramp_type": "PERPENDICULAR", "created_at": CREATED, "updated_at": CREATED,
            "ramp_running_slope_pct": Decimal(slope), "ramp_width_in": Decimal("48"), "warnings": warnings})
        ramps.append(InspectionRamp(property_id=property_id, intersection_street_1="3rd Ave",
                                    intersection_street_2="38 Street", quadrant="NE", facing_street="street_1",
                                    ramp_type="PERPENDICULAR", ramp_id_label=f"R{number}", has_attributes=True,
                                    attributes=attributes))
    return InspectionReport(corner=corner, ramp_types=["PERPENDICULAR"], ramps=ramps)


INSPECTION_SHEET = "SMART_INSP_3RD_AVE_38_STREET_NE"


class TestCornerSheet:
    def export(self, **smart) -> bytes:
        """
        Export the IDR with two ramps of CORNER imported, the SMART sheet reading what is given.
        Takes the SMART reads ({report, inspection, site_items}).
        Returns the .xlsx bytes.
        """
        content, _ = run([ped_ramp(1, SWCB_1), ped_ramp(2, SWCB_1)], smart={"report": corner_report(), **smart})
        return content

    def test_once_per_corner_before_its_inspection_pages(self):
        content = self.export()
        assert visible_sheets(content)[4:] == [INSPECTION_SHEET, "Attachments 1", "Attachments 2"]

    def test_inspection_corner_ramps_and_warnings(self):
        inspection = {"status": "approved", "inspector_name": "Ina Spector", "re_name": "Rex Resident",
                      "re_signed_at": datetime(2026, 10, 7, 3, 15, tzinfo=timezone.utc)}
        content = self.export(inspection=inspection)
        rows = rows_of(content, INSPECTION_SHEET)
        assert rows[1][0] == "SMART inspection — 3rd Ave & 38 Street NE"
        assert any(row[:4] == ["Status", "Approved", "Inspector", "Ina Spector"] for row in rows)
        assert any(row[:4] == ["RE", "Rex Resident", "Approved", "10/6/26"] for row in rows)
        assert any(row[:2] == ["Contractor name", "Benny Bowers"] for row in rows)
        heading = row_starting(rows, "Field")
        assert heading[:3] == ["Field", "Ramp 1: Perpendicular (R1)", "Ramp 2: Perpendicular (R2)"]
        slope = next(row for row in rows if row[0] and "running slope" in str(row[0]).lower()
                     and row[1] and "%" in str(row[1]))
        assert slope[1:3] == ["9.5%  (ADA)", "7%"]
        assert any(row[0] == "Ramp 1: Ramp running slope 9.5% is over 8.3%" for row in rows)

    def test_of_the_dot_forms_sections_1_to_3_only_the_construction_fields_print(self):
        rows = rows_of(self.export(), INSPECTION_SHEET)
        construction = rows.index(row_starting(rows, "Construction"))
        assert rows[construction + 1][:4] == ["Contractor name", "Benny Bowers", "Construction completion date", None]
        assert rows[construction + 2][0] == "Curb material" and rows[construction + 2][2] == "Sidewalk material"
        cut = [label for name, label in CORNER_LABELS.items() if name not in CORNER_FIELDS]
        assert [label for label in cut if any(label in row for row in rows)] == []
        assert not any(isinstance(row[0], str) and row[0][:2] in ("1.", "2.", "3.") for row in rows if row)

    def test_without_an_inspection_says_so(self):
        rows = rows_of(self.export(), INSPECTION_SHEET)
        assert any(row[0] == "No inspection of this corner yet." for row in rows)

    def test_the_pay_items_on_its_ramps_sites(self):
        item = site_item(EXTRA_ITEM, "2", property_id=UUID(int=0x9A01))
        other = site_item(LF_ITEM, "5", property_id=UUID(int=0xBEEF))
        rows = rows_of(self.export(site_items=[item, other]), INSPECTION_SHEET)
        items_at = rows.index(row_starting(rows, "Pay items for this site"))
        assert rows[items_at + 1][:4] == ["Item", "Description", None, "Quantity"]
        assert rows[items_at + 2][:4] == ["6.59 P (SE0944)", "Pedestrian ramp", None, "2.00 EA"]
        assert rows[items_at + 3][0] == "Total" and len(rows) == items_at + 4

    def test_portrait_letter_on_one_page(self):
        xml = sheet_xml(self.export(), INSPECTION_SHEET)
        setup = re.search(r"<pageSetup\b[^>]*/>", xml).group(0)
        for attribute in ('paperSize="1"', 'orientation="portrait"', 'fitToWidth="1"', 'fitToHeight="1"'):
            assert attribute in setup
        assert "<rowBreaks" not in xml

    def test_a_corner_that_cant_be_read_gets_no_sheet(self):
        content, _ = run([ped_ramp(1, SWCB_1)], smart={"report": None})
        assert not any(is_smart_sheet(s) for s in visible_sheets(content))


class TestAddedSheets:
    def test_an_added_sheet_keeps_its_print_area_and_breaks_when_moved(self):
        workbook = WorkbookTemplate(TEMPLATE_PATH)
        workbook.add_sheet("SMART_PIR_X", {"A": 10, "B": 20})
        workbook.set_cell("SMART_PIR_X", "A2", "hello")
        workbook.set_print_layout("SMART_PIR_X", True, "A1:B40", [20])
        workbook.set_print_layout("SMART_PIR_X", False, "A1:B50", [25, 30])  # set again: replaced, not added
        workbook.move_sheet("SMART_PIR_X", after="Gen Fr")
        workbook.show_only(["Gen Fr", "SMART_PIR_X"])
        book = openpyxl.load_workbook(io.BytesIO(workbook.to_bytes()))
        sheet = book["SMART_PIR_X"]
        assert sheet["A2"].value == "hello"
        assert sheet.print_area == "'SMART_PIR_X'!$A$1:$B$50"
        assert [b.id for b in sheet.row_breaks.brk] == [25, 30]
        assert sheet.page_setup.orientation == "portrait" and sheet.page_setup.fitToHeight == 0
        assert sheet.column_dimensions["B"].width == 20

    def test_a_taken_name_is_refused(self):
        workbook = WorkbookTemplate(TEMPLATE_PATH)
        with pytest.raises(ValueError, match="already exists"):
            workbook.add_sheet("Gen Fr", {"A": 10})
