"""
Smoke tests for the Conc Cyl page of the export template (templates/report_forms.xlsx) and its export module.
"""

from pathlib import Path

import openpyxl

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "report_forms.xlsx"


def test_conc_cyl_sheet_exists_in_template():
    assert "Conc Cyl" in openpyxl.load_workbook(TEMPLATE, read_only=True).sheetnames


def test_conc_cyl_sheet_dimensions():
    # The DDC form is 60 rows by 41 columns: a copy that lost its cells would be a bare sheet
    sheet = openpyxl.load_workbook(TEMPLATE, read_only=True)["Conc Cyl"]
    assert sheet.max_row >= 60
    assert sheet.max_column >= 41


def test_conc_cyl_module_imports():
    from api.services import export_conc_cyl

    assert export_conc_cyl.SHEET_NAME == "Conc Cyl"
