"""
The Assessment Summary (R1i): a project's assessable quantities and amounts per property and month of construction,
an .xlsx workbook with one tab per month that has any, for the OE or RE to send DOT each month. Slice C adds the same
figures a row per site, summed over a range of days (summary_sites), for the As-Builts & Assessments page and its
tailored one-sheet download.

Each flag is measured exactly as the Construction Report measures it (construction_measure.flag_areas), so a
property's row adds up to its PDFs. A flag's day is its property's Date Constructed (page 1's title block), or the
day its IDR was approved when that isn't typed (counted in the tab's footer); its month is that day's. 7" assessed as
4" counts in the 4" columns at the 4" rate, noted in Remarks. Values are written whole; the cells show two decimals.

The Assessment Summary polish (ML1a-c, ML4b): the monthly workbook has a tab for every month of its range, an empty one
a shell; both workbooks can be narrowed to a task order's properties, come as CSV too, and carry a footer on every tab
(who compiled it and when, the project's Project Engineer, and that only approved IDRs count).
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Callable, Hashable, Optional
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from api.schemas.property import display_address_of
from api.services.budget_rollup import money
from api.services.construction_measure import ZERO, flag_areas
from api.services.export_common import signed_date

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONTH = re.compile(r"^(\d{4})-(\d{2})$")
SF_FORMAT = "#,##0.00"
MONEY_FORMAT = "$#,##0.00"
DASH = "—"
COLUMNS = ("Property Address", "CB", '4" Assessable (SF)', '4" Non-Assessable (SF)', '4" Total (SF)', '4" Amount ($)',
           '7" Assessable (SF)', '7" Non-Assessable (SF)', '7" Total (SF)', '7" Amount ($)', "Grand Total ($)",
           "Outside PL (SF)", "Remarks")
COLUMN_WIDTHS = (34, 8, 14, 16, 12, 14, 14, 16, 12, 14, 15, 14, 40)
# Where the amounts and the square feet are, by column number
MONEY_COLUMNS = (6, 10, 11)
SF_COLUMNS = (3, 4, 5, 7, 8, 9, 12)
# The tab's note when any row has a part outside the property line (Slice Y2)
OUTSIDE_PL_NOTE = ("Outside PL (SF): the part of the flags outside the sidewalk strip drawn at the property line, "
                   "clipped at approval and not assessed (it is counted in Non-Assessable). Contractor billing is not "
                   "clipped.")
TOTALS_FILL = PatternFill("solid", fgColor="D9D9D9")
# A tab's warning when any of its rows takes a measurement waiting to be re-measured (Slice Y3b-2), and its column
REMEASURE_BANNER = ("\u26a0 One or more measurements need to be re-measured by OE. Numbers below use the pre-remesh "
                    "outer dimensions.")
BANNER_FILL = PatternFill("solid", fgColor="FFF2CC")
REMEASURE_COLUMN = "Remeasure Pending"
REMEASURE_WIDTH = 12
PENDING = "\u26a0"
# The workbook's last tab: every month's rows summed per property (Slice Y3b-2)
PROJECT_TOTAL = "Project Total"
NO_ACTIVITY = "No activity"
EMPTY_PROJECT = "No approved IDRs in this project yet."
EMPTY_RANGE = "No approved IDRs in the selected range."
EMPTY_VIEW = "No constructed sites in this view."
CSV = "text/csv; charset=utf-8"
# A tab's footer (ML4b): the Assessment Summary reads only flags approved IDRs claimed
APPROVED_ONLY = "Approved IDRs only"
NOT_ASSIGNED = "not assigned"
# The figures a row (and a totals row) sums: square feet, and the amounts, which are None without a rate
SF_FIELDS = ("four_assessable", "four_non_assessable", "four_total", "seven_assessable", "seven_non_assessable",
             "seven_total", "outside_pl")
AMOUNT_FIELDS = ("four_amount", "seven_amount", "grand_total")


class MonthError(ValueError):
    """A month filter that isn't YYYY-MM, or a range that runs backwards; the route's 400."""


def parse_month(value: Optional[str], name: str) -> Optional[date]:
    """
    Read a month filter.
    Takes the text sent (None or blank: no filter) and its name, for the message.
    Returns the month's first day, or None; raises MonthError for anything but YYYY-MM with a month 01 to 12.
    """
    if value is None or not value.strip():
        return None
    match = MONTH.match(value.strip())
    if not match or not 1 <= int(match.group(2)) <= 12:
        raise MonthError(f"{name} must be YYYY-MM")
    return date(int(match.group(1)), int(match.group(2)), 1)


@dataclass
class SummaryRow:
    """One property (in one month, or over a range of days): its square feet and amounts by pay type, and what Remarks says."""

    property_id: UUID
    address: str
    community_board: Optional[str] = None
    kind: Optional[str] = None
    block: Optional[str] = None
    lot: Optional[str] = None
    violation_no: Optional[str] = None
    four_assessable: Decimal = ZERO
    four_non_assessable: Decimal = ZERO
    four_total: Decimal = ZERO
    seven_assessable: Decimal = ZERO
    seven_non_assessable: Decimal = ZERO
    seven_total: Decimal = ZERO
    # The part of its flags outside the sidewalk strip, clipped at approval and not assessed (Slice Y2)
    outside_pl: Decimal = ZERO
    built_7_billed_4: Decimal = ZERO
    used_approval_date: bool = False
    # Whether any of its flags' measurements waits to be re-measured (Slice Y3b-2)
    remeasure_pending: bool = False
    four_amount: Optional[Decimal] = None
    seven_amount: Optional[Decimal] = None
    grand_total: Optional[Decimal] = None
    remarks: list[str] = field(default_factory=list)


def _day_of(flag: dict[str, Any]) -> tuple[Optional[date], bool]:
    """
    The day a flag counts on: its property's Date Constructed, else the day its IDR was approved in New York.
    Takes the flag row.
    Returns (the day, whether the approval date stood in); (None, True) with neither.
    """
    if flag["date_constructed"] is not None:
        return flag["date_constructed"], False
    if flag["re_signed_at"] is not None:
        return signed_date(flag["re_signed_at"]), True
    return None, True


def _money(sf: Decimal, rate: Optional[Decimal]) -> Optional[Decimal]:
    """
    An amount: square feet at a rate.
    Takes the square feet and the rate (None when not set).
    Returns the amount, or None without a rate.
    """
    return None if rate is None else sf * Decimal(rate)


def _new_row(flag: dict[str, Any]) -> SummaryRow:
    """
    Start a property's row from one of its flags.
    Takes the flag row (with its property's fields).
    Returns the empty row.
    """
    address = display_address_of(flag["property_kind"], flag["address"], flag["intersection_street_1"],
                                 flag["intersection_street_2"], flag["quadrant"], flag["facing_street"])
    return SummaryRow(flag["property_id"], address, flag.get("community_board"), flag["property_kind"],
                      flag.get("block"), flag.get("lot"), flag.get("violation_no"))


def _rows_by(flags: list[dict[str, Any]], rates: dict[str, Optional[Decimal]], first: Optional[date],
             last: Optional[date], key: Callable[[date, dict[str, Any]], Hashable]) -> list[tuple[Hashable, SummaryRow]]:
    """
    Sum claimed flags into rows, measured as the Construction Report measures them, with their amounts and remarks.
    Takes the flags (list_assessment_flags), the project's rates (cost_per_sf_4in, cost_per_sf_7in), the days to keep (first and last, inclusive; None: open) and what makes a row (a function of a flag's day and the flag).
    Returns (key, row) pairs in no set order; amounts None without a rate.
    """
    rows: dict[Hashable, SummaryRow] = {}
    for flag in flags:
        day, fallback = _day_of(flag)
        if day is None or (first and day < first) or (last and day > last):
            continue
        row = rows.get(key(day, flag))
        if row is None:
            row = rows[key(day, flag)] = _new_row(flag)
        row.used_approval_date = row.used_approval_date or fallback
        row.remeasure_pending = row.remeasure_pending or bool(flag.get("needs_remeasure"))
        total, assessable, non_assessable = flag_areas(flag)
        if flag["surveyed_is_assessable"]:
            row.outside_pl += Decimal(flag["quantity"]) - assessable  # the clip, as flag_areas applied it
        if flag["surveyed_thickness"] == "7in":
            row.seven_total += total
            row.seven_assessable += assessable
            row.seven_non_assessable += non_assessable
        else:  # 4", and 7" assessed as 4"
            row.four_total += total
            row.four_assessable += assessable
            row.four_non_assessable += non_assessable
            if flag["surveyed_thickness"] == "7in_as_4in":
                row.built_7_billed_4 += total

    rate_4, rate_7 = rates.get("cost_per_sf_4in"), rates.get("cost_per_sf_7in")
    for row in rows.values():
        row.four_amount, row.seven_amount = _money(row.four_assessable, rate_4), _money(row.seven_assessable, rate_7)
        amounts = [amount for amount in (row.four_amount, row.seven_amount) if amount is not None]
        row.grand_total = sum(amounts, ZERO) if amounts else None
        if row.built_7_billed_4:
            row.remarks.append(f'{row.built_7_billed_4:,.2f} SF built at 7" billed at 4"')
        if rate_4 is None and row.four_total:
            row.remarks.append('4" rate not set')
        if rate_7 is None and row.seven_total:
            row.remarks.append('7" rate not set')
    return list(rows.items())


@dataclass
class SummaryFooter:
    """Who compiled a download and when, and the project's Project Engineers (its REs), for every tab's footer (ML4b)."""

    compiled_by: str
    day: date
    engineers: list[str] = field(default_factory=list)

    def lines(self) -> list[str]:
        """
        The footer's lines.
        Takes nothing.
        Returns 'Compiled by: ...', 'Date: YYYY-MM-DD', 'Project Engineer: ...' (the names, comma-separated, else 'not assigned') and 'Approved IDRs only'.
        """
        return [f"Compiled by: {self.compiled_by}", f"Date: {self.day.isoformat()}",
                f"Project Engineer: {', '.join(self.engineers) or NOT_ASSIGNED}", APPROVED_ONLY]


def person_name(row: dict[str, Any]) -> str:
    """
    A user's name as the footer prints it: first and last name, else the email.
    Takes a row with first_name, last_name and email.
    Returns the name.
    """
    return " ".join(part for part in (row.get("first_name"), row.get("last_name")) if part) or row.get("email") or ""


def project_engineers(roles: list[dict[str, Any]]) -> list[str]:
    """
    The project's Project Engineers: whoever holds its RE role, each once, in the roles' order.
    Takes the project's roles (list_project_roles).
    Returns their names.
    """
    names: list[str] = []
    for row in roles:
        if row.get("role") == "re" and person_name(row) not in names:
            names.append(person_name(row))
    return names


def flags_on_properties(flags: list[dict[str, Any]], property_ids: Optional[set[UUID]]) -> list[dict[str, Any]]:
    """
    Keep the flags of some properties (a task order's bundle, ML1b).
    Takes the flags (list_assessment_flags) and the properties' ids (None: all).
    Returns the flags kept, in order.
    """
    return flags if property_ids is None else [flag for flag in flags if flag["property_id"] in property_ids]


def _month_end(month: Optional[date]) -> Optional[date]:
    """
    A month's last day.
    Takes the month's first day (None: open).
    Returns the next month's first day less one, or None.
    """
    return None if month is None else date(month.year + month.month // 12, month.month % 12 + 1, 1) - timedelta(days=1)


def summary_months(flags: list[dict[str, Any]], rates: dict[str, Optional[Decimal]],
                   first: Optional[date] = None, last: Optional[date] = None) -> dict[date, list[SummaryRow]]:
    """
    Group a project's claimed flags by month and property, measured as the Construction Report measures them.
    Takes the flags (list_assessment_flags), the project's rates (cost_per_sf_4in, cost_per_sf_7in) and the months to keep (first and last, each a month's first day, inclusive; None: open).
    Returns {month: rows}, the months in order and each month's rows by address, any case; amounts None without a rate.
    """
    def by_month(day: date, flag: dict[str, Any]) -> Hashable:
        return day.replace(day=1), flag["property_id"]

    rows = _rows_by(flags, rates, first, _month_end(last), by_month)
    months: dict[date, list[SummaryRow]] = {}
    for (month, _), row in sorted(rows, key=lambda item: (item[0][0], item[1].address.lower())):
        months.setdefault(month, []).append(row)
    return months


def _next_month(month: date) -> date:
    """
    The month after one.
    Takes a month's first day.
    Returns the next month's first day.
    """
    return date(month.year + month.month // 12, month.month % 12 + 1, 1)


def every_month(months: dict[date, list[SummaryRow]], first: Optional[date] = None,
                last: Optional[date] = None) -> dict[date, list[SummaryRow]]:
    """
    Give the monthly workbook a tab for every month of its range (ML1a): from the month asked for, else the first with rows, to the month asked for, else the last with rows; a month without rows gets none (an empty tab).
    Takes the months with rows (summary_months) and the months asked for (first and last, each a month's first day; None: open).
    Returns {month: rows} for every month of the range, in order; {} when neither end is known (no range and no rows).
    """
    start = first or (min(months) if months else last)
    end = last or (max(months) if months else first)
    if start is None or end is None:
        return {}
    filled: dict[date, list[SummaryRow]] = {}
    month = start
    while month <= end:
        filled[month] = months.get(month, [])
        month = _next_month(month)
    return filled


def summary_sites(flags: list[dict[str, Any]], rates: dict[str, Optional[Decimal]], first: Optional[date] = None,
                  last: Optional[date] = None, property_ids: Optional[set[UUID]] = None) -> list[SummaryRow]:
    """
    Sum a project's claimed flags a row per property, over a range of days, measured as the Construction Report measures them.
    Takes the flags (list_assessment_flags), the project's rates, the days to keep (first and last, inclusive; None: open) and the properties to keep (None: all).
    Returns the rows by address, any case; amounts None without a rate.
    """
    kept = [flag for flag in flags if property_ids is None or flag["property_id"] in property_ids]
    rows = _rows_by(kept, rates, first, last, lambda day, flag: flag["property_id"])
    return sorted((row for _, row in rows), key=lambda row: row.address.lower())


def summary_project_total(flags: list[dict[str, Any]], rates: dict[str, Optional[Decimal]],
                          first: Optional[date] = None, last: Optional[date] = None) -> list[SummaryRow]:
    """
    The Project Total tab's rows (Slice Y3b-2): every property the month tabs have, summed over the same months.
    Takes the flags (list_assessment_flags), the project's rates and the months kept (first and last, each a month's first day, inclusive; None: open).
    Returns the rows by address, any case (summary_sites over those months' days).
    """
    return summary_sites(flags, rates, first, _month_end(last))


def project_total_label(first: Optional[date], last: Optional[date]) -> str:
    """
    The Project Total tab's label: the months it sums.
    Takes the months asked for (first and last, each a month's first day; None: open).
    Returns 'Project Total', or with its months ('Project Total, October 2026 to December 2026', 'from' or 'through' one).
    """
    if first and last:
        return f"{PROJECT_TOTAL}, {first:%B %Y} to {last:%B %Y}"
    if first:
        return f"{PROJECT_TOTAL} from {first:%B %Y}"
    if last:
        return f"{PROJECT_TOTAL} through {last:%B %Y}"
    return PROJECT_TOTAL


def total(rows: list[SummaryRow], attribute: str) -> Optional[Decimal]:
    """
    Sum one figure over rows, as a totals row shows it.
    Takes the rows and the attribute.
    Returns the sum of the values that are known; None when none is.
    """
    known = [getattr(row, attribute) for row in rows if getattr(row, attribute) is not None]
    return sum(known, ZERO) if known else None


def _figures(values: Callable[[str], Optional[Decimal]]) -> dict[str, Optional[float]]:
    """
    Round a row's (or the totals') figures for the page: square feet and amounts to 2 decimal places, halves up.
    Takes how to read a figure by its name.
    Returns them keyed by name; an amount without a rate stays None.
    """
    return {name: money(values(name) or ZERO) for name in SF_FIELDS} | {name: money(values(name)) for name in AMOUNT_FIELDS}


def summary_payload(rows: list[SummaryRow], rates: dict[str, Optional[Decimal]], first: Optional[date],
                    last: Optional[date]) -> dict[str, Any]:
    """
    Shape the page's view of the sites for its response (AssessmentSummary).
    Takes the rows (summary_sites), the project's rates and the range asked for.
    Returns {range, cost_per_sf_4in, cost_per_sf_7in, sites, totals}, figures rounded to the cent.
    """
    sites = [{
        "property_id": row.property_id, "kind": row.kind, "display_address": row.address, "block": row.block,
        "lot": row.lot, "violation_no": row.violation_no, "community_board": row.community_board,
        "remarks": row.remarks, "used_approval_date": row.used_approval_date,
        "remeasure_pending": row.remeasure_pending,
        **_figures(lambda name, row=row: getattr(row, name)),
    } for row in rows]
    return {
        "range": {"from": first, "to": last},
        **{name: None if rates.get(name) is None else float(rates[name]) for name in ("cost_per_sf_4in", "cost_per_sf_7in")},
        "sites": sites,
        "totals": {**_figures(lambda name: total(rows, name)),
                   "remeasure_pending": any(row.remeasure_pending for row in rows)},
    }


def range_label(first: Optional[date], last: Optional[date]) -> str:
    """
    Name a range of days, as the tailored download's sheet is named.
    Takes the first and last day (None: open).
    Returns '2026-10-01 to 2026-10-31', 'From 2026-10-01', 'Through 2026-10-31' or 'All dates'.
    """
    if first and last:
        return f"{first.isoformat()} to {last.isoformat()}"
    if first:
        return f"From {first.isoformat()}"
    if last:
        return f"Through {last.isoformat()}"
    return "All dates"


# ---------------------------------------------------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------------------------------------------------

def _rate_label(size: str, rate: Optional[Decimal]) -> str:
    """
    A tab header's rate cell.
    Takes the pay type ('4"' or '7"') and its rate (None when not set).
    Returns '4" rate: $11.54', or the label with a dash.
    """
    return f"{size} rate: " + (DASH if rate is None else f"${Decimal(rate):,.2f}")


def _header(sheet: Worksheet, project: dict[str, Any], label: str) -> None:
    """
    Write a tab's first two rows: the project, and the month (or range) with the two rates.
    Takes the sheet, the project row and the tab's label.
    Returns nothing.
    """
    parts = [project["project_name"] or "", f"Project ID: {project['project_id']}"]
    parts += [f"Reg: {project['registration_code']}"] if project.get("registration_code") else []
    parts += [f"Borough: {project['borough']}"] if project.get("borough") else []
    sheet["A1"] = " · ".join(part for part in parts if part)
    sheet["A1"].font = Font(bold=True, size=12)
    sheet["A2"] = label
    sheet["A2"].font = Font(bold=True)
    sheet["C2"] = _rate_label('4"', project.get("cost_per_sf_4in"))
    sheet["G2"] = _rate_label('7"', project.get("cost_per_sf_7in"))
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width


def _sf(value: Decimal) -> Any:
    """
    A square-feet cell's value: the number, or a dash for none (for scanning the column).
    Takes the square feet.
    Returns what to write.
    """
    return value if value else DASH


def _write_row(sheet: Worksheet, row_number: int, values: list[Any], bold: bool = False) -> None:
    """
    Write one row of a tab, numbers formatted by their column: SF to two decimals, amounts in dollars.
    Takes the sheet, the row's number, its values (the thirteen columns, and the remeasure column's when the tab has it; None leaves a cell blank) and whether it is the totals row.
    Returns nothing.
    """
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column, value)
        if column in MONEY_COLUMNS:
            cell.number_format = MONEY_FORMAT
        elif column in SF_COLUMNS:
            cell.number_format = SF_FORMAT
            if value == DASH:
                cell.alignment = Alignment(horizontal="right")
        if bold:
            cell.font = Font(bold=True)
            cell.fill = TOTALS_FILL


def _footer(sheet: Worksheet, first_row: int, footer: Optional[SummaryFooter]) -> None:
    """
    Write a tab's footer (ML4b), a line a row from first_row down, in column A: Compiled by, Date, Project Engineer and Approved IDRs only.
    Takes the sheet, the row it starts on and the footer (None: none).
    Returns nothing.
    """
    if footer is None:
        return
    for offset, text in enumerate(footer.lines()):
        sheet.cell(first_row + offset, 1, text).font = Font(bold=offset == 0)


def _empty_total(attribute: str, project: dict[str, Any]) -> Optional[Decimal]:
    """
    A figure of an empty month's totals row (ML1a): 0 square feet, and $0 for an amount whose rate is set.
    Takes the attribute and the project row (its rates).
    Returns 0, or None for an amount without a rate.
    """
    if attribute in SF_FIELDS:
        return ZERO
    rates = {"four_amount": [project.get("cost_per_sf_4in")], "seven_amount": [project.get("cost_per_sf_7in")],
             "grand_total": [project.get("cost_per_sf_4in"), project.get("cost_per_sf_7in")]}[attribute]
    return ZERO if any(rate is not None for rate in rates) else None


def _rows_tab(workbook: Workbook, project: dict[str, Any], name: str, label: str, rows: list[SummaryRow],
              footer: Optional[SummaryFooter] = None) -> None:
    """
    Add a tab: the header, a row per property, the totals row and, when any row went by its IDR's approval date, a note saying how many, and when any row has a part outside the property line, a note saying what that column is (Slice Y2). When any row takes a measurement waiting to be re-measured (Slice Y3b-2), a banner across the tab above the column titles and a fourteenth column marking those rows; otherwise neither. A month without rows (ML1a) is the column titles and a totals row of 0. The footer (ML4b) follows a blank row.
    Takes the workbook, the project row, the tab's name and label, its rows and the footer (None: none).
    Returns nothing.
    """
    sheet = workbook.create_sheet(name)
    _header(sheet, project, label)
    pending = any(row.remeasure_pending for row in rows)
    columns = COLUMNS + ((REMEASURE_COLUMN,) if pending else ())
    titles = 3
    if pending:
        banner = sheet.cell(3, 1, REMEASURE_BANNER)
        banner.font = Font(bold=True)
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(columns))
        for column in range(1, len(columns) + 1):
            sheet.cell(3, column).fill = BANNER_FILL
        sheet.column_dimensions[sheet.cell(1, len(columns)).column_letter].width = REMEASURE_WIDTH
        titles = 4
    for column, title in enumerate(columns, start=1):
        sheet.cell(titles, column, title).font = Font(bold=True)
        sheet.cell(titles, column).alignment = Alignment(wrap_text=True, vertical="top")
    first_row = titles + 1
    sheet.freeze_panes = f"A{first_row}"
    for offset, row in enumerate(rows):
        _write_row(sheet, first_row + offset, [
            row.address, row.community_board, _sf(row.four_assessable), _sf(row.four_non_assessable), _sf(row.four_total), row.four_amount,
            _sf(row.seven_assessable), _sf(row.seven_non_assessable), _sf(row.seven_total), row.seven_amount,
            row.grand_total, _sf(row.outside_pl), "; ".join(row.remarks) or None,
        ] + ([PENDING if row.remeasure_pending else None] if pending else []))

    totals_row = first_row + len(rows)

    def summed(attribute: str) -> Optional[Decimal]:
        """A totals-row figure: the rows' sum, or an empty month's 0 (_empty_total)."""
        return total(rows, attribute) if rows else _empty_total(attribute, project)

    _write_row(sheet, totals_row, [
        "Totals", None, summed("four_assessable"), summed("four_non_assessable"), summed("four_total"),
        summed("four_amount"), summed("seven_assessable"), summed("seven_non_assessable"),
        summed("seven_total"), summed("seven_amount"), summed("grand_total"), summed("outside_pl"),
        None,
    ] + ([PENDING] if pending else []), bold=True)
    notes = []
    fallback = sum(row.used_approval_date for row in rows)
    if fallback:
        notes.append(f"{fallback} {'row' if fallback == 1 else 'rows'} used IDR approval date "
                     "(no page-level Date Constructed).")
    if any(row.outside_pl for row in rows):
        notes.append(OUTSIDE_PL_NOTE)
    for offset, text in enumerate(notes):
        sheet.cell(totals_row + 2 + offset, 1, text).font = Font(italic=True)
    _footer(sheet, totals_row + 2 + len(notes) + (1 if notes else 0), footer)


def _empty_tab(workbook: Workbook, project: dict[str, Any], name: str, message: str,
               footer: Optional[SummaryFooter] = None) -> None:
    """
    Add a tab with the header and a line saying why there are no rows, then the footer (ML4b) after a blank row.
    Takes the workbook, the project row, the tab's name (also its label), the line and the footer (None: none).
    Returns nothing.
    """
    sheet = workbook.create_sheet(name)
    _header(sheet, project, name)
    sheet["A4"] = message
    sheet["A4"].font = Font(italic=True)
    _footer(sheet, 6, footer)


def _saved(workbook: Workbook) -> bytes:
    """
    Save a workbook in memory.
    Takes the workbook.
    Returns the .xlsx file's bytes.
    """
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_workbook(project: dict[str, Any], months: dict[date, list[SummaryRow]], filtered: bool,
                   total_rows: list[SummaryRow], total_label: str = PROJECT_TOTAL,
                   footer: Optional[SummaryFooter] = None) -> bytes:
    """
    Build the Assessment Summary workbook: one tab per month (named 'Nov 2026'; every month of its range, ML1a, one without rows an empty tab), or a single 'No activity' tab saying why there are none; then, always last, the Project Total tab (Slice Y3b-2): the same properties summed over the same months, or the same line saying why there are none. Every tab ends with the footer (ML4b).
    Takes the project row, the months (every_month over summary_months), whether a month range was asked for, the Project Total's rows (summary_project_total, over the same months), its label (project_total_label) and the footer (None: none).
    Returns the .xlsx file's bytes.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for month, rows in months.items():
        _rows_tab(workbook, project, month.strftime("%b %Y"), month.strftime("%B %Y"), rows, footer)
    empty = EMPTY_RANGE if filtered else EMPTY_PROJECT
    if not months:
        _empty_tab(workbook, project, NO_ACTIVITY, empty, footer)
    if total_rows:
        _rows_tab(workbook, project, PROJECT_TOTAL, total_label, total_rows, footer)
    else:
        _empty_tab(workbook, project, PROJECT_TOTAL, empty, footer)
    return _saved(workbook)


def build_tailored_workbook(project: dict[str, Any], rows: list[SummaryRow], first: Optional[date],
                            last: Optional[date], footer: Optional[SummaryFooter] = None) -> bytes:
    """
    Build the tailored download: the rows of the page's view on one tab named for its range of days, ending with the footer (ML4b).
    Takes the project row, the rows (summary_sites), the range (first and last; None: open) and the footer (None: none).
    Returns the .xlsx file's bytes; the tab says so when there are no rows.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    name = range_label(first, last)
    if rows:
        _rows_tab(workbook, project, name, name, rows, footer)
    else:
        _empty_tab(workbook, project, name, EMPTY_VIEW, footer)
    return _saved(workbook)


# ---------------------------------------------------------------------------------------------------------------------
# The CSVs (ML1c): the workbooks' rows, flat, for sorting and pivoting; no totals, notes or footer
# ---------------------------------------------------------------------------------------------------------------------

# Remeasure Pending always present, so the column count never changes
CSV_COLUMNS = COLUMNS + (REMEASURE_COLUMN,)
_CSV_FIGURES = ("four_assessable", "four_non_assessable", "four_total", "four_amount", "seven_assessable",
                "seven_non_assessable", "seven_total", "seven_amount", "grand_total", "outside_pl")


def _csv_number(value: Optional[Decimal]) -> str:
    """
    A figure as the CSVs write it: to the cent, halves up, without separators.
    Takes the figure (None: an amount without a rate).
    Returns the text; '' for None.
    """
    return "" if value is None else f"{money(value):.2f}"


def _csv_values(row: SummaryRow) -> list[str]:
    """
    One row's values for a CSV: the workbook's thirteen columns and Remeasure Pending ('Yes' or '').
    Takes the row.
    Returns the values as text.
    """
    return [row.address, row.community_board or "", *(_csv_number(getattr(row, name)) for name in _CSV_FIGURES),
            "; ".join(row.remarks), "Yes" if row.remeasure_pending else ""]


def _csv_bytes(header: list[str], lines: list[list[str]]) -> bytes:
    """
    Write a CSV: UTF-8 with a byte-order mark (so Excel reads its characters), CRLF line ends.
    Takes the header and the lines.
    Returns the file's bytes.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(header)
    writer.writerows(lines)
    return buffer.getvalue().encode("utf-8-sig")


def monthly_csv(months: dict[date, list[SummaryRow]]) -> bytes:
    """
    The monthly workbook as one CSV: a row per month and property, led by its Month (YYYY-MM); an empty month has none.
    Takes the months (as build_workbook takes them).
    Returns the file's bytes (the header alone when there are no rows).
    """
    return _csv_bytes(["Month", *CSV_COLUMNS],
                      [[month.strftime("%Y-%m"), *_csv_values(row)] for month, rows in months.items() for row in rows])


def sites_csv(rows: list[SummaryRow]) -> bytes:
    """
    The tailored workbook as a CSV: a row per site over its range of days.
    Takes the rows (summary_sites).
    Returns the file's bytes (the header alone when there are no rows).
    """
    return _csv_bytes(list(CSV_COLUMNS), [_csv_values(row) for row in rows])


def summary_filename(project_id: str, now: Optional[datetime] = None, tailored: bool = False,
                     extension: str = "xlsx") -> str:
    """
    The download's file name: the project and today's date in New York.
    Takes the project id, for tests the moment it is (now, UTC, by default), whether it is the tailored download and its extension (xlsx or csv).
    Returns '<project_id>_assessment_summary_<YYYY-MM-DD>.xlsx', '..._assessment_summary_tailored_<YYYY-MM-DD>.csv' and so on.
    """
    today = signed_date(now or datetime.now(timezone.utc))
    return f"{project_id}_assessment_summary_{'tailored_' if tailored else ''}{today.isoformat()}.{extension}"
